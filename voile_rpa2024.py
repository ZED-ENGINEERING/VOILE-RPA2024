#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
VOILE RPA 2024 — Ferraillage, ductilité locale et méthode des contraintes (tous étages)
=====================================================================================
Application de bureau (Tkinter) :
  * import des efforts « Pier Forces » d'ETABS (xlsx / xlsm / csv) et choix du pier,
  * enveloppe des moments fléchissants RPA 2024 (éq. 7.22 et 7.23, décalage de hcr),
  * condition d'élancement du voile (RPA 2024 Art. 7.7.4 : hw/lw > 2,0 → enveloppe décalée de hcr ; sinon moments
    du calcul de structure non modifiés),
  * dimensions des voiles (L, e, h) calculées depuis l'export ETABS « Pier Section Properties » (xlsx / csv),
  * enveloppe de l'effort tranchant RPA 2024 (Art. 7.7.2 : 1,4·V en zone critique, retour à V en hw/3),
  * méthode des contraintes sur tous les étages (S.E.C / S.P.C / S.E.T, As, béton, cisaillement)
    avec colonne FACULTATIVE « section adoptée » (nombre et diamètre des barres, au choix de l'ingénieur)
    et calcul de la section choisie (As choisie, comparaison à As requise),
  * moment résistant d'une section rectangulaire en flexion composée (BAEL 91 modifié 99 / RPA 2024 :
    ELU par les trois pivots et ELS par section fissurée homogénéisée),
  * comparaison des cas de calcul (durable / accidentelle / personnalisé),
  * ferraillage des rives et ductilité locale (RPA 2024 Art. 7.7.5) pour le voile / étage choisi,
  * schémas : coupe en plan, détail des rives et ÉLÉVATION du voile avec enveloppe des moments,
  * note de calcul A4 au format PDF ou WORD (.docx), générée à la demande,
  * sauvegarde / ouverture de projets (.json).

Dépendances :  pip install matplotlib reportlab python-docx openpyxl
Lancement   :  python voile_rpa2024.py
Mode console (sans interface) :
    python voile_rpa2024.py --etabs extrait.xlsm --pier V4 --pdf note.pdf --docx note.docx [--projet p.json]

Les formules reprennent fidèlement les classeurs « Ductilite_Locale_Voile_RPA2024_Revision_3.xlsx »,
« Calcul_Voiles_Methode_Contraintes.xlsx », « Enveloppe_Moment_Effort_Tranchant_RPA2024.xlsx » et
« Fiche_Moment_Resistant_BAEL91_RPA2024.xlsx ».
Les écarts volontaires (corrections) sont signalés par « [CORR.] » dans le code.

Unités d'entrée : longueurs voile en m (he, lw, bw, hcr), éléments de rive en cm,
diamètres en mm, efforts en kN / kN.m, contraintes en MPa.
"""
import io
import json
import webbrowser
import math
import os
import sys
import datetime as dt

# ---------------------------------------------------------------------------
#  1. DONNÉES PAR DÉFAUT (issues du classeur Excel)
# ---------------------------------------------------------------------------
DEFAULTS = {
    # --- informations du projet
    "projet": "Projet — à renseigner",
    "element": "Voile V1",
    "niveau": "RDC",
    "auteur": "",
    "date": dt.date.today().strftime("%d/%m/%Y"),
    # --- données générales
    "he": 3.06, "lw": 5.83, "bw": 0.35, "d": 2.5,
    "fc": 30.0, "fe": 500.0, "gb": 1.2, "gs": 1.0,
    "hcr": 5.83, "R": 3.0, "Qf": 1.05, "T0": 0.821, "T2": 0.6,
    "N": 5908.0, "MED": 26641.0, "MRD": 17423.0, "V": 1708.0,
    # --- âme
    "phi_v": 10.0, "st_v": 15.0, "phi_h": 12.0, "st_h": 15.0,
}
for _s in ("g", "d"):
    DEFAULTS.update({
        "phi_l" + _s: 20.0,   # mm  barres verticales de rive
        "lc" + _s: 160.0,     # cm  longueur de l'élément de rive retenue
        "bc" + _s: 35.0,      # cm  largeur de l'élément de rive retenue
        "Avcalc" + _s: 125.6, # cm² section calculée (RDM)
        "phi_c" + _s: 8.0,    # mm  cadres de confinement
        "nbr" + _s: 7,        # -   nombre de branches
        "n_conf" + _s: 36,    # -   nombre de barres confinées (n)
        "bi" + _s: 8.76,      # cm  entraxe des barres confinées
        "L_conf" + _s: 937.8, # cm  longueur cumulée des branches
        "nbc" + _s: 4,        # -   barres sur chaque petite face (schéma)
        "nl" + _s: 0,         # -   barres sur chaque grande face (0 = auto)
    })

# --- données du projet multi-étages (source ETABS, enveloppe, méthode des contraintes)
DEFAULTS.update({
    "pier": "", "etabs_path": "",
    "cas": "Accidentelle", "dref": "h étage", "kV": 1.4, "kd": 0.9, "ktau": 0.2,
    "rho_t": 0.002, "rho_g": 0.0015,
    "env_auto": 1, "env_mbase": 19228.0, "env_msom": 1081.0, "env_ref": "base", "m_mode": "auto",
})
for _i, (_n, _fc, _fe, _gb, _gs, _th) in enumerate(
        [("Durable / transitoire", 30.0, 500.0, 1.5, 1.15, 1.0),
         ("Accidentelle", 30.0, 500.0, 1.2, 1.0, 1.0),
         ("Personnalisé", 30.0, 500.0, 1.2, 1.0, 1.0)], start=1):
    DEFAULTS.update({f"c{_i}_name": _n, f"c{_i}_fc": _fc, f"c{_i}_fe": _fe, f"c{_i}_gb": _gb,
                     f"c{_i}_gs": _gs, f"c{_i}_th": _th})

# --- effort tranchant de calcul pour la méthode des contraintes : « kV » = kV × V2 ETABS (comportement d'origine),
#     « enveloppe » = V corrigé de l'enveloppe de l'effort tranchant (RPA 2024 Art. 7.7.2)
DEFAULTS.update({"v_mode": "kV"})

# --- moment résistant (classeur « Fiche_Moment_Resistant_BAEL91_RPA2024 » — valeurs de l'exemple)
DEFAULTS.update({
    "mr_sit": "Durable / transitoire", "mr_theta": 1.0, "mr_gb_m": 1.5, "mr_gs_m": 1.15,
    "mr_fc": 30.0, "mr_fe": 500.0, "mr_Es": 210000.0, "mr_n": 15.0,
    "mr_b": 0.35, "mr_h": 5.80, "mr_dp": 0.50, "mr_c": 0.50,      # m
    "mr_Ap": 56.28, "mr_A": 56.28,                                 # cm²  (A' sup. / A inf.)
    "mr_elu": "Oui", "mr_Nu": 2359.0,
    "mr_els": "Non", "mr_Ns": 1700.0, "mr_fiss": "Peu préjudiciable", "mr_eta": 1.6,
    "mr_Mu": "", "mr_Ms": "",                                      # vérification facultative (vide = non utilisée)
    "mr_nulim": 0.3, "mr_report": 1,
})
# --- cadres de confinement (choix de l'ingénieur) : type, nombre et longueur unitaire par ligne, pour chaque rive
CADRE_TYPES = ["Cadre fermé", "Épingle transversale", "Épingle longitudinale", "Cadre intermédiaire", "Autre"]
CAD_ROWS = 6
for _s in ("g", "d"):
    for _i in range(1, CAD_ROWS + 1):
        DEFAULTS.update({f"cad_t{_i}{_s}": "", f"cad_n{_i}{_s}": "", f"cad_l{_i}{_s}": ""})
DEFAULTS["cad_auto"] = 1          # envoi automatique de L cumulée vers « Rives et âme »
# --- hypothèses sismiques et matériaux (classeur « Note_de_calcul_Hypotheses » — RPA 2024), liées aux données générales
DEFAULTS.update({
    "seis_link": 1, "mat_link": 1,
    "sz_zone": "Zone VI", "sz_groupe": "Groupe 2", "sz_site": "S3",
    "sz_sys": "6 — Système à ossature à noyau ou à effet noyau",
    "sz_ct": 0.05, "sz_hn": 30.26,
    "sz_struct": "Voile (béton armé / maçonnerie)", "sz_remp": "Dense",
    "mt_fc_fond": 30.0, "mt_fc": 30.0, "mt_dens": 2.5, "mt_fe_ha": 500.0, "mt_fe_tr": 500.0, "mt_fe_doux": 235.0,
    "mt_sit": "Accidentelle (séisme)", "mt_gb_m": 1.2, "mt_gs_m": 1.0,
})
for _i in range(1, 8):
    DEFAULTS[f"sz_pq{_i}"] = "Oui"          # critères de qualité observés ? (Table 3.19)
DEFAULTS["sz_pq5"] = "Non"                  # exemple du classeur : plan irrégulier → QF = 1,05
TEXT_KEYS = ("sz_zone", "sz_groupe", "sz_site", "sz_sys", "sz_struct", "sz_remp", "mt_sit",
             "sz_pq1", "sz_pq2", "sz_pq3", "sz_pq4", "sz_pq5", "sz_pq6", "sz_pq7", "projet", "element", "niveau", "auteur", "date", "pier", "etabs_path", "cas", "dref", "env_ref",
             "m_mode", "c1_name", "c2_name", "c3_name", "v_mode", "mr_sit", "mr_elu", "mr_els", "mr_fiss",
             "mr_Mu", "mr_Ms") + tuple(k for k in DEFAULTS if k.startswith("cad_") and k != "cad_auto")


# ---------------------------------------------------------------------------
#  2. OUTILS
# ---------------------------------------------------------------------------
def xround(x, n=0):
    """ARRONDI Excel (demi supérieur, symétrique)."""
    m = 10 ** n
    return math.copysign(math.floor(abs(x) * m + 0.5) / m, x)


def ceil_(x):
    return math.ceil(x - 1e-9)


def fr(x, nd=2):
    """Nombre au format français."""
    if isinstance(x, str):
        return x
    if x is None:
        return ""
    if isinstance(x, int):
        return str(x)
    return f"{x:,.{nd}f}".replace(",", " ").replace(".", ",")


class Row:
    """Ligne de résultat : libellé, expression, valeurs (1 ou 2 colonnes), unité, statut."""
    def __init__(self, label, expr, v1, v2=None, unit="", ok1=None, ok2=None, nd=2):
        self.label, self.expr, self.v1, self.v2 = label, expr, v1, v2
        self.unit, self.ok1, self.ok2, self.nd = unit, ok1, ok2, nd

    def texts(self):
        return fr(self.v1, self.nd), ("" if self.v2 is None else fr(self.v2, self.nd))


# ---------------------------------------------------------------------------
#  3. MOTEUR DE CALCUL
# ---------------------------------------------------------------------------
def _side(P, s):
    """Calcul complet d'une rive (s = 'g' ou 'd'). Longueurs en cm, diamètres en cm."""
    he, lw, bw, d = P["he"], P["lw"], P["bw"], P["d"]
    fc, fe, gb, gs = P["fc"], P["fe"], P["gb"], P["gs"]
    hcr, N = P["hcr"], P["N"]
    Vd = 1.4 * P["V"]
    phi_l = P["phi_l" + s] / 10.0
    phi_c = P["phi_c" + s] / 10.0
    lc, bc = P["lc" + s], P["bc" + s]
    nbr = int(P["nbr" + s])
    r = {"phi_l": phi_l, "phi_c": phi_c, "lc": lc, "bc": bc, "nbr": nbr}

    # --- dimensions réglementaires (Art. 7.7.5)
    r["lc_min"] = xround(max(0.15 * lw, 1.5 * bw), 2) * 100
    # [CORR.] le classeur utilise MAX(lc,min ; lc,choisi) à gauche et lc,min seul à droite :
    # on applique la même règle (celle de gauche, plus logique) aux deux rives.
    cond = max(r["lc_min"], lc) / 100 > max(2 * bw, 0.2 * lw)
    r["bc_min"] = max(he / 10 if cond else he / 15, 0.2) * 100
    r["ok_lc"] = lc >= r["lc_min"]
    r["ok_bc"] = bc >= r["bc_min"]
    r["ok_phi_l"] = phi_l >= 1.2

    # --- armatures verticales de rive
    nb = int(lc / 20) + 1 + (int(bc / 20.1) if bc > 20 else 0)
    r["Avmin"] = max(0.005 * lc * bc, nb * 2 * math.pi * 0.6 ** 2)
    r["Avcalc"] = P["Avcalc" + s]
    r["ok_Avcalc"] = r["Avcalc"] >= r["Avmin"]
    # [CORR.] Avj = 1,1·Ṽ/fe : Ṽ en kN et fe en MPa → résultat en kN/MPa = 10 cm² (facteur 10 omis dans le classeur)
    r["Avj"] = 1.1 * Vd / fe * 10
    r["Av"] = max(r["Avcalc"], r["Avmin"], r["Avj"])

    # --- confinement
    r["b0"] = bc - 2 * d - phi_c
    r["h0"] = lc - 2 * d - phi_c
    if r["b0"] <= 0 or r["h0"] <= 0:
        raise ValueError("Noyau confiné nul ou négatif : vérifier lc, bc, enrobage et Ø des cadres.")
    r["Ag"] = lc * bc
    r["Ac"] = r["h0"] * r["b0"]
    r["St"] = math.floor(min(r["b0"] / 3, 12.5, 6 * phi_l))
    if r["St"] <= 0:
        raise ValueError("Espacement des cadres nul : vérifier Øl et b0.")
    r["At"] = max(0.09 * r["St"] * r["b0"] * fc / fe,
                  0.3 * r["St"] * r["b0"] * (r["Ag"] / r["Ac"] - 1) * fc / fe)
    r["Atcalc"] = math.pi * phi_c ** 2 / 4 * nbr
    r["ok_At"] = r["Atcalc"] >= r["At"]

    # --- ductilité locale
    n_conf, bi, L = P["n_conf" + s], P["bi" + s], P["L_conf" + s]
    r["n_conf"], r["bi"], r["L"] = n_conf, bi, L
    r["Nnbr"] = math.floor(hcr * 100 / r["St"])
    r["nd"] = N * 1e3 / (bw * lw * 1e6 * fc)
    r["fyd_fcd"] = (fe / fc) * (gb / gs)
    r["Vnc"] = r["h0"] * r["b0"] * hcr * 100
    r["Vst"] = L * (math.pi * phi_c ** 2 / 4) * r["Nnbr"] * r["fyd_fcd"]
    r["w_raw"] = r["Vst"] / r["Vnc"]
    r["wwd"] = max(r["w_raw"], 0.12)
    r["ok_wwd"] = r["w_raw"] > 0.12
    return r


def compute(P):
    """Retourne un dictionnaire de résultats + sections de lignes pour affichage / PDF."""
    he, lw, bw, d = P["he"], P["lw"], P["bw"], P["d"]
    fc, fe = P["fc"], P["fe"]
    hcr = P["hcr"]
    Vd = 1.4 * P["V"]
    G, D = _side(P, "g"), _side(P, "d")

    # --- âme
    phi_v, phi_h = P["phi_v"] / 10, P["phi_h"] / 10
    st_v, st_h = P["st_v"], P["st_h"]
    if st_v <= 0 or st_h <= 0:
        raise ValueError("Espacements de l'âme : valeurs strictement positives requises.")
    W = {}
    W["stmax_v"] = min(25, 25 * phi_v)
    W["stmax_h"] = min(25, 25 * phi_h)       # [CORR.] le classeur utilisait Øv pour les horizontales
    W["phi_min"] = 1.0
    W["phi_max"] = bw * 100 / 8
    # NB : terme « 4·bw » conservé tel quel dans le classeur (bw en m, effet ≈ 1,4 cm).
    W["nv"] = ceil_((lw * 100 - G["lc"] - D["lc"] - 4 * bw) / st_v)
    W["Asv"] = W["nv"] * (math.pi * phi_v ** 2 / 4) * 2
    W["z"] = lw * 100 - G["lc"] / 2 - D["lc"] / 2
    W["Ah"] = st_h * Vd / (W["z"] * fe) * 10
    W["Ahcalc"] = math.pi * phi_h ** 2 / 4 * 2
    W["Ash"] = ceil_(hcr * 100 / st_h) * (math.pi * phi_h ** 2 / 4) * 2
    W["tau_adm"] = 0.2 * fc
    W["tau_b"] = xround(abs(Vd) * 1e-3 / (lw * 0.9 * bw), 2)
    W["ok_stv"] = st_v <= W["stmax_v"]
    W["ok_sth"] = st_h <= W["stmax_h"]
    W["ok_phiv"] = W["phi_min"] <= phi_v <= W["phi_max"]
    W["ok_phih"] = phi_h <= W["phi_max"]
    W["ok_Ah"] = W["Ahcalc"] >= W["Ah"]
    W["ok_tau"] = W["tau_b"] <= W["tau_adm"]

    # --- ductilité (calculs communs puis par rive)
    T0, T2, Rc, Qf = P["T0"], P["T2"], P["R"], P["Qf"]
    if P["MRD"] <= 0 or Qf <= 0 or T0 <= 0:
        raise ValueError("MRD, Qf et T0 doivent être strictement positifs.")
    ratio = Rc / Qf * P["MED"] / P["MRD"]
    m_phi = 2 * ratio - 1 if T0 >= T2 else 1 + 2 * (ratio - 1) * T2 / T0
    eps_sy = fe / 210000
    for r in (G, D):
        r["m_phi"], r["eps_sy"] = m_phi, eps_sy
        r["wn"] = (W["Asv"] / ((lw * 100 - 2 * r["lc"]) * bw * 100)) * r["fyd_fcd"]
        r["alpha_n"] = 1 - (r["n_conf"] * r["bi"] ** 2) / (6 * r["b0"] * r["h0"])
        r["alpha_s"] = (1 - r["St"] / (2 * r["b0"])) * (1 - r["St"] / (2 * r["h0"]))
        r["alpha"] = r["alpha_n"] * r["alpha_s"]
        r["eps_cu"] = 0.0035
        r["awd"] = r["alpha"] * r["wwd"]
        r["awd_min"] = xround(30 * m_phi * (r["nd"] + r["wn"]) * eps_sy * (r["bc"] / r["b0"]) - 0.035, 4)
        r["ok_awd"] = r["awd"] >= r["awd_min"]
        r["eps_cuc"] = 0.0035 + 0.1 * max(r["awd"], r["awd_min"])
        r["c"] = (r["nd"] + r["wn"]) * (lw * 100 * r["bc"] / r["b0"])
        r["lc_calc"] = xround(r["c"] * (1 - r["eps_cu"] / r["eps_cuc"]), 2)
        r["ok_lc_calc"] = r["lc_calc"] <= r["lc"]
        r["ok_lw"] = 2 * max(r["lc_calc"], r["lc"]) < lw * 100
        r["lc_adopt"] = (f"Lc = {fr(r['lc'], 0)} cm" if r["lc_calc"] < r["lc"]
                         else f"Recalculer avec lc ≥ {fr(r['lc_calc'])} cm")

    # --- disposition des barres (schéma)
    for s, r in (("g", G), ("d", D)):
        n = ceil_(round(r["Av"] / (math.pi * r["phi_l"] ** 2 / 4), 3))
        nbc = max(2, int(P["nbc" + s]))
        nl_auto = max(2, ceil_((n - 2 * (nbc - 2)) / 2))
        nl_in = int(P["nl" + s])
        nl = min(40, nl_in if nl_in > 0 else nl_auto)
        nl = max(nl, 2)
        o = d + r["phi_c"] + r["phi_l"] / 2
        el = (r["lc"] - 2 * o) / (nl - 1)
        eb = (r["bc"] - 2 * o) / (nbc - 1)
        r["lay"] = dict(n=n, nbc=nbc, nl=nl, n_ret=2 * nl + 2 * (nbc - 2), o=o, el=el, eb=eb)
        r["lay"]["Av_f"] = r["lay"]["n_ret"] * math.pi * r["phi_l"] ** 2 / 4
        r["lay"]["ok_n"] = r["lay"]["n_ret"] >= n
        r["lay"]["ok_max"] = max(el, eb) <= 20
        r["lay"]["ok_free"] = min(el, eb) - r["phi_l"] >= max(r["phi_l"], 2)

    res = {"G": G, "D": D, "W": W, "Vd": Vd, "m_phi": m_phi, "eps_sy": eps_sy}
    res["sections"] = _sections(P, res)
    res["checks"] = _collect_checks(res)
    res["all_ok"] = all(ok for _, ok in res["checks"])
    return res


def cadres_cumul(P, s):
    """Longueur cumulée des branches (cm) d'une rive (s = 'g' ou 'd') d'après les lignes saisies :
    Σ (nombre × longueur unitaire). Une ligne incomplète est ignorée et signalée dans 'warn'."""
    rows, warn, total = [], [], 0.0
    for i in range(1, CAD_ROWS + 1):
        t = str(P.get(f"cad_t{i}{s}", "") or "").strip()
        n, l = _num(P.get(f"cad_n{i}{s}", "")), _num(P.get(f"cad_l{i}{s}", ""))
        if not t and n is None and l is None:
            continue
        if n is None or l is None or n < 0 or l < 0:
            warn.append(f"ligne {i} : nombre et longueur (≥ 0) requis")
            continue
        sub = n * l
        total += sub
        rows.append({"i": i, "type": t or "—", "n": n, "l": l, "sub": sub})
    return {"rows": rows, "warn": warn, "total": total}


def cadre_longueur_auto(typ, r):
    """Longueur unitaire proposée (cm) = périmètre / portée à l'axe des cadres (b0, h0) + 2 crochets de 10·Øc.
    r : résultats d'une rive (b0, h0, phi_c en cm). None si le type n'a pas de formule (« Autre »…)."""
    hook = 10 * r["phi_c"]
    t = str(typ).strip().lower()
    if t.startswith("cadre fermé"):
        return 2 * (r["b0"] + r["h0"]) + 2 * hook
    if t.startswith("épingle transv"):
        return r["b0"] + 2 * hook
    if t.startswith("épingle long"):
        return r["h0"] + 2 * hook
    return None


def _collect_checks(res):
    G, D, W = res["G"], res["D"], res["W"]
    out = []
    for nm, r in (("rive gauche", G), ("rive droite", D)):
        out += [(f"Ølmin ≥ 12 mm ({nm})", r["ok_phi_l"]),
                (f"lc,choisi ≥ lc,min ({nm})", r["ok_lc"]),
                (f"bc,choisi ≥ bc,min ({nm})", r["ok_bc"]),
                (f"Avcalc ≥ Avmin ({nm})", r["ok_Avcalc"]),
                (f"Atcalc ≥ At ({nm})", r["ok_At"]),
                (f"Taux de confinement ωwd > 0,12 ({nm})", r["ok_wwd"]),
                (f"αωwd ≥ (αωwd)min ({nm})", r["ok_awd"]),
                (f"lc,calcul ≤ lc,choisi ({nm})", r["ok_lc_calc"]),
                (f"2·lc < lw ({nm})", r["ok_lw"]),
                (f"Nombre de barres de rive ≥ n requis ({nm})", r["lay"]["ok_n"]),
                (f"Entraxe des barres de rive ≤ 20 cm ({nm})", r["lay"]["ok_max"]),
                (f"Espacement libre des barres de rive ({nm})", r["lay"]["ok_free"])]
    out += [("Âme : St,v ≤ St,max", W["ok_stv"]), ("Âme : St,h ≤ St,max", W["ok_sth"]),
            ("Âme : Øv dans [Ømin ; bw/8]", W["ok_phiv"]), ("Âme : Øh ≤ bw/8", W["ok_phih"]),
            ("Âme : Ah,calc ≥ Ah", W["ok_Ah"]), ("Âme : τb ≤ τadm", W["ok_tau"])]
    return out


def _ok(b, yes="Vérifiée", no="Non vérifiée"):
    return yes if b else no


def _sections(P, res):
    G, D, W = res["G"], res["D"], res["W"]
    he, lw, bw = P["he"], P["lw"], P["bw"]
    S = []

    S.append(("1. DONNÉES D'ENTRÉE", ("Valeur", ""), [
        Row("he — Hauteur d'étage", "", P["he"], None, "m"),
        Row("lw — Longueur du voile", "", P["lw"], None, "m"),
        Row("bw — Épaisseur du voile", "", P["bw"], None, "m"),
        Row("d — Enrobage des armatures", "", P["d"], None, "cm"),
        Row("fc28 — Résistance du béton", "", P["fc"], None, "MPa", nd=1),
        Row("fe — Limite élastique de l'acier", "", P["fe"], None, "MPa", nd=0),
        Row("γb — Coefficient de sécurité du béton", "", P["gb"], None, "-"),
        Row("γs — Coefficient de sécurité de l'acier", "", P["gs"], None, "-"),
        Row("hcr — Hauteur de la zone critique", "", P["hcr"], None, "m"),
        Row("R — Coefficient de comportement", "", P["R"], None, "-"),
        Row("Qf — Facteur de qualité", "", P["Qf"], None, "-"),
        Row("T0 — Période", "", P["T0"], None, "s", nd=3),
        Row("T2 — Période caractéristique du site", "", P["T2"], None, "s", nd=3),
        Row("N — Effort normal sollicitant", "", P["N"], None, "kN", nd=0),
        Row("MED — Moment sollicitant", "", P["MED"], None, "kN.m", nd=0),
        Row("MRD — Moment résistant", "", P["MRD"], None, "kN.m", nd=0),
        Row("V — Effort tranchant sismique", "", P["V"], None, "kN", nd=0),
        Row("Ṽ — Effort tranchant majoré", "Ṽ = 1,4 × V", res["Vd"], None, "kN", nd=1),
    ]))

    S.append(("2. ÉLÉMENTS DE RIVE — DIMENSIONS ET ARMATURES VERTICALES (Art. 7.7.5)",
              ("Rive gauche", "Rive droite"), [
        Row("Ølmin — Diamètre des barres de rive", "≥ 12 mm", G["phi_l"] * 10, D["phi_l"] * 10, "mm",
            G["ok_phi_l"], D["ok_phi_l"], 0),
        Row("lc,min — Longueur minimale", "max(0,15·lw ; 1,5·bw)", G["lc_min"], D["lc_min"], "cm", nd=0),
        Row("lc,choisi — Longueur retenue", "lc ≥ lc,min", G["lc"], D["lc"], "cm", G["ok_lc"], D["ok_lc"], 0),
        Row("bc,min — Largeur minimale", "he/10 ou he/15 (≥ 20 cm)", G["bc_min"], D["bc_min"], "cm"),
        Row("bc,choisi — Largeur retenue", "bc ≥ bc,min", G["bc"], D["bc"], "cm", G["ok_bc"], D["ok_bc"], 1),
        Row("Avmin — Section minimale", "max(0,5 % lc·bc ; barres tous les 20 cm)", G["Avmin"], D["Avmin"], "cm²"),
        Row("Avcalc — Section calculée (RDM)", "Avcalc ≥ Avmin", G["Avcalc"], D["Avcalc"], "cm²", G["ok_Avcalc"], D["ok_Avcalc"]),
        Row("Avj — Section requise par l'effort tranchant", "1,1·Ṽ / fe", G["Avj"], D["Avj"], "cm²"),
        Row("Av — Section verticale de rive retenue", "max(Avcalc ; Avmin ; Avj)", G["Av"], D["Av"], "cm²"),
    ]))

    S.append(("3. CONFINEMENT DES ÉLÉMENTS DE RIVE", ("Rive gauche", "Rive droite"), [
        Row("Øc — Diamètre des cadres", "", G["phi_c"] * 10, D["phi_c"] * 10, "mm", nd=0),
        Row("Nbr — Nombre de branches", "", G["nbr"], D["nbr"], "-"),
        Row("b0 — Largeur du noyau confiné", "bc − 2d − Øc", G["b0"], D["b0"], "cm"),
        Row("h0 — Longueur du noyau confiné", "lc − 2d − Øc", G["h0"], D["h0"], "cm"),
        Row("Ag — Section brute de la rive", "lc × bc", G["Ag"], D["Ag"], "cm²", nd=0),
        Row("Ac — Section confinée (noyau)", "h0 × b0", G["Ac"], D["Ac"], "cm²", nd=0),
        Row("St,max — Espacement des cadres", "⌊min(b0/3 ; 12,5 ; 6·Ølmin)⌋", G["St"], D["St"], "cm", nd=0),
        Row("At — Section de confinement requise",
            "max(0,09·St·b0·fc/fe ; 0,3·St·b0·(Ag/Ac−1)·fc/fe)", G["At"], D["At"], "cm²"),
        Row("Atcalc — Section de confinement fournie", "π·Øc²/4 × Nbr ≥ At", G["Atcalc"], D["Atcalc"], "cm²",
            G["ok_At"], D["ok_At"]),
    ]))

    S.append(("4. ÂME (PARTIE COURANTE DU VOILE)", ("Verticales", "Horizontales"), [
        Row("Ø — Diamètre des barres", "", P["phi_v"], P["phi_h"], "mm", nd=0),
        Row("St,choisi — Espacement retenu", "", P["st_v"], P["st_h"], "cm", nd=1),
        Row("St,max — Espacement maximal", "min(25 ; 25·Ø)", W["stmax_v"], W["stmax_h"], "cm", W["ok_stv"], W["ok_sth"], 1),
        Row("Ømin / Ømax — Diamètres limites", "10 mm ≤ Ø ≤ bw/8", W["phi_max"] * 10, W["phi_max"] * 10, "mm",
            W["ok_phiv"], W["ok_phih"], 1),
        Row("nv — Nombre de barres verticales par face", "⌈(lw − lcg − lcd − 4·bw)/St⌉", W["nv"], None, "-"),
        Row("Asv — Section verticale de l'âme fournie", "nv · π·Ø²/4 × 2 faces", W["Asv"], None, "cm²"),
        Row("z — Bras de levier", "lw − lcg/2 − lcd/2", W["z"], None, "cm", nd=1),
        Row("Ah — Section horizontale requise (effort tranchant)", "St·Ṽ / (z·fe)", None, W["Ah"], "cm²"),
        Row("Ah,calc — Section fournie par nappe", "π·Ø²/4 × 2 ≥ Ah", None, W["Ahcalc"], "cm²", None, W["ok_Ah"]),
        Row("Ash — Section horizontale totale sur hcr", "⌈hcr/St⌉·π·Ø²/4 × 2", None, W["Ash"], "cm²"),
        Row("τadm — Contrainte admissible", "0,2·fc28", W["tau_adm"], None, "MPa"),
        Row("τb — Contrainte de cisaillement", "Ṽ / (0,9·lw·bw)", W["tau_b"], None, "MPa", W["ok_tau"]),
    ]))

    S.append(("5. DUCTILITÉ LOCALE — PARAMÈTRES DE CONFINEMENT", ("Rive gauche", "Rive droite"), [
        Row("n — Nombre de barres confinées", "donnée", G["n_conf"], D["n_conf"], "-"),
        Row("bi — Entraxe des barres confinées", "donnée", G["bi"], D["bi"], "cm"),
        Row("L — Longueur cumulée des branches", "donnée", G["L"], D["L"], "cm", nd=1),
        Row("Nnbr — Nombre de nappes sur hcr", "⌊hcr/St⌋", G["Nnbr"], D["Nnbr"], "-"),
        Row("nd — Effort normal réduit", "N / (bw·lw·fc28)", G["nd"], D["nd"], "-", nd=4),
        Row("Fyd/Fcd — Rapport des résistances", "(fe/fc28)·(γb/γs)", G["fyd_fcd"], D["fyd_fcd"], "-"),
        Row("Vnc — Volume du noyau confiné", "h0·b0·hcr", G["Vnc"], D["Vnc"], "cm³", nd=0),
        Row("Vst — Volume équivalent d'acier", "L·(π·Øc²/4)·Nnbr·Fyd/Fcd", G["Vst"], D["Vst"], "cm³", nd=0),
        Row("ωwd — Taux mécanique de confinement", "max(Vst/Vnc ; 0,12), exigé > 0,12", G["wwd"], D["wwd"], "-",
            G["ok_wwd"], D["ok_wwd"], 4),
        Row("ωn — Taux mécanique de l'âme", "Asv/((lw−2lc)·bw)·Fyd/Fcd", G["wn"], D["wn"], "-", nd=4),
    ]))

    S.append(("6. DUCTILITÉ EN COURBURE ET EFFICACITÉ DU CONFINEMENT", ("Rive gauche", "Rive droite"), [
        Row("mφ — Ductilité en courbure demandée", "T0 ≥ T2 : 2·(R/Qf)·MED/MRD − 1", G["m_phi"], D["m_phi"], "-"),
        Row("εsy — Déformation élastique de l'acier", "fe / 210000", G["eps_sy"], D["eps_sy"], "-", nd=5),
        Row("αn — Efficacité transversale", "1 − n·bi²/(6·b0·h0)", G["alpha_n"], D["alpha_n"], "-", nd=4),
        Row("αs — Efficacité longitudinale", "(1−St/2b0)·(1−St/2h0)", G["alpha_s"], D["alpha_s"], "-", nd=4),
        Row("α — Efficacité globale", "αn × αs", G["alpha"], D["alpha"], "-", nd=4),
        Row("εcu — Déformation ultime du béton", "", G["eps_cu"], D["eps_cu"], "-", nd=4),
    ]))

    S.append(("7. LONGUEUR DE CONFINEMENT ET CONCLUSION", ("Rive gauche", "Rive droite"), [
        Row("αωwd — Taux de confinement effectif", "α × ωwd", G["awd"], D["awd"], "-", nd=4),
        Row("(αωwd)min — Taux minimal requis", "30·mφ·(nd+ωn)·εsy·bc/b0 − 0,035", G["awd_min"], D["awd_min"], "-",
            nd=4),
        Row("Vérification αωwd ≥ (αωwd)min", "", _ok(G["ok_awd"]), _ok(D["ok_awd"]), "", G["ok_awd"], D["ok_awd"]),
        Row("εcu,c — Déformation ultime du béton confiné", "0,0035 + 0,1·max(αωwd ; (αωwd)min)", G["eps_cuc"],
            D["eps_cuc"], "-", nd=5),
        Row("c — Profondeur de l'axe neutre", "(nd+ωn)·lw·bc/b0", G["c"], D["c"], "cm"),
        Row("lc,calcul — Longueur de confinement calculée", "c·(1 − εcu/εcu,c)", G["lc_calc"], D["lc_calc"], "cm"),
        Row("Vérification lc,calcul ≤ lc,choisi", "", _ok(G["ok_lc_calc"]), _ok(D["ok_lc_calc"]), "",
            G["ok_lc_calc"], D["ok_lc_calc"]),
        Row("Longueur de confinement à adopter", "", G["lc_adopt"], D["lc_adopt"], ""),
        Row("Vérification 2·lc < lw", "", _ok(G["ok_lw"]), _ok(D["ok_lw"]), "", G["ok_lw"], D["ok_lw"]),
    ]))

    S.append(("8. DISPOSITION DES BARRES DE RIVE (SCHÉMA)", ("Rive gauche", "Rive droite"), [
        Row("n — Nombre minimal de barres (Av / section d'une barre)", "⌈Av / (π·Øl²/4)⌉", G["lay"]["n"],
            D["lay"]["n"], "-"),
        Row("nbc — Barres par petite face (coins inclus)", "saisie", G["lay"]["nbc"], D["lay"]["nbc"], "-"),
        Row("nl — Barres par grande face (coins inclus)", "auto : ⌈(n − 2(nbc−2))/2⌉", G["lay"]["nl"],
            D["lay"]["nl"], "-"),
        Row("n retenu", "2·nl + 2·(nbc − 2) ≥ n", G["lay"]["n_ret"], D["lay"]["n_ret"], "-", G["lay"]["ok_n"],
            D["lay"]["ok_n"]),
        Row("Av,fournie", "n retenu × π·Øl²/4", G["lay"]["Av_f"], D["lay"]["Av_f"], "cm²"),
        Row("o — Distance face béton → axe des barres", "d + Øc + Øl/2", G["lay"]["o"], D["lay"]["o"], "cm"),
        Row("el — Entraxe selon lc", "(lc − 2o)/(nl − 1)", G["lay"]["el"], D["lay"]["el"], "cm"),
        Row("eb — Entraxe selon bc", "(bc − 2o)/(nbc − 1)", G["lay"]["eb"], D["lay"]["eb"], "cm"),
        Row("Vérification entraxe maximal", "≤ 20 cm", _ok(G["lay"]["ok_max"]), _ok(D["lay"]["ok_max"]), "",
            G["lay"]["ok_max"], D["lay"]["ok_max"]),
        Row("Vérification espacement libre", "≥ max(Øl ; 2 cm)", _ok(G["lay"]["ok_free"]), _ok(D["lay"]["ok_free"]),
            "", G["lay"]["ok_free"], D["lay"]["ok_free"]),
    ]))
    return S


# ---------------------------------------------------------------------------
#  4. SCHÉMA DE FERRAILLAGE (matplotlib)
# ---------------------------------------------------------------------------
def draw_ductile_elevation(ax, P, res):
    """Élévation du voile sur la zone ductile (hauteur hcr) : rives confinées (barres + cadres) et âme
    (barres verticales et horizontales), avec cotes lcg / âme / lcd et légendes. Unités de tracé : cm."""
    from matplotlib.patches import Rectangle
    G, D, W = res["G"], res["D"], res["W"]
    lw, hz = P["lw"] * 100, P["hcr"] * 100
    lcg, lcd = G["lc"], D["lc"]
    web = lw - lcg - lcd
    NAVY, ORANGE, RED = "#1f4e79", "#f0a020", "#c0392b"
    ax.set_title(f"Élévation du voile — zone ductile (hcr = {fr(P['hcr'])} m)", fontsize=9, fontweight="bold")
    ax.add_patch(Rectangle((lcg, 0), max(web, 0), hz, fc="white", ec="none", zorder=1))
    # --- âme : barres horizontales (rouge) puis verticales (noir)
    st_h = max(P["st_h"], 1.0)
    nh = int(hz // st_h) + 1
    for j in range(min(nh, 120)):
        y = (j + 0.5) * st_h
        if y < hz and web > 0:
            ax.plot([lcg, lcg + web], [y, y], color=RED, lw=1.0, zorder=2, solid_capstyle="butt")
    nv = min(max(W["nv"], 1), 80)
    xs_v = [lcg + (i + 0.5) * web / W["nv"] for i in range(nv)] if web > 0 and W["nv"] > 0 else []
    for x in xs_v:
        ax.plot([x, x], [0, hz], color="#111111", lw=0.6, zorder=3)
    # --- rives : fond sombre, barres verticales (traits blancs) et cadres (traits fins)
    for x0, r in ((0.0, G), (lw - lcd, D)):
        ax.add_patch(Rectangle((x0, 0), r["lc"], hz, fc="#2b2b2b", ec="none", zorder=4))
        L = r["lay"]
        for i in range(L["nl"]):
            x = x0 + L["o"] + i * L["el"]
            ax.plot([x, x], [0, hz], color="white", lw=0.7, ls=(0, (4, 3)), zorder=5)
        St = max(r["St"], 1)
        for k in range(min(int(hz // St) + 1, 200)):
            y = (k + 0.5) * St
            if y < hz:
                ax.plot([x0, x0 + r["lc"]], [y, y], color="white", lw=0.35, alpha=0.55, zorder=5)
    ax.add_patch(Rectangle((0, 0), lw, hz, fc="none", ec="#333333", lw=1.6, zorder=6))
    # --- cotes
    yd = -0.11 * lw
    for xa, xb, txt in ((0, lcg, f"lcg = {fr(lcg, 0)} cm"), (lcg, lw - lcd, f"Âme = {fr(web, 0)} cm"),
                        (lw - lcd, lw, f"lcd = {fr(lcd, 0)} cm")):
        ax.annotate("", (xa, yd), (xb, yd), arrowprops=dict(arrowstyle="<->", color=NAVY, lw=1.1))
        ax.text((xa + xb) / 2, yd - 0.025 * lw, txt, ha="center", va="top", fontsize=7.5, color=NAVY, fontweight="bold")
    ax.annotate("", (-0.05 * lw, 0), (-0.05 * lw, hz), arrowprops=dict(arrowstyle="<->", color=RED, lw=1.0))
    ax.text(-0.07 * lw, hz / 2, f"hcr = {fr(P['hcr'])} m", rotation=90, ha="right", va="center", fontsize=7,
            color=RED, fontweight="bold")

    def label(txt, xy, xyt, ha="center", va="bottom"):
        ax.annotate(txt, xy=xy, xytext=xyt, ha=ha, va=va, fontsize=7.5, color=NAVY, fontweight="bold", zorder=8,
                    arrowprops=dict(arrowstyle="-|>", color=ORANGE, lw=1.3, shrinkA=1, shrinkB=0))
    ytop = hz + 0.17 * lw
    label(f"Rive G : {G['lay']['n_ret']} Ø{fr(G['phi_l'] * 10, 0)}", (lcg / 2, hz), (lcg / 2, ytop))
    label(f"Aciers verticaux de\nl'âme Ø{fr(P['phi_v'], 0)}", (lcg + web / 2, hz), (lcg + web / 2, ytop))
    label(f"Rive D : {D['lay']['n_ret']} Ø{fr(D['phi_l'] * 10, 0)}", (lw - lcd / 2, hz), (lw - lcd / 2, ytop))
    yh = hz * 0.8
    label(f"Aciers horizontaux\nØ{fr(P['phi_h'], 0)} — St = {fr(P['st_h'], 0)} cm", (lw - lcd, yh),
          (lw + 0.04 * lw, yh), ha="left", va="center")
    label(f"Cadres Ø{fr(G['phi_c'] * 10, 0)}\nSt = {fr(G['St'], 0)} cm", (0, hz * 0.2), (-0.12 * lw, hz * 0.2),
          ha="right", va="center")
    # espacements des barres de l'âme : St vertical (en tête) et St horizontal (au milieu)
    if len(xs_v) >= 2:
        ya = hz + 0.045 * lw
        ax.annotate("", (xs_v[1], ya), (xs_v[0], ya), arrowprops=dict(arrowstyle="<|-|>", color=NAVY, lw=0.9))
        ax.text(xs_v[0] - 0.01 * lw, ya + 0.015 * lw, f"St = {fr(P['st_v'], 0)} cm", ha="right", va="bottom",
                fontsize=7, color=NAVY, fontweight="bold")
    if web > 0 and st_h < hz:
        xm, ym = lcg + web * 0.55, hz * 0.45
        y1 = (int(ym // st_h) + 0.5) * st_h
        ax.annotate("", (xm, y1 + st_h), (xm, y1), arrowprops=dict(arrowstyle="<|-|>", color=NAVY, lw=0.9), zorder=9)
        ax.text(xm + 0.02 * lw, y1 + st_h / 2, f"St = {fr(P['st_h'], 0)} cm", ha="left", va="center", fontsize=7,
                color=NAVY, fontweight="bold", zorder=9, bbox=dict(fc="white", ec="none", pad=1.5))
    ax.set_xlim(-0.30 * lw, 1.32 * lw)
    ax.set_ylim(-0.24 * lw, hz + 0.30 * lw)
    ax.set_aspect("equal", adjustable="box")
    ax.axis("off")


def draw_schema(fig, P, res, layout="portrait"):
    from matplotlib.patches import Rectangle
    fig.clear()
    lw, bw, d = P["lw"] * 100, P["bw"] * 100, P["d"]
    G, D, W = res["G"], res["D"], res["W"]
    if layout == "landscape":              # écran : élévation à gauche, plan et rives à droite
        gs = fig.add_gridspec(3, 2, width_ratios=[1.35, 1.0], height_ratios=[0.8, 1.15, 1.15], hspace=0.25,
                              wspace=0.04, left=0.03, right=0.99, top=0.95, bottom=0.02)
        ax_el, ax_plan, row_r = fig.add_subplot(gs[:, 0]), fig.add_subplot(gs[0, 1]), (lambda c: gs[1 + c, 1])
    else:                                  # note de calcul : empilé
        gs = fig.add_gridspec(4, 1, height_ratios=[1.75, 0.7, 1.0, 1.0], hspace=0.22,
                              left=0.04, right=0.98, top=0.97, bottom=0.02)
        ax_el, ax_plan, row_r = fig.add_subplot(gs[0, 0]), fig.add_subplot(gs[1, 0]), (lambda c: gs[2 + c, 0])
    draw_ductile_elevation(ax_el, P, res)
    # ---------- coupe en plan
    ax = ax_plan
    ax.set_title("Coupe en plan du voile — zones de rive et âme", fontsize=9, fontweight="bold")
    ax.add_patch(Rectangle((0, 0), lw, bw, fc="#eef0f2", ec="black", lw=1))
    for (x0, r, c) in ((0, G, "#cfe3f7"), (lw - D["lc"], D, "#cfe3f7")):
        ax.add_patch(Rectangle((x0, 0), r["lc"], bw, fc=c, ec="#1f4e79", lw=1))
        ax.add_patch(Rectangle((x0 + d, d), r["lc"] - 2 * d, bw - 2 * d, fc="none", ec="#c0392b", lw=0.9))
    # barres de rive
    for x0, r in ((0, G), (lw - D["lc"], D)):
        L = r["lay"]
        xs = [x0 + L["o"] + i * L["el"] for i in range(L["nl"])]
        ax.scatter(xs + xs, [L["o"]] * len(xs) + [r["bc"] - L["o"]] * len(xs), s=6, c="#111", zorder=3)
        ys = [L["o"] + j * L["eb"] for j in range(1, L["nbc"] - 1)]
        for xx in (x0 + L["o"], x0 + r["lc"] - L["o"]):
            ax.scatter([xx] * len(ys), ys, s=6, c="#111", zorder=3)
    # barres d'âme
    nv = W["nv"]
    web = lw - G["lc"] - D["lc"]
    if web > 0 and nv > 0:
        xs = [G["lc"] + (i + 0.5) * web / nv for i in range(min(nv, 40))]
        ov = d + P["phi_h"] / 10 + P["phi_v"] / 20
        ax.scatter(xs + xs, [ov] * len(xs) + [bw - ov] * len(xs), s=4, c="#2e7d32", zorder=3)
    ax.set_xlim(-0.03 * lw, 1.03 * lw)
    ax.set_ylim(-0.8 * bw, 1.9 * bw)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.annotate("", (0, -0.35 * bw), (lw, -0.35 * bw), arrowprops=dict(arrowstyle="<->", lw=0.7))
    ax.text(lw / 2, -0.75 * bw, f"lw = {fr(lw, 0)} cm   (âme : {W['nv']} Ø{fr(P['phi_v'], 0)} /face, "
            f"St = {fr(P['st_v'], 0)} cm ; HA Ø{fr(P['phi_h'], 0)} / {fr(P['st_h'], 0)} cm)",
            ha="center", va="top", fontsize=6.5)
    ax.text(G["lc"] / 2, 1.15 * bw, f"lcg={fr(G['lc'], 0)}", ha="center", fontsize=6.5, color="#1f4e79")
    ax.text(lw - D["lc"] / 2, 1.15 * bw, f"lcd={fr(D['lc'], 0)}", ha="center", fontsize=6.5, color="#1f4e79")
    ax.text(lw / 2, 1.15 * bw, "âme", ha="center", fontsize=6.5, color="#2e7d32")
    ax.text(lw / 2, 1.7 * bw, "Barres non à l'échelle (taille fixe)", ha="center", fontsize=5.5,
            color="gray", style="italic")

    # ---------- détail des rives
    for col, (nm, r) in enumerate((("RIVE GAUCHE", G), ("RIVE DROITE", D))):
        a = fig.add_subplot(row_r(col))
        L = r["lay"]
        lc, bc = r["lc"], r["bc"]
        a.add_patch(Rectangle((0, 0), lc, bc, fc="#eef0f2", ec="black", lw=1.2))
        a.add_patch(Rectangle((d, d), lc - 2 * d, bc - 2 * d, fc="none", ec="#c0392b", lw=1.4))
        # branches intermédiaires (Nbr − 2) accrochées aux barres des grandes faces
        nint = max(0, r["nbr"] - 2)
        if nint and L["nl"] > 2:
            idx = sorted({int(round(k)) for k in
                          [1 + i * (L["nl"] - 3) / max(nint - 1, 1) for i in range(nint)]}) if nint > 1 else [L["nl"] // 2]
            for i in idx:
                x = L["o"] + i * L["el"]
                a.plot([x, x], [d, bc - d], color="#c0392b", lw=0.9)
        xs = [L["o"] + i * L["el"] for i in range(L["nl"])]
        ys_b = [L["o"] + j * L["eb"] for j in range(1, L["nbc"] - 1)]
        rad = r["phi_l"] / 2
        from matplotlib.patches import Circle
        for x in xs:
            for y in (L["o"], bc - L["o"]):
                a.add_patch(Circle((x, y), rad, fc="#111", ec="none", zorder=3))
        for x in (L["o"], lc - L["o"]):
            for y in ys_b:
                a.add_patch(Circle((x, y), rad, fc="#111", ec="none", zorder=3))
        a.set_xlim(-0.08 * lc, 1.08 * lc)
        a.set_ylim(-0.55 * bc, 1.45 * bc)
        a.set_aspect("equal")
        a.axis("off")
        a.set_title(f"{nm} — {L['n_ret']} HA{fr(r['phi_l'] * 10, 0)}", fontsize=8.5, fontweight="bold")
        a.annotate("", (0, -0.2 * bc), (lc, -0.2 * bc), arrowprops=dict(arrowstyle="<->", lw=0.7))
        a.text(lc / 2, -0.28 * bc, f"lc = {fr(lc, 0)} cm", ha="center", va="top", fontsize=7)
        a.annotate("", (-0.04 * lc, 0), (-0.04 * lc, bc), arrowprops=dict(arrowstyle="<->", lw=0.7))
        a.text(-0.06 * lc, bc / 2, f"bc = {fr(bc, 0)}", ha="right", va="center", fontsize=7, rotation=90)
        a.text(lc / 2, 1.1 * bc,
               f"Cadres HA{fr(r['phi_c'] * 10, 0)} St = {fr(r['St'], 0)} cm — {r['nbr']} branches", ha="center",
               fontsize=7, color="#c0392b")
        a.text(lc / 2, -0.62 * bc,
               f"lcg : 2×{L['nl']}  |  bcg : 2×{L['nbc'] - 2}  (e = {fr(L['el'])} / {fr(L['eb'])} cm)",
               ha="center", fontsize=6.5)
    return fig




# ---------------------------------------------------------------------------
#  4b. SOURCE DES EFFORTS : EXTRAIT « PIER FORCES » D'ETABS
# ---------------------------------------------------------------------------
import re
import csv

_ALIASES = {
    "story": ("story", "niveau", "etage", "étage"),
    "pier": ("pier", "voile", "trumeau"),
    "load": ("load", "loadcase", "load case/combo", "output case", "combo", "case"),
    "loc": ("loc", "location"),
    "P": ("p",), "V2": ("v2",), "V3": ("v3",), "T": ("t",), "M2": ("m2",), "M3": ("m3",),
    # feuille de synthèse
    "Pmin_s": ("p min", "pmin"), "V2_s": ("v2 (abs max)", "v2 abs max"),
    "M3_s": ("m3 (abs max)", "m3 abs max"), "Pc_s": ("p correspondant (m3 abs max)", "p correspondant"),
}


def _norm(x):
    return re.sub(r"\s+", " ", str(x if x is not None else "")).strip().lower()


def _num(x):
    try:
        return float(str(x).replace(",", ".").replace(" ", ""))
    except ValueError:
        return None


def story_elevation(name, order):
    """Cote déduite du nom du niveau (« NIV+27.2 » → 27.2). À défaut : rang d'apparition."""
    m = re.findall(r"[+-]\s*\d+(?:[.,]\d+)?", str(name))
    if m:
        v = _num(m[-1].replace(" ", ""))
        if v is not None:
            return v
    m = re.findall(r"\d+(?:[.,]\d+)?", str(name))
    return _num(m[-1]) if m else float(order)


def _rows_from_table(table):
    """table : liste de lignes (listes). Détecte l'en-tête et renvoie ('raw'|'synth', rows) ou None."""
    for hi, hdr in enumerate(table[:25]):
        h = [_norm(c) for c in hdr]
        if "pier" not in h:
            continue
        pi = h.index("pier")
        cols = {}
        # niveau : colonne « story » ou, à défaut, celle qui précède « pier »
        for i, c in enumerate(h):
            if c in _ALIASES["story"]:
                cols["story"] = i
        cols.setdefault("story", max(pi - 1, 0))
        cols["pier"] = pi
        for key in ("load", "loc", "P", "V2", "V3", "T", "M2", "M3", "Pmin_s", "V2_s", "M3_s", "Pc_s"):
            for i, c in enumerate(h):
                if c in _ALIASES[key]:
                    cols.setdefault(key, i)
        raw = all(k in cols for k in ("P", "V2", "M3"))
        synth = all(k in cols for k in ("Pmin_s", "V2_s", "M3_s", "Pc_s"))
        if not (raw or synth):
            continue
        out = []
        for r in table[hi + 1:]:
            if len(r) <= pi or r[pi] in (None, ""):
                continue
            st = str(r[cols["story"]]).strip() if r[cols["story"]] is not None else ""
            if not st:
                continue
            if raw:
                d = {"story": st, "pier": str(r[pi]).strip(),
                     "load": str(r[cols["load"]]).strip() if "load" in cols else "",
                     "loc": str(r[cols["loc"]]).strip() if "loc" in cols else "Bottom"}
                ok = True
                for k in ("P", "V2", "V3", "T", "M2", "M3"):
                    v = _num(r[cols[k]]) if k in cols else 0.0
                    if v is None:
                        ok = False
                        break
                    d[k] = v
                if ok:
                    out.append(d)
            else:
                vals = [_num(r[cols[k]]) for k in ("Pmin_s", "V2_s", "M3_s", "Pc_s")]
                if None in vals:
                    continue
                out.append({"story": st, "pier": str(r[pi]).strip(), "Pmin": vals[0], "V2": vals[1],
                            "M3": vals[2], "Pc": vals[3], "load": "synthèse"})
        if out:
            return ("raw" if raw else "synth"), out
    return None


def load_etabs(path):
    """Lit un export ETABS (xlsx / xlsm / csv). Retourne {'kind','rows','sheet','piers','stories'}."""
    ext = os.path.splitext(path)[1].lower()
    found = []
    if ext in (".csv", ".txt", ".tsv"):
        with open(path, encoding="utf-8-sig", errors="replace") as fh:
            sample = fh.read(4096)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            except csv.Error:
                dialect = csv.excel
            found.append((os.path.basename(path), _rows_from_table([list(r) for r in csv.reader(fh, dialect)])))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            tbl = [list(r) for r in ws.iter_rows(values_only=True)]
            found.append((ws.title, _rows_from_table(tbl)))
        wb.close()
    best = None
    for name, res in found:
        if res and (best is None or (res[0] == "raw" and best[1][0] != "raw")):
            best = (name, res)
    if best is None:
        raise ValueError("Aucune table « Pier Forces » reconnue (colonnes attendues : Story, Pier, "
                         "Load, Loc, P, V2, V3, T, M2, M3).")
    sheet, (kind, rows) = best
    order = {}
    for r in rows:
        order.setdefault(r["story"], len(order))
    stories = sorted(order, key=lambda s: story_elevation(s, order[s]))
    piers = []
    for r in rows:
        if r["pier"] not in piers:
            piers.append(r["pier"])
    return {"kind": kind, "rows": rows, "sheet": sheet, "piers": piers, "stories": stories}


def etabs_governing(data, pier):
    """Efforts déterminants du pier, par niveau (section du bas), comme la feuille « Synthese_Bottom » :
    P min ; V2 de valeur absolue maximale ; M3 de valeur absolue maximale et P correspondant."""
    rows = [r for r in data["rows"] if r["pier"] == pier]
    if not rows:
        raise ValueError(f"Pier « {pier} » absent du fichier.")
    order = {}
    for r in data["rows"]:
        order.setdefault(r["story"], len(order))
    out = []
    for st in sorted({r["story"] for r in rows}, key=lambda s: story_elevation(s, order[s])):
        rs = [r for r in rows if r["story"] == st]
        if data["kind"] == "raw":
            bot = [r for r in rs if _norm(r["loc"]) in ("bottom", "bot", "bas", "")] or rs
            m3 = max(bot, key=lambda r: abs(r["M3"]))
            v2 = max(bot, key=lambda r: abs(r["V2"]))
            out.append({"story": st, "elev": story_elevation(st, order[st]), "Pmin": min(r["P"] for r in bot),
                        "V2": v2["V2"], "M3": m3["M3"], "Pc": m3["P"], "combo_M3": m3["load"],
                        "combo_V2": v2["load"]})
        else:
            r = rs[0]
            out.append({"story": st, "elev": story_elevation(st, order[st]), "Pmin": r["Pmin"], "V2": r["V2"],
                        "M3": r["M3"], "Pc": r["Pc"], "combo_M3": "synthèse", "combo_V2": "synthèse"})
    return out


def stories_from_etabs(data, pier, P):
    """Construit la liste d'étages (modifiable) du projet à partir des efforts ETABS du pier."""
    gov = etabs_governing(data, pier)
    out = []
    for k, g in enumerate(gov):
        if k == 0:
            h = P["he"]                       # 1re hauteur = he (hauteur libre), comme le classeur
        elif k + 1 < len(gov):
            h = gov[k + 1]["elev"] - gov[k]["elev"]
        else:
            h = out[-1]["h"]                  # dernier niveau : même hauteur que le précédent
        out.append({"name": g["story"], "h": round(h, 3), "L": P["lw"], "e": P["bw"], "N": g["Pc"],
                    "V2": g["V2"], "M3": g["M3"], "Pmin": g["Pmin"]})
    return out


# ---------------------------------------------------------------------------
#  4c-ter. CONDITION D'ÉLANCEMENT DU VOILE (RPA 2024 — Art. 7.7.4)
# ---------------------------------------------------------------------------
ELANCEMENT_LIM = 2.0


def elancement(hw, lw):
    """Art. 7.7.4 — voile élancé si hw/lw > 2,0 : le diagramme de calcul des moments est l'enveloppe linéaire du
    diagramme issu du calcul de structure, décalée verticalement de hcr (Fig. 7.13). Voile court (hw/lw ≤ 2,0) :
    les moments du calcul de structure n'ont pas à être modifiés."""
    ratio = hw / lw if lw > 0 else float("inf")
    el = ratio > ELANCEMENT_LIM
    return {"hw": hw, "lw": lw, "ratio": ratio, "elance": el, "lim": ELANCEMENT_LIM,
            "cls": "Voile élancé" if el else "Voile court",
            "regle": ("enveloppe linéaire des moments décalée verticalement de hcr (Fig. 7.13)" if el else
                      "moments du calcul de structure non modifiés")}


# ---------------------------------------------------------------------------
#  4b-bis. DIMENSIONS DES VOILES DEPUIS L'EXPORT ETABS « PIER SECTION PROPERTIES »
#          (reprend « Calculation_Sheet_Voiles_ETABS.xlsx » : L = max(WidthBot, WidthTop), e = max(ThickBot,
#           ThickTop), h = CGTopZ − CGBotZ)
# ---------------------------------------------------------------------------
_SEC_ALIASES = {
    "story": ("story", "niveau", "etage", "étage"), "pier": ("pier", "voile", "trumeau"),
    "wb": ("widthbot", "width bot", "width bottom"), "tb": ("thickbot", "thick bot", "thickness bottom"),
    "wt": ("widthtop", "width top"), "tt": ("thicktop", "thick top", "thickness top"),
    "zb": ("cgbotz", "cg bot z"), "zt": ("cgtopz", "cg top z"), "ang": ("axisangle", "axis angle"),
}


def _nat_key(x):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(x))]


def _rows_from_section_table(table):
    for hi, hdr in enumerate(table[:25]):
        h = [_norm(c) for c in hdr]
        cols = {}
        for key, al in _SEC_ALIASES.items():
            for i, c in enumerate(h):
                if c in al:
                    cols.setdefault(key, i)
        if not all(k in cols for k in ("story", "pier", "wb", "tb", "wt", "tt", "zb", "zt")):
            continue
        out = []
        for r in table[hi + 1:]:
            if len(r) <= max(cols.values()):
                continue
            st, pr = r[cols["story"]], r[cols["pier"]]
            if st in (None, "") or pr in (None, ""):
                continue
            v = {k: _num(r[cols[k]]) for k in ("wb", "tb", "wt", "tt", "zb", "zt")}
            if None in v.values():
                continue                                    # ligne d'unités ou ligne vide
            L, e = max(v["wb"], v["wt"]), max(v["tb"], v["tt"])
            out.append({"story": str(st).strip(), "pier": str(pr).strip(), "L": L, "e": e,
                        "h": round(v["zt"] - v["zb"], 3), "S": L * e, "zb": v["zb"], "zt": v["zt"],
                        "ang": _num(r[cols["ang"]]) if "ang" in cols else None})
        if out:
            return out
    return None


def load_pier_sections(path):
    """Lit un export « Pier Section Properties » d'ETABS (xlsx / xlsm / csv)."""
    ext = os.path.splitext(path)[1].lower()
    found = []
    if ext in (".csv", ".txt", ".tsv"):
        with open(path, encoding="utf-8-sig", errors="replace") as fh:
            sample = fh.read(4096)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            except csv.Error:
                dialect = csv.excel
            found.append((os.path.basename(path), _rows_from_section_table([list(r) for r in csv.reader(fh, dialect)])))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            found.append((ws.title, _rows_from_section_table([list(r) for r in ws.iter_rows(values_only=True)])))
        wb.close()
    for name, rows in found:
        if rows:
            break
    else:
        raise ValueError("Aucune table « Pier Section Properties » reconnue (colonnes attendues : Story, Pier, "
                         "WidthBot, ThickBot, WidthTop, ThickTop, CGBotZ, CGTopZ).")
    zmin = {}
    for r in rows:
        zmin[r["story"]] = min(zmin.get(r["story"], r["zb"]), r["zb"])
    stories = sorted(zmin, key=lambda t: (zmin[t], _nat_key(t)))
    piers = sorted({r["pier"] for r in rows}, key=_nat_key)
    return {"rows": rows, "sheet": name, "stories": stories, "piers": piers}


def dimension_tables(data):
    """Tableaux de synthèse : matrices L et e (pier × niveau), synthèse par niveau, récapitulatif par pier."""
    rows, stories, piers = data["rows"], data["stories"], data["piers"]
    mat_L = {p: {} for p in piers}
    mat_e = {p: {} for p in piers}
    for r in rows:
        mat_L[r["pier"]][r["story"]] = mat_L[r["pier"]].get(r["story"], 0.0) + r["L"]
        mat_e[r["pier"]][r["story"]] = mat_e[r["pier"]].get(r["story"], 0.0) + r["e"]
    levels = []
    for st in stories:
        rs = [r for r in rows if r["story"] == st]
        zb, zt = min(r["zb"] for r in rs), max(r["zt"] for r in rs)
        levels.append({"name": st, "zb": zb, "zt": zt, "h": round(zt - zb, 3), "n": len(rs),
                       "sumL": sum(r["L"] for r in rs), "emin": min(r["e"] for r in rs),
                       "emax": max(r["e"] for r in rs), "sumS": sum(r["S"] for r in rs)})
    recap = []
    for p in piers:
        rs = sorted((r for r in rows if r["pier"] == p), key=lambda r: r["zb"])
        hw = sum(r["h"] for r in rs)
        lw0 = rs[0]["L"]                                  # lw de référence : longueur au niveau le plus bas du pier
        ratio = hw / lw0 if lw0 > 0 else float("inf")
        recap.append({"pier": p, "n": len(rs), "Lmin": min(r["L"] for r in rs), "Lmax": max(r["L"] for r in rs),
                      "emin": min(r["e"] for r in rs), "emax": max(r["e"] for r in rs), "hw": hw,
                      "S": sum(r["S"] for r in rs), "lw0": lw0, "ratio": ratio, "elance": ratio > ELANCEMENT_LIM})
    return {"mat_L": mat_L, "mat_e": mat_e, "levels": levels, "recap": recap,
            "hw_total": sum(l["h"] for l in levels), "S_total": sum(l["sumS"] for l in levels)}


def pier_levels(data, pier):
    """Dimensions d'un pier, niveau par niveau, du bas vers le haut."""
    rs = sorted((r for r in data["rows"] if r["pier"] == pier), key=lambda r: r["zb"])
    if not rs:
        raise ValueError(f"Pier « {pier} » absent du fichier de sections.")
    return rs


def apply_dims_to_stories(stories, data, pier):
    """Écrase h, L, e des niveaux (liste de dicts) par les dimensions du pier, par nom de niveau. Retourne le nombre
    de niveaux mis à jour."""
    by = {_norm(r["story"]): r for r in pier_levels(data, pier)}
    n = 0
    for s in stories:
        r = by.get(_norm(s["name"]))
        if r:
            s["h"], s["L"], s["e"] = round(r["h"], 3), r["L"], r["e"]
            n += 1
    return n


def stories_from_dims(data, pier):
    """Liste d'étages (sans efforts) construite depuis les dimensions du pier."""
    return [{"name": r["story"], "h": r["h"], "L": r["L"], "e": r["e"], "N": 0.0, "V2": 0.0, "M3": 0.0,
             "Pmin": 0.0, "sec": ""} for r in pier_levels(data, pier)]


def export_dimensions_xlsx(data, path):
    """Exporte les tableaux de dimensions (valeurs) dans un classeur Excel : données, L, e, synthèse, récapitulatif."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    T = dimension_tables(data)
    wb = Workbook()
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")

    def sheet(title, header, rows, widths=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = hdr_font, hdr_fill, Alignment(horizontal="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "B2"
        return ws

    sheet("Données ETABS", ["Story", "Pier", "Longueur L (m)", "Épaisseur e (m)", "Cote base (m)", "Cote sommet (m)",
                            "Hauteur étage h (m)", "Surface L×e (m²)"],
          [[r["story"], r["pier"], r["L"], r["e"], r["zb"], r["zt"], r["h"], round(r["S"], 4)] for r in data["rows"]],
          [12, 8, 14, 14, 12, 14, 16, 16], first=True)
    for nm, mat, ttl in (("Longueur L", T["mat_L"], "Pier"), ("Épaisseur e", T["mat_e"], "Pier")):
        sheet(nm, [ttl] + data["stories"],
              [[p] + [mat[p].get(st) for st in data["stories"]] for p in data["piers"]], [8] + [11] * len(data["stories"]))
    sheet("Synthèse niveau", ["Niveau", "Cote base (m)", "Cote sommet (m)", "Hauteur étage h (m)", "Nb piers",
                              "Σ Longueur L (m)", "Épaisseur min (m)", "Épaisseur max (m)", "Σ Surface L×e (m²)"],
          [[l["name"], l["zb"], l["zt"], l["h"], l["n"], round(l["sumL"], 3), l["emin"], l["emax"], round(l["sumS"], 4)]
           for l in T["levels"]] + [["Total / Hauteur totale", None, None, round(T["hw_total"], 3), None, None, None, None,
                                     round(T["S_total"], 4)]], [16, 12, 14, 16, 9, 16, 14, 14, 18])
    sheet("Récap par pier", ["Pier", "Nb niveaux", "L min (m)", "L max (m)", "e min (m)", "e max (m)", "hw = Σ h (m)",
                             "lw base (m)", "hw/lw", "Classe (Art. 7.7.4)", "Σ Surface L×e (m²)"],
          [[q["pier"], q["n"], q["Lmin"], q["Lmax"], q["emin"], q["emax"], round(q["hw"], 3), q["lw0"],
            round(q["ratio"], 2), "élancé (> 2,0)" if q["elance"] else "court (≤ 2,0)", round(q["S"], 4)]
           for q in T["recap"]], [8, 11, 10, 10, 10, 10, 13, 12, 8, 18, 16])
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
#  Coordonnées du bureau d'études et site d'aide
# ---------------------------------------------------------------------------
CONTACT = {"bureau": "ZED ENGINEERING", "ingenieur": "ZEBIRI Fouad", "mail": "zed.engineering3@gmail.com",
           "tels": ["+213 555 49 12 18", "+213 797 537 625"]}
# Adresse du site d'aide : à renseigner ici lorsque le site sera en ligne (ou via la variable d'environnement
# VOILE_HELP_URL). Tant qu'elle est vide, le bouton « Aide » l'indique à l'utilisateur.
HELP_URL = os.environ.get("VOILE_HELP_URL", "")


# ---------------------------------------------------------------------------
#  4e. HYPOTHÈSES SISMIQUES ET MATÉRIAUX (RPA 2024 — tables du classeur « Note_de_calcul_Hypotheses »)
# ---------------------------------------------------------------------------
RPA_ZONES = {"Zone 0": ("Très faible", 0.0), "Zone I": ("Faible", 0.07), "Zone II": ("Faible à moyenne", 0.10),
             "Zone III": ("Moyenne", 0.15), "Zone IV": ("Moyenne à élevée", 0.20), "Zone V": ("Élevée", 0.25),
             "Zone VI": ("Élevée", 0.30)}                                              # Table 3.3
RPA_GROUPES = {"Groupe 1A": 1.4, "Groupe 1B": 1.2, "Groupe 2": 1.0, "Groupe 3": 0.8}   # Table 3.10
RPA_SITES = {"S1": (1.0, 0.10, 0.4, 2.0), "S2": (1.2, 0.10, 0.5, 2.0), "S3": (1.3, 0.15, 0.6, 2.0),
             "S4": (1.35, 0.15, 0.7, 2.0)}                                             # Table 3.3 : S, T1, T2, T3
RPA_SYSTEMES = {1: ("Système à ossature", 5.5, "a"),
                2: ("Système à contreventement mixte, équivalent à une ossature", 5.5, "a"),
                3: ("Système à ossature ou mixte équivalent à ossature avec remplissage en maçonnerie rigide", 3.5, "a"),
                4: ("Système à contreventement mixte, équivalent à des voiles", 4.5, "a"),
                5: ("Système de contreventement constitué par des voiles", 4.5, "b"),
                6: ("Système à ossature à noyau ou à effet noyau", 3.0, "b"),
                7: ("Système fonctionnant en console verticale à masses réparties prédominantes", 3.0, "c"),
                8: ("Système en pendule inversé", 2.0, "c"),
                9: ("Système de voiles de grandes dimensions en béton peu armé", 1.5, "c")}   # Table 3.18
RPA_PQ = [("a", "1. Régularité en plan", 0.05), ("a", "2. Régularité en élévation", 0.20),
          ("a", "3. Conditions minimales sur le nombre d'étages", 0.20),
          ("a", "4. Conditions minimales sur les travées", 0.10),
          ("b", "1. Régularité en plan", 0.05), ("b", "2. Régularité en élévation", 0.20),
          ("b", "3. Redondance en plan", 0.05)]                                          # Table 3.19
RPA_AMORT = {"Portique – béton armé": (6, 7), "Portique – acier": (4, 5),
             "Voile (béton armé / maçonnerie)": (10, 10)}                               # Table 3.5 : léger, dense
MAT_SITUATIONS = {"Durable / transitoire": (1.5, 1.15), "Accidentelle (séisme)": (1.2, 1.0), "Manuel": None}
SYS_LABELS = [f"{n} — {v[0]}" for n, v in RPA_SYSTEMES.items()]


def _sys_num(txt):
    m = re.match(r"\s*(\d+)", str(txt))
    if not m or int(m.group(1)) not in RPA_SYSTEMES:
        raise ValueError(f"Système de contreventement invalide : « {txt} »")
    return int(m.group(1))


def seismic_params(P):
    """Paramètres sismiques RPA 2024 choisis par l'ingénieur : A, I, S, T1-T3, R, QF, ξ, η, T = CT·hN^(3/4)."""
    if P["sz_zone"] not in RPA_ZONES or P["sz_groupe"] not in RPA_GROUPES or P["sz_site"] not in RPA_SITES:
        raise ValueError("Zone sismique, groupe d'importance ou site invalide.")
    if P["sz_struct"] not in RPA_AMORT or P["sz_remp"] not in ("Léger", "Dense"):
        raise ValueError("Type de structure ou remplissage invalide.")
    niv, A = RPA_ZONES[P["sz_zone"]]
    I = RPA_GROUPES[P["sz_groupe"]]
    S, T1, T2, T3 = RPA_SITES[P["sz_site"]]
    n = _sys_num(P["sz_sys"])
    desc, R, cat = RPA_SYSTEMES[n]
    pq = []
    for i, (c, lab, v) in enumerate(RPA_PQ, start=1):
        obs = str(P[f"sz_pq{i}"]).strip().lower() != "non"
        pq.append({"cat": c, "label": lab, "pq": v, "obs": obs, "actif": c == cat, "pen": v if (c == cat and not obs) else 0.0})
    QF = 1.0 + sum(q["pen"] for q in pq)
    if P["sz_hn"] <= 0 or P["sz_ct"] <= 0:
        raise ValueError("CT et hN doivent être positifs.")
    T = P["sz_ct"] * P["sz_hn"] ** 0.75
    xi = RPA_AMORT[P["sz_struct"]][1 if P["sz_remp"] == "Dense" else 0]
    eta = math.sqrt(7.0 / (2.0 + xi))
    return {"niv": niv, "A": A, "I": I, "S": S, "T1": T1, "T2": T2, "T3": T3, "n_sys": n, "sys": desc, "R": R,
            "cat": cat, "pq": pq, "QF": QF, "T": T, "xi": xi, "eta": eta, "Sa_min": 0.2 * A * I}


def material_params(P):
    """Béton (fc28, Ev, Ei), densité, aciers et coefficients γb, γs selon la situation de calcul."""
    for k in ("mt_fc_fond", "mt_fc", "mt_fe_ha"):
        if P[k] <= 0:
            raise ValueError("fc28 et fe doivent être positifs.")
    ev = lambda fc: 3700.0 * fc ** (1 / 3)
    ei = lambda fc: 11000.0 * fc ** (1 / 3)
    sit = P["mt_sit"]
    if sit not in MAT_SITUATIONS:
        raise ValueError(f"Situation de calcul invalide : « {sit} »")
    gb, gs = MAT_SITUATIONS[sit] or (P["mt_gb_m"], P["mt_gs_m"])
    return {"fc_fond": P["mt_fc_fond"], "fc": P["mt_fc"], "Ev": ev(P["mt_fc"]), "Ei": ei(P["mt_fc"]),
            "Ev_fond": ev(P["mt_fc_fond"]), "Ei_fond": ei(P["mt_fc_fond"]), "dens": P["mt_dens"],
            "fe": P["mt_fe_ha"], "fe_tr": P["mt_fe_tr"], "fe_doux": P["mt_fe_doux"], "gb": gb, "gs": gs,
            "fbu": 0.85 * P["mt_fc"] / gb, "sigs": P["mt_fe_ha"] / gs, "sit": sit}


def link_values(P):
    """Valeurs des données générales imposées par les rubriques « Données sismiques » et « Matériaux » lorsque la
    liaison est active : {clé: valeur}. Fonction pure (aucune modification de P)."""
    out = {}
    if int(P.get("seis_link", 0)):
        sp = seismic_params(P)
        out.update({"R": sp["R"], "Qf": round(sp["QF"], 4), "T2": sp["T2"], "T0": round(sp["T"], 4)})
    if int(P.get("mat_link", 0)):
        mp = material_params(P)
        out.update({"fc": mp["fc"], "fe": mp["fe"], "gb": mp["gb"], "gs": mp["gs"], "mr_fc": mp["fc"],
                    "mr_fe": mp["fe"], "c1_fc": mp["fc"], "c1_fe": mp["fe"], "c2_fc": mp["fc"], "c2_fe": mp["fe"]})
    return out


def apply_links(P):
    """Applique la liaison à P (dictionnaire de calcul) et retourne les valeurs imposées."""
    v = link_values(P)
    P.update(v)
    return v


# ---------------------------------------------------------------------------
#  4c. ENVELOPPE DES MOMENTS FLÉCHISSANTS (RPA 2024 — éq. 7.22 et 7.23)
# ---------------------------------------------------------------------------
def envelope(he, lw, heights, mbase, msom, ref="base"):
    """heights : hauteurs d'étage (m) de bas en haut. M enveloppe linéaire décalée de hcr."""
    n = len(heights)
    hw = sum(heights)
    if n == 0 or hw <= 0:
        raise ValueError("Enveloppe : au moins un étage de hauteur > 0 est requis.")
    hcr1 = max(lw, hw / 6)                                  # éq. 7.22
    limite = min(2 * lw, he if n <= 6 else 2 * he)          # éq. 7.23
    hcr = min(hcr1, limite)
    pente = (mbase - msom) / hw
    m_top = mbase - pente * (hw - hcr)

    def m_at(H):
        H = min(max(H, 0.0), hw)
        return mbase if H <= hcr + 1e-9 else mbase - pente * (H - hcr)

    cum = [0.0]
    for h in heights:
        cum.append(round(cum[-1] + h, 4))
    table = [{"k": k, "H": cum[k], "M_lin": mbase - pente * cum[k], "M_corr": m_at(cum[k]),
              "zone": "Zone critique (M = M base)" if cum[k] <= hcr + 1e-9 else "Enveloppe décalée"}
             for k in range(n + 1)]
    # moment appliqué à l'étage k : au sommet (comme le classeur) ou à la base (section du bas, sécuritaire)
    m_story = [m_at(cum[k + 1] if ref == "sommet" else cum[k]) for k in range(n)]
    return {"n": n, "he": he, "lw": lw, "hw": hw, "mbase": mbase, "msom": msom, "hcr1": hcr1, "limite": limite,
            "hcr": hcr, "pente": pente, "m_top": m_top, "cum": cum, "table": table, "m_story": m_story,
            "m_at": m_at, "ref": ref}


# ---------------------------------------------------------------------------
#  4c-bis. ENVELOPPE DE L'EFFORT TRANCHANT (RPA 2024 — Art. 7.7.2)
# ---------------------------------------------------------------------------
K_AMPL_V = 1.4          # facteur d'amplification de l'effort tranchant en zone critique (capacity design)


def envelope_v(env, stories):
    """Enveloppe de calcul de l'effort tranchant, niveau par niveau (section du bas de chaque étage).

    V corrigé(H) = 1,4·V(H)                       si H ≤ hcr
                 = interpolation linéaire entre (hcr ; 1,4·V(hcr)) et (hw/3 ; V(hw/3))   si hcr < H < hw/3
                 = V(H)                           si H ≥ hw/3
    [CORR.] Le classeur calcule V(hcr) et V(hw/3) avec TREND(), c.-à-d. par RÉGRESSION linéaire sur tous les
    niveaux, alors que l'en-tête de la feuille annonce une « interpolation linéaire sur le tableau des niveaux » :
    on applique ici l'interpolation linéaire par morceaux entre niveaux (valeurs bornées aux extrémités)."""
    n, cum, hcr, hw = env["n"], env["cum"], env["hcr"], env["hw"]
    h3 = hw / 3.0
    pts = [(cum[k], abs(stories[k]["V2"])) for k in range(n)]

    def v_an(H):
        if H <= pts[0][0]:
            return pts[0][1]
        for (h0, v0), (h1, v1) in zip(pts, pts[1:]):
            if H <= h1 + 1e-12:
                return v1 if h1 <= h0 else v0 + (v1 - v0) * (H - h0) / (h1 - h0)
        return pts[-1][1]

    v_hcr, v_h3 = v_an(hcr), v_an(h3)

    def v_at(H, ti):
        if H <= hcr + 1e-9:
            return K_AMPL_V * ti, "Zone critique (1,4·V)"
        if H < h3 - 1e-9:
            return (K_AMPL_V * v_hcr + (H - hcr) * (v_h3 - K_AMPL_V * v_hcr) / (h3 - hcr), "Transition linéaire")
        return ti, "Analyse (non majoré)"

    table = []
    for k in range(n):
        H, ti = cum[k], pts[k][1]
        vc, zone = v_at(H, ti)
        table.append({"k": k, "H": H, "Ti": ti, "T14": K_AMPL_V * ti, "Vc": vc, "zone": zone})
    return {"n": n, "hcr": hcr, "hw": hw, "h3": h3, "v_hcr": v_hcr, "v_h3": v_h3, "table": table,
            "v_story": [t["Vc"] for t in table], "k": K_AMPL_V}


# ---------------------------------------------------------------------------
#  4d. MÉTHODE DES CONTRAINTES — TOUS LES ÉTAGES
# ---------------------------------------------------------------------------
CAS_DEFAUT = [("Durable / transitoire", 30.0, 500.0, 1.5, 1.15, 1.0),
              ("Accidentelle", 30.0, 500.0, 1.2, 1.0, 1.0),
              ("Personnalisé", 30.0, 500.0, 1.2, 1.0, 1.0)]


def _cas_list(P):
    out = []
    for i in (1, 2, 3):
        out.append({"name": str(P[f"c{i}_name"]), "fc": P[f"c{i}_fc"], "fe": P[f"c{i}_fe"], "gb": P[f"c{i}_gb"],
                    "gs": P[f"c{i}_gs"], "th": P[f"c{i}_th"]})
    for c in out:
        c["fbu"] = 0.85 * c["fc"] / (c["th"] * c["gb"])
        c["sigs"] = c["fe"] / c["gs"]
    return out


# --- section d'acier ADOPTÉE (facultative) : « 12HA20 », « 12 T 20 », « 12Ø20 », « 6HA20 + 4HA16 »…
_SEC_RE = re.compile(r"^\s*(\d+)\s*(?:HA|H\.A\.?|T|Ø|ø|φ|Φ|ϕ|PHI|x|×|\*)\s*(\d+(?:[.,]\d+)?)\s*$", re.I)


def parse_section(txt):
    """Texte → liste de couples (nombre de barres, diamètre en mm). Vide → []. Lève ValueError si illisible."""
    txt = str(txt if txt is not None else "").strip()
    if not txt:
        return []
    terms = []
    for part in re.split(r"\s*[+;]\s*", txt):
        if not part.strip():
            continue
        m = _SEC_RE.match(part)
        if not m:
            raise ValueError(f"section adoptée illisible : « {txt} » (écrire par ex. 12HA20 ou 6HA20+4HA16)")
        n, phi = int(m.group(1)), float(m.group(2).replace(",", "."))
        if n <= 0 or not (5.0 <= phi <= 50.0):
            raise ValueError(f"section adoptée « {part.strip()} » : nombre > 0 et diamètre entre 5 et 50 mm requis")
        terms.append((n, phi))
    return terms


def section_area(terms):
    """Section totale (cm²) d'une liste (n, Ø mm)."""
    return sum(n * math.pi * phi ** 2 / 4 for n, phi in terms) / 100.0


def section_text(terms):
    return " + ".join(f"{n}HA{phi:g}" for n, phi in terms)


def propose_section(As, phi):
    """Nombre minimal de barres de diamètre phi (mm) pour couvrir As (cm²), au moins 2."""
    a1 = math.pi * phi ** 2 / 400.0
    return f"{max(2, ceil_(As / a1))}HA{phi:g}"


def stress_level(N, V2, M, L, e, h, cas, P, kV=None):
    """Calcul d'un niveau (N<0 compression ; kN, kN.m, m, MPa). kV : surcharge du coefficient kV (sinon P['kV'])."""
    A = L * e
    I = e * L ** 3 / 12
    v = L / 2
    s1 = (N / A + M * v / I) / 1000
    s2 = (N / A - M * v / I) / 1000
    smax, smin = max(s1, s2), min(s1, s2)
    if s1 < 0 and s2 < 0:
        typ = "S.E.C"
    elif s1 > 0 and s2 > 0:
        typ = "S.E.T"
    else:
        typ = "S.P.C"
    if typ == "S.P.C":
        Lt = L * smax / (smax - smin) if smax != smin else 0.0
        Ft = smax * 1000 * e * Lt / 2
    elif typ == "S.E.T":
        Lt = L
        Ft = (s1 + s2) * 1000 * e * Lt / 2
    else:
        Lt, Ft = 0.0, 0.0
    As_calc = Ft * 10 / cas["sigs"]                                    # cm²
    As_min = (P["rho_g"] * e * L if typ == "S.E.C" else P["rho_t"] * e * Lt) * 10000
    As = max(As_calc, As_min)
    sc = max(0.0, -min(s1, s2))
    d = P["kd"] * (L if str(P["dref"]).strip().lower().startswith("l") else h)
    kVe = P["kV"] if kV is None else kV
    tau = abs(kVe * V2 / (d * e * 1000))
    taulim = P["ktau"] * cas["fc"]
    r = {"A": A, "I": I, "s1": s1, "s2": s2, "type": typ, "Lt": Lt, "Ft": Ft, "As_calc": As_calc,
         "As_min": As_min, "As": As, "As_Lt": (As / Lt if Lt else 0.0), "sc": sc, "sc_fbu": sc / cas["fbu"],
         "ok_c": sc / cas["fbu"] <= 1, "d": d, "tau": tau, "taulim": taulim, "tau_lim": tau / taulim,
         "ok_t": tau < taulim}
    r["ok"] = r["ok_c"] and r["ok_t"]
    return r


def compute_multi(P):
    """Enveloppe des moments + méthode des contraintes sur tous les étages (cas actif et comparaison)."""
    S = P.get("stories") or []
    if not S:
        raise ValueError("Aucun étage défini (importer un fichier ETABS ou ajouter des niveaux).")
    for s in S:
        if s["L"] <= 0 or s["e"] <= 0:
            raise ValueError(f"Niveau {s['name']} : L et e doivent être strictement positifs.")
    heights = [s["h"] for s in S]
    if P["env_auto"]:
        mbase, msom = abs(S[0]["M3"]), abs(S[-1]["M3"])
    else:
        mbase, msom = abs(P["env_mbase"]), abs(P["env_msom"])
    env = envelope(P["he"], P["lw"], heights, mbase, msom, str(P["env_ref"]))
    envv = envelope_v(env, S)
    v_mode = str(P.get("v_mode", "kV"))
    secs = []
    for s in S:
        try:
            secs.append(parse_section(s.get("sec", "")))
        except ValueError as e:
            raise ValueError(f"Niveau {s['name']} : {e}")
    cas_all = _cas_list(P)
    names = [c["name"] for c in cas_all]
    act = names.index(P["cas"]) if P["cas"] in names else 1
    cas = cas_all[act]
    mode = str(P["m_mode"])
    elanc = elancement(env["hw"], P["lw"])
    # « auto » : voile élancé → enveloppe décalée de hcr ; voile court → moments du calcul de structure (Art. 7.7.4)
    mode_eff = ("enveloppe" if elanc["elance"] else "etabs") if mode == "auto" else mode

    def run(c):
        rows = []
        for k, s in enumerate(S):
            mc = env["m_story"][k]
            M = {"enveloppe": mc, "max": max(mc, abs(s["M3"])), "etabs": abs(s["M3"])}.get(mode_eff, mc)
            if v_mode == "enveloppe":                    # V corrigé de l'enveloppe (1,4 déjà inclus en zone critique)
                r = stress_level(s["N"], envv["v_story"][k], M, s["L"], s["e"], s["h"], c, P, kV=1.0)
                r["V_calc"] = envv["v_story"][k]
            else:                                        # kV × V2 ETABS (comportement d'origine)
                r = stress_level(s["N"], s["V2"], M, s["L"], s["e"], s["h"], c, P)
                r["V_calc"] = P["kV"] * abs(s["V2"])
            r.update({"name": s["name"], "N": s["N"], "V2": s["V2"], "M": M, "M_etabs": s["M3"], "L": s["L"],
                      "e": s["e"], "h": s["h"], "M_env": mc})
            # section adoptée (facultative) : section choisie et comparaison à la section requise
            terms = secs[k]
            r["sec"] = section_text(terms)
            r["As_ch"] = section_area(terms) if terms else None
            r["ok_as"] = None if not terms else (r["As_ch"] >= r["As"] - 1e-9)
            r["ratio_as"] = None if not terms or r["As"] <= 0 else r["As_ch"] / r["As"]
            rows.append(r)
        return rows

    rows = run(cas)
    comp = []
    for c in cas_all:
        rr = run(c)
        comp.append({"cas": c, "sumAs": sum(r["As"] for r in rr), "n_ng": sum(not r["ok"] for r in rr),
                     "max_sc": max(r["sc_fbu"] for r in rr), "max_tau": max(r["tau_lim"] for r in rr)})
    n_as_ko = sum(r["ok_as"] is False for r in rows)
    return {"env": env, "envv": envv, "v_mode": v_mode, "rows": rows, "cas": cas, "cas_all": cas_all, "comp": comp,
            "mode": mode, "mode_eff": mode_eff, "elanc": elanc, "n_ng": sum(not r["ok"] for r in rows), "n_as_ko": n_as_ko,
            "n_sec": sum(r["ok_as"] is not None for r in rows), "has_sec": any(r["ok_as"] is not None for r in rows),
            "all_ok": all(r["ok"] for r in rows) and n_as_ko == 0,
            "As_max": max(r["As"] for r in rows), "sumAs": sum(r["As"] for r in rows),
            "sumAs_ch": sum(r["As_ch"] for r in rows if r["As_ch"] is not None),
            "max_sc": max(r["sc_fbu"] for r in rows), "max_tau": max(r["tau_lim"] for r in rows)}


# ---------------------------------------------------------------------------
#  4d-bis. MOMENT RÉSISTANT — FLEXION COMPOSÉE (BAEL 91 modifié 99 / RPA 2024)
#          Reprend le classeur « Fiche_Moment_Resistant_BAEL91_RPA2024.xlsx » (onglets Calcul_ELU / Calcul_ELS)
# ---------------------------------------------------------------------------
MR_SITUATIONS = ["Durable / transitoire", "Accidentelle (RPA 2024)", "Saisie manuelle"]
MR_FISSURATION = ["Peu préjudiciable", "Préjudiciable", "Très préjudiciable"]


def _bisect(f, lo, hi, n=200):
    """Dichotomie sur une fonction croissante f : renvoie x tel que f(x) ≈ 0 (comme les onglets Iter_ELU / Iter_ELS)."""
    for _ in range(n):
        mid = 0.5 * (lo + hi)
        if f(mid) < 0:
            lo = mid
        else:
            hi = mid
    return 0.5 * (lo + hi)


def _mr_elu_state(y, g):
    """État de la section à l'ELU pour un axe neutre à la profondeur y (m, depuis la fibre supérieure).
    Béton parabole-rectangle, acier élasto-plastique parfait, règle des trois pivots, béton tendu négligé.
    Compression > 0 ; N en kN ; M en kN.m par rapport au centre de gravité du béton (h/2)."""
    b, h, dp, c, Ap, A = g["b"], g["h"], g["dp"], g["c"], g["Ap"], g["A"]
    fbu, fsu, Es = g["fbu"], g["fsu"], g["Es"]
    d = h - c
    if y <= 3.5 / 13.5 * d:
        K = 10 / (d - y)                    # pivot A
    elif y <= h:
        K = 3.5 / y                         # pivot B
    else:
        K = 2 / (y - 3 * h / 7)             # pivot C
    zmax = min(y, h)
    z2 = y - 2 / K                          # profondeur où ε = 2 ‰ (début du palier)
    zR = max(0.0, min(z2, zmax))            # hauteur du palier rectangulaire
    za = max(0.0, z2)                       # début de la zone parabolique
    ea = min(2.0, K * (y - za))
    eb = K * (y - zmax)
    if za < zmax:
        Fp = ((ea ** 2 / 2 - ea ** 3 / 12) - (eb ** 2 / 2 - eb ** 3 / 12)) / K
        Mp = ((y * (ea ** 2 / 2 - ea ** 3 / 12) - (ea ** 3 / 3 - ea ** 4 / 16) / K)
              - (y * (eb ** 2 / 2 - eb ** 3 / 12) - (eb ** 3 / 3 - eb ** 4 / 16) / K)) / K
    else:
        Fp = Mp = 0.0
    Nc = b * fbu * (zR + Fp) * 1000
    Mt = b * fbu * (zR ** 2 / 2 + Mp) * 1000
    es1, es2 = K * (y - dp), K * (y - d)    # ‰ (compression > 0)
    ss1 = max(-fsu, min(fsu, Es * es1 / 1000))
    ss2 = max(-fsu, min(fsu, Es * es2 / 1000))
    Fsp, Fst = Ap * ss1 / 10, A * ss2 / 10
    N = Nc + Fsp + Fst
    M = Nc * h / 2 - Mt + Fsp * (h / 2 - dp) + Fst * (h / 2 - d)
    return {"y": y, "K": K, "N": N, "M": M, "Nc": Nc, "Fsp": Fsp, "Fst": Fst, "es1": es1, "es2": es2,
            "ss1": ss1, "ss2": ss2, "d": d, "yAB": 3.5 / 13.5 * d,
            "pivot": "Pivot A" if y <= 3.5 / 13.5 * d else ("Pivot B" if y <= h else "Pivot C")}


def _mr_els_state(x, g):
    """État de la section fissurée homogénéisée (n = 15) à l'ELS, axe neutre à la profondeur x."""
    b, h, dp, c, Ap, A, n = g["b"], g["h"], g["dp"], g["c"], g["Ap"], g["A"], g["n"]
    d = h - c
    sbc = min(g["sbc_lim"], g["ss_lim"] * x / (n * (d - x))) if x < d else g["sbc_lim"]
    Nc = 0.5 * b * x * sbc * 1000
    sp = n * sbc * (x - dp) / x
    st = n * sbc * (d - x) / x
    Fsp, Fst = Ap * sp / 10, -A * st / 10
    N = Nc + Fsp + Fst
    M = Nc * (h / 2 - x / 3) + Fsp * (h / 2 - dp) + Fst * (h / 2 - d)
    return {"x": x, "N": N, "M": M, "sbc": sbc, "sp": sp, "st": st, "Nc": Nc, "Fsp": Fsp, "Fst": Fst, "d": d}


def compute_mr(P):
    """Moment résistant ELU / ELS d'une section rectangulaire sous effort normal donné.
    Retourne un dictionnaire de résultats, des sections de lignes (affichage / note) et les vérifications."""
    b, h, dp, c = P["mr_b"], P["mr_h"], P["mr_dp"], P["mr_c"]
    Ap, A = P["mr_Ap"], P["mr_A"]
    fc, fe, Es, n = P["mr_fc"], P["mr_fe"], P["mr_Es"], P["mr_n"]
    if min(b, h, fc, fe, Es, n) <= 0:
        raise ValueError("Moment résistant : b, h, fc28, fe, Es et n doivent être strictement positifs.")
    if not (0 < dp < h and 0 < c < h):
        raise ValueError("Moment résistant : d' et c doivent être compris entre 0 et h.")
    if Ap < 0 or A < 0:
        raise ValueError("Moment résistant : A et A' doivent être positifs ou nuls.")
    sit = str(P["mr_sit"])
    if sit == "Saisie manuelle":
        gb, gs = P["mr_gb_m"], P["mr_gs_m"]
    elif sit.startswith("Accidentelle"):
        gb, gs = 1.15, 1.0
    else:
        gb, gs = 1.5, 1.15
    th = P["mr_theta"]
    if min(gb, gs, th) <= 0:
        raise ValueError("Moment résistant : γb, γs et θ doivent être strictement positifs.")
    do_elu = str(P["mr_elu"]).strip().lower().startswith("o")
    do_els = str(P["mr_els"]).strip().lower().startswith("o")
    Nu, Ns = P["mr_Nu"], P["mr_Ns"]
    Mu_in, Ms_in = _num(P.get("mr_Mu", "")), _num(P.get("mr_Ms", ""))
    fbu = 0.85 * fc / (th * gb)
    fsu = fe / gs
    g = {"b": b, "h": h, "dp": dp, "c": c, "Ap": Ap, "A": A, "fbu": fbu, "fsu": fsu, "Es": Es, "n": n}
    out = {"gb": gb, "gs": gs, "fbu": fbu, "fsu": fsu, "d": h - c, "do_elu": do_elu, "do_els": do_els,
           "elu": None, "els": None, "checks": [], "Mu_in": Mu_in, "Ms_in": Ms_in}

    # ---------------- ELU
    nu = Nu / (b * h * fc * 1000)
    out["nu"], out["nu_lim"] = nu, P["mr_nulim"]
    out["ok_nu"] = nu <= P["mr_nulim"]
    if do_elu:
        y = _bisect(lambda t: _mr_elu_state(t, g)["N"] - Nu, 0.0001, 100.0)
        st = _mr_elu_state(y, g)
        st["conv"] = abs(st["N"] - Nu) < 0.05
        st["Mru"] = st["M"]
        out["elu"] = st
        if Mu_in is not None:
            out["ratio_u"] = abs(Mu_in) / st["Mru"] if st["Mru"] else float("inf")
            out["ok_u"] = st["conv"] and abs(Mu_in) <= st["Mru"]
            out["checks"].append(("Moment résistant ELU : |Mu| ≤ Mru", out["ok_u"]))

    # ---------------- ELS
    if do_els:
        sbc_lim = 0.6 * fc
        ftj = 0.6 + 0.06 * fc
        eta = P["mr_eta"]
        fis = str(P["mr_fiss"])
        if fis == "Peu préjudiciable":
            ss_lim = fe
        elif fis == "Préjudiciable":
            ss_lim = min(2 / 3 * fe, max(0.5 * fe, 110 * math.sqrt(eta * ftj)))
        else:
            ss_lim = min(0.5 * fe, 90 * math.sqrt(eta * ftj))
        g.update({"sbc_lim": sbc_lim, "ss_lim": ss_lim})
        x = _bisect(lambda t: _mr_els_state(t, g)["N"] - Ns, 0.0001, h)
        st = _mr_els_state(x, g)
        st["conv"] = abs(st["N"] - Ns) < 0.05
        st["Mrs"] = st["M"]
        st["crit"] = "Contrainte de l'acier tendu" if st["sbc"] < sbc_lim - 1e-6 else "Contrainte du béton comprimé"
        st.update({"sbc_lim": sbc_lim, "ss_lim": ss_lim, "ftj": ftj})
        out["els"] = st
        if Ms_in is not None:
            out["ratio_s"] = abs(Ms_in) / st["Mrs"] if st["Mrs"] else float("inf")
            out["ok_s"] = st["conv"] and abs(Ms_in) <= st["Mrs"]
            out["checks"].append(("Moment résistant ELS : |Ms| ≤ Mrs", out["ok_s"]))
    out["all_ok"] = all(ok for _, ok in out["checks"])
    out["g"] = g
    out["sections"] = _mr_sections(P, out)
    return out


def _mr_sections(P, mr):
    """Sections de lignes (titre, [(libellé, expression, valeur, unité, ok)]) — partagées par l'écran et la note."""
    S = []
    S.append(("DONNÉES DE LA SECTION", [
        ("b — Largeur", "", fr(P["mr_b"]), "m", None),
        ("h — Hauteur", "", fr(P["mr_h"]), "m", None),
        ("d' — Position cdg A' (depuis la fibre sup.)", "", fr(P["mr_dp"]), "m", None),
        ("c — Position cdg A (depuis la fibre inf.)", "", fr(P["mr_c"]), "m", None),
        ("d — Hauteur utile", "h − c", fr(mr["d"]), "m", None),
        ("A' — Armatures supérieures", "", fr(P["mr_Ap"]), "cm²", None),
        ("A — Armatures inférieures", "", fr(P["mr_A"]), "cm²", None),
        ("fc28 — Résistance du béton", "", fr(P["mr_fc"], 1), "MPa", None),
        ("fe — Limite élastique de l'acier", "", fr(P["mr_fe"], 0), "MPa", None),
        ("Situation de calcul", "", str(P["mr_sit"]), "", None),
        ("γb / γs", "", f"{fr(mr['gb'])} / {fr(mr['gs'])}", "-", None),
        ("fbu — Contrainte de calcul du béton", "0,85·fc28 / (θ·γb)", fr(mr["fbu"]), "MPa", None),
        ("fsu — Contrainte de calcul de l'acier", "fe / γs", fr(mr["fsu"], 1), "MPa", None),
    ]))
    e = mr["elu"]
    if e:
        rows = [
            ("Nu — Effort normal (compression > 0)", "donnée", fr(P["mr_Nu"], 1), "kN", None),
            ("y — Position de l'axe neutre", "équilibre N(y) = Nu (dichotomie)", fr(e["y"], 4), "m", None),
            ("Pivot", "A : y ≤ 3,5/13,5·d ; B : y ≤ h ; C : y > h", e["pivot"], "", None),
            ("K — Courbure", "", fr(e["K"], 4), "‰/m", None),
            ("ε' — Déformation des armatures supérieures", "compression > 0", fr(e["es1"], 3), "‰", None),
            ("ε — Déformation des armatures inférieures", "compression > 0", fr(e["es2"], 3), "‰", None),
            ("σ's / σs — Contraintes des armatures", "plafonnées à ± fsu", f"{fr(e['ss1'], 1)} / {fr(e['ss2'], 1)}", "MPa",
             None),
            ("Nbc — Résultante du béton comprimé", "", fr(e["Nc"], 1), "kN", None),
            ("A'·σ's — Force des armatures supérieures", "", fr(e["Fsp"], 1), "kN", None),
            ("A·σs — Force des armatures inférieures", "", fr(e["Fst"], 1), "kN", None),
            ("Mru — Moment résistant ELU", "M(y) / cdg de la section", fr(e["Mru"], 1), "kN.m", None),
            ("Convergence", "|N(y) − Nu| < 0,05 kN", "OK" if e["conv"] else "Nu hors domaine de résistance", "",
             e["conv"]),
            ("ν — Effort normal réduit", "Nu / (b·h·fc28)", fr(mr["nu"], 4), "-", None),
            ("Contrôle ν ≤ νlim (informatif)", f"νlim = {fr(mr['nu_lim'])}",
             "ν ≤ νlim" if mr["ok_nu"] else "ν > νlim — vérifier la ductilité", "", mr["ok_nu"]),
        ]
        if mr["Mu_in"] is not None:
            rows += [("Mu / Mru", "|Mu| / Mru", fr(mr["ratio_u"], 3), "-", mr["ok_u"]),
                     ("Vérification à l'ELU", f"Mu = {fr(abs(mr['Mu_in']), 1)} kN.m",
                      "Section vérifiée" if mr["ok_u"] else "Section NON vérifiée", "", mr["ok_u"])]
        S.append(("RÉSULTATS À L'ELU — BAEL 91 (trois pivots)", rows))
    e = mr["els"]
    if e:
        rows = [
            ("Ns — Effort normal (compression > 0)", "donnée", fr(P["mr_Ns"], 1), "kN", None),
            ("Fissuration / η", "", f"{P['mr_fiss']} / {fr(P['mr_eta'])}", "", None),
            ("σ̄bc — Contrainte limite du béton", "0,6·fc28", fr(e["sbc_lim"], 1), "MPa", None),
            ("σ̄s — Contrainte limite de l'acier tendu", "selon la fissuration", fr(e["ss_lim"], 1), "MPa", None),
            ("x — Position de l'axe neutre", "équilibre N(x) = Ns (dichotomie)", fr(e["x"], 4), "m", None),
            ("σbc — Contrainte du béton retenue", "", fr(e["sbc"], 2), "MPa", None),
            ("σ's / σs — Contraintes des armatures", "n·σbc·(x−d')/x ; n·σbc·(d−x)/x",
             f"{fr(e['sp'], 1)} / {fr(e['st'], 1)}", "MPa", None),
            ("Mrs — Moment résistant ELS", "M(x) / cdg de la section", fr(e["Mrs"], 1), "kN.m", None),
            ("Critère déterminant", "", e["crit"], "", None),
            ("Convergence", "|N(x) − Ns| < 0,05 kN", "OK" if e["conv"] else "Ns hors domaine", "", e["conv"]),
        ]
        if mr["Ms_in"] is not None:
            rows += [("Ms / Mrs", "|Ms| / Mrs", fr(mr["ratio_s"], 3), "-", mr["ok_s"]),
                     ("Vérification à l'ELS", f"Ms = {fr(abs(mr['Ms_in']), 1)} kN.m",
                      "Section vérifiée" if mr["ok_s"] else "Section NON vérifiée", "", mr["ok_s"])]
        S.append(("RÉSULTATS À L'ELS — section fissurée homogénéisée", rows))
    return S


def draw_mr(fig, P, mr):
    """Section, axe neutre et diagramme des déformations à l'ELU (équivalent de l'onglet « Dessin » du classeur)."""
    from matplotlib.patches import Rectangle, Polygon
    fig.clear()
    e = mr["elu"]
    if not e:
        ax = fig.add_subplot(111)
        ax.axis("off")
        ax.text(0.5, 0.5, "Calcul à l'ELU désactivé", ha="center", va="center", fontsize=9, color="gray")
        return fig
    b, h, dp, c = P["mr_b"], P["mr_h"], P["mr_dp"], P["mr_c"]
    y = e["y"]
    gs = fig.add_gridspec(1, 2, width_ratios=[1, 1.15], wspace=0.12, left=0.04, right=0.97, top=0.88, bottom=0.05)
    # ---- section (hauteur en ordonnée, 0 = fibre inférieure)
    ax = fig.add_subplot(gs[0, 0])
    ax.set_title("Section et axe neutre", fontsize=8.5, fontweight="bold")
    ax.add_patch(Rectangle((0, 0), b, h, fc="#eef0f2", ec="black", lw=1.1))
    ycomp = max(0.0, h - min(y, h))                       # ordonnée (depuis le bas) de l'axe neutre
    ax.add_patch(Rectangle((0, h - min(y, h)), b, min(y, h), fc="#cfe3f7", ec="none", hatch="///", zorder=1))
    ax.plot([-0.45 * b, 1.45 * b], [ycomp, ycomp], color="#c0392b", lw=1.3, ls="--")
    ax.text(1.5 * b, ycomp, f"axe neutre y = {fr(y, 3)} m ({e['pivot']})", fontsize=6.5, va="center", color="#c0392b")
    ax.plot([0.15 * b, 0.85 * b], [h - dp, h - dp], color="#111", lw=2.4)
    ax.plot([0.15 * b, 0.85 * b], [c, c], color="#111", lw=2.4)
    ax.text(1.5 * b, h - dp, f"A' = {fr(P['mr_Ap'])} cm²", fontsize=6.5, va="center")
    ax.text(1.5 * b, c, f"A = {fr(P['mr_A'])} cm²", fontsize=6.5, va="center")
    ax.annotate("", (-0.35 * b, 0), (-0.35 * b, h), arrowprops=dict(arrowstyle="<->", lw=0.7))
    ax.text(-0.5 * b, h / 2, f"h = {fr(h)} m", rotation=90, va="center", ha="right", fontsize=6.5)
    ax.text(b / 2, -0.04 * h, f"b = {fr(b)} m", ha="center", va="top", fontsize=6.5)
    ax.set_xlim(-1.6 * b, 8.5 * b)
    ax.set_ylim(-0.1 * h, 1.04 * h)
    ax.axis("off")
    # ---- diagramme des déformations (compression > 0, en ‰)
    av = fig.add_subplot(gs[0, 1])
    av.set_title("Diagramme des déformations ε (‰) — compression > 0", fontsize=8.5, fontweight="bold")
    e_top, e_bot = e["K"] * y, e["K"] * (y - h)
    av.add_patch(Polygon([(0, 0), (e_bot, 0), (e_top, h), (0, h)], closed=True, fc="#e8eef5", ec="#1f4e79", lw=1))
    av.axvline(0, color="black", lw=0.8)
    av.plot([e["es1"]], [h - dp], "ko", ms=4)
    av.plot([e["es2"]], [c], "ko", ms=4)
    av.text(e_top, h, f" {fr(e_top, 2)}", fontsize=6.5, va="bottom", ha="left")
    av.text(e_bot, 0, f" {fr(e_bot, 2)}", fontsize=6.5, va="top", ha="center")
    av.text(e["es1"], h - dp, f"  ε' = {fr(e['es1'], 2)}", fontsize=6.3, va="bottom")
    av.text(e["es2"], c, f"  ε = {fr(e['es2'], 2)}", fontsize=6.3, va="bottom")
    av.axhline(ycomp, color="#c0392b", lw=0.9, ls="--")
    lo, hi = min(e_bot, 0) * 1.25, max(e_top, 0) * 1.25 + 1
    av.set_xlim(lo - 1, hi + 1)
    av.set_ylim(-0.1 * h, 1.1 * h)
    av.set_yticks([])
    av.tick_params(labelsize=6.5)
    for sp in ("top", "right", "left"):
        av.spines[sp].set_visible(False)
    fig.suptitle(f"Mru = {fr(e['Mru'], 0)} kN.m pour Nu = {fr(P['mr_Nu'], 0)} kN", fontsize=8, y=0.985)
    return fig


# ---------------------------------------------------------------------------
#  4e. SCHÉMA EN ÉLÉVATION DU VOILE + ENVELOPPE DES MOMENTS
# ---------------------------------------------------------------------------
def draw_elevation(fig, P, res, multi):
    """Élévation du voile sur tous les étages (rives, zone critique, As par niveau) + diagramme des moments."""
    from matplotlib.patches import Rectangle
    fig.clear()
    env, rows = multi["env"], multi["rows"]
    G, D = res["G"], res["D"]
    hw = env["hw"]
    el_ = multi["elanc"]
    Lmax = max(r["L"] for r in rows)
    gs = fig.add_gridspec(1, 2, width_ratios=[1.3, 1.0], wspace=0.05, left=0.05, right=0.985, top=0.93, bottom=0.06)
    ax = fig.add_subplot(gs[0, 0])
    am = fig.add_subplot(gs[0, 1], sharey=ax)
    lcg, lcd = G["lc"] / 100, D["lc"] / 100
    hcr = env["hcr"]
    xlab = Lmax + 0.35
    for k, r in enumerate(rows):
        y0, y1 = env["cum"][k], env["cum"][k + 1]
        L = r["L"]
        ax.add_patch(Rectangle((0, y0), L, y1 - y0, fc="#f4f5f7", ec="black", lw=0.9, zorder=1))
        for (x0, w) in ((0, min(lcg, L / 2)), (L - min(lcd, L / 2), min(lcd, L / 2))):
            crit = y1 <= hcr + 1e-9
            ax.add_patch(Rectangle((x0, y0), w, y1 - y0, fc="#cfe3f7", ec="#1f4e79", lw=0.7,
                                   hatch="////" if crit else "..", zorder=2))
        col = "#1b7a3d" if r["ok"] else "#b3261e"
        ax.text(xlab, (y0 + y1) / 2 + 0.28, f"{r['name']}   {'OK' if r['ok'] else 'NG'}", va="center", fontsize=6.6,
                color=col, fontweight="bold")
        ax.text(xlab, (y0 + y1) / 2 - 0.34, f"e={fr(r['e'] * 100, 0)} cm · {r['type']} · As req.={fr(r['As'], 1)} cm²",
                va="center", fontsize=6.0, color=col)
        if r.get("ok_as") is not None:                   # section adoptée (facultative)
            cas_ = "#1b7a3d" if r["ok_as"] else "#b3261e"
            ax.text(xlab, (y0 + y1) / 2 - 0.92, f"Adopté {r['sec'].replace(' + ', '+')} = {fr(r['As_ch'], 1)} cm² "
                    f"{'✓' if r['ok_as'] else '✗ insuffisant'}", va="center", fontsize=5.8, color=cas_,
                    fontweight="bold")
    # zone critique
    ax.add_patch(Rectangle((-0.12, 0), Lmax + 0.24, min(hcr, hw), fc="#e74c3c", alpha=0.10, ec="#c0392b", ls="--",
                           lw=1.0, zorder=3))
    ax.annotate("", (-0.5, 0), (-0.5, min(hcr, hw)), arrowprops=dict(arrowstyle="<->", color="#c0392b", lw=0.8))
    ax.text(-0.58, min(hcr, hw) / 2, f"hcr = {fr(hcr)} m", rotation=90, va="center", ha="right", fontsize=7,
            color="#c0392b")
    ax.annotate("", (-1.45, 0), (-1.45, hw), arrowprops=dict(arrowstyle="<->", lw=0.8))
    ax.text(-1.53, hw / 2, f"hw = {fr(hw)} m", rotation=90, va="center", ha="right", fontsize=7)
    ax.set_xlim(-2.2, xlab + 5.2)
    ax.set_ylim(-1.9, hw + 0.9)
    ax.set_aspect("equal", adjustable="box")
    ax.axis("off")
    ax.set_title(f"Élévation du voile — tous les étages\nhw/lw = {fr(el_['ratio'])} "
                 f"{'>' if el_['elance'] else '≤'} {fr(el_['lim'], 1)} : {el_['cls'].lower()}",
                 fontsize=8.5, fontweight="bold")
    ax.text(0, -0.55, f"Rives : lcg = {fr(G['lc'], 0)} cm, lcd = {fr(D['lc'], 0)} cm — "
            f"HA{fr(G['phi_l'] * 10, 0)} ; cadres HA{fr(G['phi_c'] * 10, 0)} St = {fr(G['St'], 0)} cm",
            fontsize=6.3, va="top")
    ax.text(0, -1.25, "Hachures : zone critique (confinement) — points : rives courantes",
            fontsize=5.8, va="top", color="gray", style="italic")

    # diagramme des moments
    el = multi["elanc"]
    am.set_title("Enveloppe des moments" if el["elance"] else "Moments du calcul de structure (voile court)",
                 fontsize=9, fontweight="bold")
    Hs = [0, hw]
    if el["elance"] or multi["mode_eff"] != "etabs":
        am.plot([env["mbase"], env["msom"]], Hs, ls="--", color="gray", lw=1.1, label="Enveloppe linéaire")
        am.plot([env["mbase"], env["mbase"], env["m_top"]], [0, hcr, hw], color="#c0392b", lw=1.8,
                label="Enveloppe de calcul (décalée de hcr)")
    base_pts = [(abs(r["M_etabs"]), env["cum"][k]) for k, r in enumerate(rows)]
    am.scatter([p[0] for p in base_pts], [p[1] for p in base_pts], s=14, c="#1f4e79", zorder=4,
               label="M3 ETABS (base des étages)")
    am.scatter([r["M"] for r in rows],
               [env["cum"][k + (1 if env["ref"] == "sommet" else 0)] for k, r in enumerate(rows)],
               s=22, marker="s", fc="none", ec="#1b7a3d", zorder=4, label="M de calcul retenu")
    for yy in env["cum"]:
        am.axhline(yy, color="#dddddd", lw=0.5, zorder=0)
    am.axhline(hcr, color="#c0392b", lw=0.8, ls=":")
    am.set_xlim(0, max(env["mbase"], max(abs(r["M_etabs"]) for r in rows)) * 1.08)
    am.set_xlabel("M (kN·m)", fontsize=7.5)
    am.tick_params(labelsize=6.5, labelleft=False)
    am.grid(False)
    am.legend(fontsize=5.8, loc="upper right", frameon=True)
    for sp in ("top", "right"):
        am.spines[sp].set_visible(False)
    return fig


# ---------------------------------------------------------------------------
#  5. NOTE DE CALCUL A4 — MODÈLE COMMUN, RENDU PDF (reportlab) ET WORD (python-docx)
# ---------------------------------------------------------------------------
def fig_png(draw, figsize, *args, dpi=200):
    """Dessine une figure matplotlib et retourne (octets PNG, largeur, hauteur en pouces)."""
    from matplotlib.figure import Figure
    fig = Figure(figsize=figsize, dpi=dpi)
    draw(fig, *args)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi)
    return buf.getvalue(), figsize[0], figsize[1]


def _txt(x, nd=2):
    return x if isinstance(x, str) else fr(x, nd)


def report_blocks(P, res, multi=None):
    """Construit la note de calcul sous forme de blocs indépendants du format de sortie."""
    B = []
    G, D, Wd = res["G"], res["D"], res["W"]
    mr = None
    if int(P.get("mr_report", 0)):
        try:
            mr = compute_mr(P)
        except ValueError:
            mr = None                       # données du moment résistant invalides : rubrique omise de la note
    all_ok = res["all_ok"] and (multi["all_ok"] if multi else True) and (mr["all_ok"] if mr else True)
    ko = [n for n, ok in res["checks"] if not ok]
    if mr:
        ko = [n for n, ok in mr["checks"] if not ok] + ko
    if multi:
        ko = [f"Section adoptée insuffisante — {r['name']}" for r in multi["rows"] if r["ok_as"] is False] + ko
        ko = [f"Méthode des contraintes — {r['name']}" for r in multi["rows"] if not r["ok"]] + ko
    lt = iter("ABCDEFGH")                   # lettres des parties de la note

    # ---------------- page de garde
    sub = ("Voile en béton armé — enveloppes des moments et de l'effort tranchant, méthode des contraintes "
           "(tous étages),<br/>moment résistant, ferraillage des rives et ductilité locale — RPA 2024" if multi else
           "Ferraillage des éléments de rive et vérification de la ductilité locale d'un voile<br/>"
           "en béton armé — RPA 2024 (Art. 7.7.5)")
    B.append({"t": "title", "text": "NOTE DE CALCUL", "sub": sub})
    B.append({"t": "info", "rows": [
        ["Projet", P["projet"], "Élément", P["element"]],
        ["Niveau", P["niveau"], "Date", P["date"]],
        ["Établi par", P["auteur"] or "—", "Matériaux",
         f"Béton fc28 = {fr(P['fc'], 0)} MPa — Acier fe = {fr(P['fe'], 0)} MPa"]]})
    if all_ok:
        B.append({"t": "banner", "text": "CONCLUSION : toutes les vérifications sont satisfaites.", "ok": True})
    else:
        B.append({"t": "banner", "ok": False,
                  "text": f"CONCLUSION : {len(ko)} vérification(s) non satisfaite(s) — voir détail ci-dessous."})
        B.append({"t": "list", "items": ko})

    syn = [["Dimensions lc × bc (cm)", f"{fr(G['lc'], 0)} × {fr(G['bc'], 0)}", f"{fr(D['lc'], 0)} × {fr(D['bc'], 0)}"],
           ["Armatures verticales de rive",
            f"{G['lay']['n_ret']} HA{fr(G['phi_l'] * 10, 0)} ({fr(G['lay']['Av_f'])} cm²)",
            f"{D['lay']['n_ret']} HA{fr(D['phi_l'] * 10, 0)} ({fr(D['lay']['Av_f'])} cm²)"],
           ["Cadres de confinement", f"HA{fr(G['phi_c'] * 10, 0)} St = {fr(G['St'], 0)} cm, {G['nbr']} br.",
            f"HA{fr(D['phi_c'] * 10, 0)} St = {fr(D['St'], 0)} cm, {D['nbr']} br."],
           ["Âme verticale", f"{Wd['nv']} HA{fr(P['phi_v'], 0)} /face, St = {fr(P['st_v'], 0)} cm", "="],
           ["Âme horizontale", f"HA{fr(P['phi_h'], 0)}, St = {fr(P['st_h'], 0)} cm (2 nappes)", "="],
           ["Longueur de confinement", G["lc_adopt"], D["lc_adopt"]]]
    B.append({"t": "table", "title": None, "head": ["Synthèse du ferraillage (rives)", "Rive gauche", "Rive droite"],
              "rows": syn, "widths": [60, 60, 60], "align": "lcc", "span_eq": True})
    if multi:
        env = multi["env"]
        B.append({"t": "table", "title": None,
                  "head": ["Synthèse — tous étages", "Valeur", "Statut"],
                  "rows": [["Cas de calcul actif", multi["cas"]["name"], ""],
                           ["Nombre de niveaux calculés / hauteur hw", f"{env['n']} niveaux — hw = {fr(env['hw'])} m", ""],
                           ["Zone critique hcr (éq. 7.22 / 7.23)", f"{fr(env['hcr'])} m", ""],
                           ["As requise maximale (par niveau)", f"{fr(multi['As_max'])} cm²", ""],
                           ["σc/fbu max — contrainte du béton", fr(multi["max_sc"], 3), ("OK" if multi["max_sc"] <= 1 else "NG", multi["max_sc"] <= 1)],
                           ["τb/τlim max — cisaillement", fr(multi["max_tau"], 3), ("OK" if multi["max_tau"] < 1 else "NG", multi["max_tau"] < 1)],
                           ["Niveaux non vérifiés", str(multi["n_ng"]), ("OK" if multi["n_ng"] == 0 else "NG", multi["n_ng"] == 0)]]
                          + ([["Sections adoptées (choix de l'ingénieur)",
                               f"{multi['n_sec']} niveau(x) — ΣAs choisie = {fr(multi['sumAs_ch'], 1)} cm²",
                               ("OK" if multi["n_as_ko"] == 0 else "NG", multi["n_as_ko"] == 0)]]
                             if multi["has_sec"] else []),
                  "widths": [80, 70, 30], "align": "lcc"})

    # ---------------- hypothèses sismiques et matériaux
    if int(P.get("seis_link", 0)) or int(P.get("mat_link", 0)):
        try:
            hrows = []
            if int(P.get("seis_link", 0)):
                sp = seismic_params(P)
                hrows += [["Zone sismique / niveau de sismicité", f"{P['sz_zone']} — {sp['niv']}", "Table 3.3", ""],
                          ["A — Coefficient d'accélération de zone", fr(sp["A"], 2), "Table 3.3", "-"],
                          ["Groupe d'importance / I", f"{P['sz_groupe']} / {fr(sp['I'], 1)}", "Table 3.10", "-"],
                          ["Site / S", f"{P['sz_site']} / {fr(sp['S'], 2)}", "Table 3.3", "-"],
                          ["T1 ; T2 ; T3", f"{fr(sp['T1'], 2)} ; {fr(sp['T2'], 2)} ; {fr(sp['T3'], 2)}", "Table 3.3", "s"],
                          ["Système de contreventement", f"n° {sp['n_sys']} — {sp['sys']}", "Table 3.18", ""],
                          ["R — Coefficient de comportement", fr(sp["R"], 2), "Table 3.18", "-"],
                          ["QF = 1 + ΣPq (catégorie " + sp["cat"] + ")", fr(sp["QF"], 2), "Table 3.19", "-"],
                          ["T = CT·hN^(3/4)", f"{fr(sp['T'], 3)}  (CT = {fr(P['sz_ct'], 2)} ; hN = {fr(P['sz_hn'])} m)",
                           "", "s"],
                          ["ξ / η = √(7/(2+ξ))", f"{fr(sp['xi'], 0)} % / {fr(sp['eta'], 3)}", "Table 3.5", ""]]
            if int(P.get("mat_link", 0)):
                mp = material_params(P)
                hrows += [["Béton — fc28 (voiles) / Ev / Ei", f"{fr(mp['fc'], 0)} MPa / {fr(mp['Ev'], 0)} / {fr(mp['Ei'], 0)} MPa",
                           "", ""],
                          ["Aciers HA — fe", f"{fr(mp['fe'], 0)} MPa", "", ""],
                          [f"Situation de calcul : {mp['sit']}", f"γb = {fr(mp['gb'], 2)} ; γs = {fr(mp['gs'], 2)}", "", ""]]
            B.append({"t": "table", "title": f"{next(lt)}. HYPOTHÈSES SISMIQUES ET MATÉRIAUX — RPA 2024",
                      "intro": "Valeurs choisies dans les rubriques « Données sismiques » et « Matériaux », reprises dans "
                               "les données générales (R, QF, T2, T0, fc28, fe, γb, γs).",
                      "head": ["Paramètre", "Valeur", "Référence", "Unité"], "rows": hrows,
                      "widths": [66, 78, 22, 14], "align": "llcc", "fs": 7.6})
        except ValueError:
            pass

    # ---------------- parties multi-étages
    if multi:
        env, rows = multi["env"], multi["rows"]
        cas = multi["cas"]
        S = P["stories"]
        B.append({"t": "pb"})
        B.append({"t": "table", "title": f"{next(lt)}. EFFORTS SOURCES — EXTRAIT « PIER FORCES » D'ETABS",
                  "intro": (f"Fichier : {os.path.basename(P.get('etabs_path') or '') or 'saisie manuelle'} — pier : "
                            f"{P.get('pier') or '—'}. Section du bas de chaque niveau ; N = P correspondant au "
                            f"M3 de valeur absolue maximale ; V2 de valeur absolue maximale (toutes combinaisons). "
                            f"Convention : N < 0 = compression."),
                  "head": ["Niveau", "h (m)", "L (m)", "e (m)", "N [P] (kN)", "V2 (kN)", "M3 ETABS (kN.m)", "P min (kN)"],
                  "rows": [[s["name"], fr(s["h"]), fr(s["L"]), fr(s["e"]), fr(s["N"]), fr(s["V2"]), fr(s["M3"]),
                            fr(s.get("Pmin", 0.0))] for s in S],
                  "widths": [28, 16, 16, 16, 26, 24, 32, 22], "align": "lrrrrrrr", "fs": 7.4})

        ref_txt = ("sommet de l'étage (convention du classeur d'origine)" if env["ref"] == "sommet"
                   else "base de l'étage (section du bas — sécuritaire)")
        mode_txt = {"enveloppe": "moment de l'enveloppe", "max": "max(enveloppe ; M3 ETABS)",
                    "etabs": "M3 ETABS brut"}.get(multi["mode_eff"], multi["mode_eff"])
        if multi["mode"] == "auto":
            mode_txt = f"automatique (Art. 7.7.4) : {mode_txt}"
        el = multi["elanc"]
        prm = [["n — Nombre d'étages", "", str(env["n"]), "-"],
               ["he — Hauteur libre d'étage", "", fr(env["he"]), "m"],
               ["lw — Longueur du voile", "", fr(env["lw"]), "m"],
               ["hw — Hauteur totale du voile", "Σ des hauteurs d'étage", fr(env["hw"]), "m"],
               ["M base — moment max. à la base", "|M3| du niveau le plus bas" if P["env_auto"] else "saisie",
                fr(env["mbase"]), "kN.m"],
               ["M sommet — moment du dernier niveau", "|M3| du niveau le plus haut" if P["env_auto"] else "saisie",
                fr(env["msom"]), "kN.m"],
               ["hw/lw — Élancement du voile", "hw / lw", fr(el["ratio"]), "-"],
               ["Condition d'élancement (Art. 7.7.4)", f"élancé si hw/lw > {fr(el['lim'], 1)}",
                f"{el['cls']} : {el['regle']}", ""],
               ["hcr1", "max(lw ; hw/6)  [éq. 7.22]", fr(env["hcr1"]), "m"],
               ["Limite", "min(2·lw ; he si n ≤ 6 ; 2·he si n ≥ 7)  [éq. 7.23]", fr(env["limite"]), "m"],
               ["hcr retenue", "min(hcr1 ; limite)", fr(env["hcr"]), "m"],
               ["Pente de l'enveloppe", "(M base − M sommet) / hw", fr(env["pente"]), "kN.m/m"],
               ["M enveloppe décalée en tête", "M base − pente·(hw − hcr)", fr(env["m_top"]), "kN.m"],
               ["Hauteur de référence par niveau", "", ref_txt, ""],
               ["Moment de calcul retenu", "", mode_txt, ""]]
        B.append({"t": "table", "title": f"{next(lt)}. ENVELOPPE DES MOMENTS FLÉCHISSANTS — RPA 2024 (éq. 7.22 et 7.23)",
                  "head": ["Désignation", "Expression", "Valeur", "Unité"], "rows": prm,
                  "widths": [58, 66, 38, 18], "align": "llcc", "fs": 7.4})
        trow = []
        for t in env["table"]:
            lab = "Base du voile (H = 0)" if t["k"] == 0 else S[t["k"] - 1]["name"]
            hh = "" if t["k"] == 0 else fr(S[t["k"] - 1]["h"])
            trow.append([lab, hh, fr(t["H"]), fr(t["M_lin"], 1), fr(t["M_corr"], 1), t["zone"]])
        B.append({"t": "table", "title": None,
                  "head": ["Niveau", "h (m)", "H cumulée (m)", "M linéaire (kN.m)", "M corrigé (kN.m)", "Zone"],
                  "rows": trow, "widths": [40, 16, 26, 30, 30, 38], "align": "lrrrrl", "fs": 7.2})

        # --- enveloppe de l'effort tranchant
        envv = multi["envv"]
        vtxt = ("V corrigé de l'enveloppe (1,4 déjà inclus en zone critique ; kV = 1 pour τb)"
                if multi["v_mode"] == "enveloppe" else f"kV × V2 ETABS (kV = {fr(P['kV'])})")
        prv = [["hcr — Zone critique (reprise)", "éq. 7.22 / 7.23", fr(envv["hcr"]), "m"],
               ["hw/3 — Seuil de retour au diagramme non majoré", "hw / 3", fr(envv["h3"]), "m"],
               ["V(hcr) — Effort tranchant d'analyse à h = hcr", "interpolation linéaire entre niveaux",
                fr(envv["v_hcr"], 1), "kN"],
               ["V(hw/3) — Effort tranchant d'analyse à h = hw/3", "interpolation linéaire entre niveaux",
                fr(envv["v_h3"], 1), "kN"],
               ["Facteur d'amplification", "capacity design (zone critique)", fr(envv["k"]), "-"],
               ["V de calcul retenu (méthode des contraintes)", "", vtxt, ""]]
        B.append({"t": "table", "title": f"{next(lt)}. ENVELOPPE DE L'EFFORT TRANCHANT — RPA 2024 (Art. 7.7.2)",
                  "intro": ("Dans la zone critique (h ≤ hcr), V issu de l'analyse est majoré de 1,4. Entre hcr et hw/3, "
                            "transition linéaire entre (hcr ; 1,4·V(hcr)) et (hw/3 ; V(hw/3)). Au-delà de hw/3, V "
                            "d'analyse non majoré. Valeurs prises à la section du bas de chaque niveau."),
                  "head": ["Désignation", "Expression", "Valeur", "Unité"], "rows": prv,
                  "widths": [66, 52, 42, 14], "align": "llcc", "fs": 7.4})
        B.append({"t": "table", "title": None,
                  "head": ["Niveau", "H cumulée (m)", "V analyse (kN)", "1,4 × V (kN)", "V corrigé (kN)", "Zone"],
                  "rows": [[S[t["k"]]["name"], fr(t["H"]), fr(t["Ti"], 1), fr(t["T14"], 1), fr(t["Vc"], 1), t["zone"]]
                           for t in envv["table"]],
                  "widths": [34, 24, 28, 28, 28, 38], "align": "lrrrrl", "fs": 7.2})

        # --- méthode des contraintes
        cp = (f"Cas actif : {cas['name']} — fc28 = {fr(cas['fc'], 1)} MPa, fe = {fr(cas['fe'], 0)} MPa, "
              f"γb = {fr(cas['gb'])}, γs = {fr(cas['gs'])}, θ = {fr(cas['th'])} → fbu = 0,85·fc28/(θ·γb) = "
              f"{fr(cas['fbu'])} MPa ; σs = fe/γs = {fr(cas['sigs'], 1)} MPa. kV = {fr(P['kV'])}, kd = {fr(P['kd'])} "
              f"(d = kd × {'L du voile' if str(P['dref']).lower().startswith('l') else 'h étage'}), "
              f"τlim = kτ·fc28 = {fr(P['ktau'] * cas['fc'])} MPa, ρmin zone tendue = {fr(P['rho_t'] * 100, 2)} %, "
              f"ρmin S.E.C = {fr(P['rho_g'] * 100, 2)} %. Effort tranchant de calcul : {vtxt}.")
        mrows = []
        for r in rows:
            mrows.append([r["name"], fr(r["N"], 0), fr(r["V2"], 0), fr(r["M"], 0),
                          f"{fr(r['L'])}×{fr(r['e'])}", fr(r["s1"]), fr(r["s2"]), r["type"], fr(r["Lt"]),
                          fr(r["As_calc"], 1), fr(r["As_min"], 1), fr(r["As"], 1), (fr(r["sc_fbu"]), r["ok_c"]),
                          (fr(r["tau_lim"]), r["ok_t"]), ("OK" if r["ok"] else "NG", r["ok"])])
        B.append({"t": "table", "title": f"{next(lt)}. MÉTHODE DES CONTRAINTES — TOUS LES ÉTAGES", "intro": cp,
                  "head": ["Niveau", "N (kN)", "V2 (kN)", "M calc. (kN.m)", "L×e (m)", "σ1 (MPa)", "σ2 (MPa)", "Type",
                           "Lt (m)", "As cal. (cm²)", "As min (cm²)", "As requise (cm²)", "σc/fbu", "τb/τlim", "Statut"],
                  "rows": mrows,
                  "widths": [15, 11, 11, 13, 14, 11, 11, 10, 9, 11, 11, 12, 11, 13, 12], "align": "lrrrcrrcrrrrccc",
                  "fs": 6.2})
        B.append({"t": "para", "text": (
            "σ1,2 = N/A ± M·v/I (A = L·e, I = e·L³/12, v = L/2) ; Lt = L·σmax/(σmax − σmin) en S.P.C ; "
            "Ft = σmax·e·Lt/2 (S.P.C) ou (σ1+σ2)·e·L/2 (S.E.T) ; As,cal = Ft/σs ; As requise = max(As,cal ; As,min) ; "
            "|σc|max ≤ fbu ; τb = V de calcul/(d·e) ≤ τlim.")})
        if multi["has_sec"]:                                   # section adoptée : colonne facultative
            B.append({"t": "table", "title": None,
                      "intro": ("Section d'acier ADOPTÉE (choix de l'ingénieur, facultatif) : As choisie = Σ n·π·Ø²/4 "
                                "des barres retenues, comparée à As requise = max(As,cal ; As,min)."),
                      "head": ["Section adoptée — niveau", "Type", "Lt (m)", "As requise (cm²)", "Section choisie",
                               "As choisie (cm²)", "As ch. / As req.", "Statut"],
                      "rows": [[r["name"], r["type"], fr(r["Lt"]), fr(r["As"], 1), r["sec"] or "—",
                                fr(r["As_ch"], 1) if r["As_ch"] is not None else "—",
                                fr(r["ratio_as"], 2) if r["ratio_as"] is not None else "—",
                                ("—" if r["ok_as"] is None else ("OK" if r["ok_as"] else "NG", r["ok_as"]))]
                               for r in rows],
                      "widths": [30, 13, 13, 24, 34, 24, 22, 20], "align": "lcrrcrrc", "fs": 7.2})
        crow = [[c["cas"]["name"], fr(c["cas"]["gb"]), fr(c["cas"]["gs"]), fr(c["cas"]["fbu"]),
                 fr(c["cas"]["sigs"], 1), fr(c["sumAs"], 1), (str(c["n_ng"]), c["n_ng"] == 0), fr(c["max_sc"], 3),
                 fr(c["max_tau"], 3)] for c in multi["comp"]]
        B.append({"t": "table", "title": None,
                  "head": ["Comparaison des cas", "γb", "γs", "fbu (MPa)", "σs (MPa)", "ΣAs (cm²)", "Niveaux NG",
                           "σc/fbu max", "τ/τlim max"],
                  "rows": crow, "widths": [42, 12, 12, 20, 20, 22, 18, 17, 17], "align": "lccrrrccc", "fs": 7.2})

    # ---------------- moment résistant (BAEL 91 / RPA 2024)
    if mr:
        B.append({"t": "pb"})
        B.append({"t": "bar", "text": f"{next(lt)}. MOMENT RÉSISTANT — FLEXION COMPOSÉE (BAEL 91 modifié 99 / RPA 2024)"})
        B.append({"t": "para", "text": (
            "Section rectangulaire b × h, armatures supérieures A' et inférieures A. ELU : sections planes, béton "
            "tendu négligé, règle des trois pivots, béton parabole-rectangle (fbu = 0,85·fc28/(θ·γb)), acier "
            "élasto-plastique parfait (fsu = fe/γs) ; l'équilibre N(y) = Nu est résolu par dichotomie. ELS : "
            "section fissurée homogénéisée (n = 15), σbc ≤ 0,6·fc28, σs ≤ σ̄s selon la fissuration ; Mrs est le "
            "moment maximal compatible avec Ns. Moments pris par rapport au centre de gravité du béton (h/2).")})
        for title, rws in mr["sections"]:
            B.append({"t": "table", "title": title, "head": ["Désignation", "Expression", "Valeur", "Unité"],
                      "rows": [[lab, ex, (val, ok), un] for lab, ex, val, un, ok in rws],
                      "widths": [72, 54, 40, 14], "align": "llcc", "fs": 7.4, "italic_col": 1})
        if mr["elu"]:
            png = fig_png(draw_mr, (7.1, 4.4), P, mr)
            B.append({"t": "image", "png": png[0], "w": 170, "h": 170 * png[2] / png[1]})

    if multi:
        # --- élévation
        png = fig_png(draw_elevation, (7.1, 9.2), P, res, multi)
        B.append({"t": "pb"})
        B.append({"t": "bar", "text": f"{next(lt)}. SCHÉMA EN ÉLÉVATION DU VOILE ET ENVELOPPE DES MOMENTS"})
        B.append({"t": "image", "png": png[0], "w": 178, "h": 178 * png[2] / png[1]})
    if multi or mr:
        B.append({"t": "pb"})
        B.append({"t": "bar", "text": f"{next(lt)}. FERRAILLAGE DES RIVES ET DUCTILITÉ LOCALE — {P['element'].upper()}"})

    # ---------------- sections de la ductilité locale (voile / étage de calcul)
    for title, heads, rws in res["sections"]:
        data = []
        for r in rws:
            t1, t2 = r.texts()
            data.append([r.label, r.expr, (t1, r.ok1), (t2, r.ok2), r.unit])
        B.append({"t": "table", "title": title, "head": ["Désignation", "Expression", heads[0], heads[1], "Unité"],
                  "rows": data, "widths": [60, 54, 24, 24, 18], "align": "llccc", "fs": 7.6, "italic_col": 1,
                  "bg_status": True})

    cad = {sd: cadres_cumul(P, sd) for sd in ("g", "d")}
    if any(c["rows"] for c in cad.values()):
        crow = []
        for sd, nm in (("g", "Rive gauche"), ("d", "Rive droite")):
            for r in cad[sd]["rows"]:
                crow.append([nm, r["type"], fr(r["n"], 0), fr(r["l"], 1), fr(r["sub"], 1)])
            if cad[sd]["rows"]:
                crow.append([nm, "Longueur cumulée L", "", "", (fr(cad[sd]["total"], 1), True)])
        B.append({"t": "table", "title": "CADRES DE CONFINEMENT — LONGUEUR CUMULÉE DES BRANCHES (L)",
                  "intro": "Types, nombre et longueur des cadres retenus par l'ingénieur ; L = Σ (nombre × longueur).",
                  "head": ["Rive", "Type de cadre", "Nombre", "Longueur unitaire (cm)", "Σ (cm)"], "rows": crow,
                  "widths": [30, 60, 22, 36, 30], "align": "llccc", "fs": 7.6})
    png = fig_png(draw_schema, (7.1, 9.4), P, res)
    B.append({"t": "pb"})
    B.append({"t": "bar", "text": "9. SCHÉMA DE FERRAILLAGE — ÉLÉVATION DE LA ZONE DUCTILE, COUPE EN PLAN ET DÉTAIL DES RIVES"})
    B.append({"t": "image", "png": png[0], "w": 178, "h": 178 * png[2] / png[1]})
    return B


def _cell_split(c):
    return (c[0], c[1]) if isinstance(c, tuple) else (c, None)


# ------------------------------- rendu PDF ---------------------------------
def _register_fonts():
    import matplotlib
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.fonts import addMapping
    base = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
    files = {"DV": "DejaVuSans.ttf", "DV-B": "DejaVuSans-Bold.ttf", "DV-I": "DejaVuSans-Oblique.ttf",
             "DV-BI": "DejaVuSans-BoldOblique.ttf"}
    for name, f in files.items():
        pdfmetrics.registerFont(TTFont(name, os.path.join(base, f)))
    addMapping("DV", 0, 0, "DV")
    addMapping("DV", 1, 0, "DV-B")
    addMapping("DV", 0, 1, "DV-I")
    addMapping("DV", 1, 1, "DV-BI")


def render_pdf(blocks, P, path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak,
                                    KeepTogether)
    from reportlab.pdfgen import canvas as rl_canvas

    _register_fonts()
    NAVY = colors.HexColor("#1f3a5f")
    GREEN, RED = colors.HexColor("#1b7a3d"), colors.HexColor("#b3261e")
    W_TOT = 180 * mm
    AL = {"l": TA_LEFT, "c": TA_CENTER, "r": TA_RIGHT}

    def pst(name, fs, font="DV", al=TA_LEFT, color=colors.black):
        return ParagraphStyle(name, fontName=font, fontSize=fs, leading=fs * 1.28, alignment=al, textColor=color)

    title_st = ParagraphStyle("t", fontName="DV-B", fontSize=15, leading=19, alignment=TA_CENTER, textColor=NAVY)
    sub_st = pst("s", 9.5, al=TA_CENTER, color=colors.grey)
    h_st = pst("h", 10, "DV-B", color=colors.white)
    n_st = pst("n", 8.2)
    i_st = pst("i", 7.4, "DV-I", color=colors.HexColor("#333333"))

    class Numbered(rl_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._saved = []

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            n = len(self._saved)
            for state in self._saved:
                self.__dict__.update(state)
                self._decor(n)
                super().showPage()
            super().save()

        def _decor(self, n):
            w, h = A4
            self.setStrokeColor(NAVY)
            self.setLineWidth(0.8)
            self.line(15 * mm, h - 13 * mm, w - 15 * mm, h - 13 * mm)
            self.line(15 * mm, 14 * mm, w - 15 * mm, 14 * mm)
            self.setFont("DV", 7.5)
            self.setFillColor(colors.HexColor("#333333"))
            self.drawString(15 * mm, h - 11 * mm, f"{P['projet']} — {P['element']} ({P['niveau']})")
            self.drawRightString(w - 15 * mm, h - 11 * mm, "Note de calcul — Voile RPA 2024")
            self.drawString(15 * mm, 9.5 * mm, f"Établie le {P['date']}" + (f" par {P['auteur']}" if P["auteur"] else ""))
            self.drawRightString(w - 15 * mm, 9.5 * mm, f"Page {self._pageNumber} / {n}")

    def bar(text):
        t = Table([[Paragraph(text, h_st)]], colWidths=[W_TOT])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), NAVY), ("TOPPADDING", (0, 0), (-1, -1), 3),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        return t

    story = []
    for b in blocks:
        t = b["t"]
        if t == "title":
            story += [Spacer(1, 4 * mm), Paragraph(b["text"], title_st), Paragraph(b["sub"], sub_st),
                      Spacer(1, 6 * mm)]
        elif t == "info":
            cb, c = pst("cb", 7.8, "DV-B"), pst("c", 7.8)
            tb = Table([[Paragraph(str(x), cb if i % 2 == 0 else c) for i, x in enumerate(r)] for r in b["rows"]],
                       colWidths=[24 * mm, 66 * mm, 24 * mm, 66 * mm])
            tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#e8eef5")),
                                    ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#e8eef5")),
                                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
            story += [tb, Spacer(1, 5 * mm)]
        elif t == "banner":
            tb = Table([[Paragraph(f"<font color='white'><b>{b['text']}</b></font>", n_st)]], colWidths=[W_TOT])
            tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), GREEN if b["ok"] else RED),
                                    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            story += [tb, Spacer(1, 3 * mm)]
        elif t == "list":
            story += [Paragraph(f"<font color='#b3261e'>✗</font> {k}", pst("c", 7.8)) for k in b["items"]]
            story.append(Spacer(1, 3 * mm))
        elif t == "bar":
            story += [bar(b["text"]), Spacer(1, 3 * mm)]
        elif t == "para":
            story += [Paragraph(b["text"], i_st), Spacer(1, 3 * mm)]
        elif t == "pb":
            story.append(PageBreak())
        elif t == "image":
            story.append(Image(io.BytesIO(b["png"]), width=b["w"] * mm, height=b["h"] * mm))
        elif t == "table":
            fs = b.get("fs", 7.6)
            cst = {a: pst("c" + a, fs, al=AL[a]) for a in "lcr"}
            cib = pst("ci", fs - 0.5, "DV-I", color=colors.HexColor("#444444"))
            hst = {a: pst("h" + a, fs, "DV-B", al=AL[a]) for a in "lcr"}
            al = b["align"]
            data = [[Paragraph(h, hst["c" if al[j] != "l" else "l"]) for j, h in enumerate(b["head"])]]
            style = [("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#999999")),
                     ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef5")),
                     ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                     ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                     ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
            for i, row in enumerate(b["rows"], start=1):
                cells = []
                for j, c in enumerate(row):
                    txt, ok = _cell_split(c)
                    if b.get("italic_col") == j:
                        cells.append(Paragraph(txt, cib))
                        continue
                    if ok is None:
                        cells.append(Paragraph(txt, cst[al[j]]))
                    else:
                        col = "#1b7a3d" if ok else "#b3261e"
                        cells.append(Paragraph(f"<font color='{col}'><b>{txt}</b></font>", cst[al[j]]))
                        style.append(("BACKGROUND", (j, i), (j, i), colors.HexColor("#e6f4ea" if ok else "#fde7e9")))
                data.append(cells)
                if b.get("span_eq") and len(row) == 3 and row[2] == "=":
                    style.append(("SPAN", (1, i), (2, i)))
            cw = [x * mm for x in b["widths"]]
            tb = Table(data, colWidths=cw, repeatRows=1)
            tb.setStyle(TableStyle(style))
            parts = []
            if b.get("title"):
                parts.append(bar(b["title"]))
            if b.get("intro"):
                parts += [Spacer(1, 1.5 * mm), Paragraph(b["intro"], i_st), Spacer(1, 1.5 * mm)]
            parts.append(tb)
            story += [KeepTogether(parts)] if len(b["rows"]) <= 14 else parts
            story.append(Spacer(1, 4 * mm))

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=18 * mm,
                            bottomMargin=19 * mm, title=f"Note de calcul — {P['element']}", author=P["auteur"] or "")
    doc.build(story, canvasmaker=Numbered)
    return path


# ------------------------------- rendu WORD --------------------------------
def render_docx(blocks, P, path):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    NAVY = "1F3A5F"
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
    sec.header_distance = sec.footer_distance = Mm(8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(2)

    def shade(cell, hexcol):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcol)
        tcPr.append(shd)

    def margins(cell, top=30, bottom=30, left=60, right=60):
        tcPr = cell._tc.get_or_add_tcPr()
        m = OxmlElement("w:tcMar")
        for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            e = OxmlElement(f"w:{k}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            m.append(e)
        tcPr.append(m)

    def borders(table, color="999999", sz=4):
        tbl = table._tbl
        tblPr = tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{k}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
            b.append(e)
        tblPr.append(b)

    def write(cell, text, size=8, bold=False, italic=False, color=None, align="l"):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                       "r": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        text = str(text).replace("<br/>", "\n")
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold, run.italic = bold, italic
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins(cell)

    def set_widths(table, widths_mm):
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for col, w in zip(table.columns, widths_mm):
            col.width = Mm(w)
        for row in table.rows:
            for c, w in zip(row.cells, widths_mm):
                c.width = Mm(w)

    def repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement("w:tblHeader")
        e.set(qn("w:val"), "true")
        trPr.append(e)

    def no_split(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))

    def bar(text):
        t = doc.add_table(rows=1, cols=1)
        set_widths(t, [180])
        c = t.rows[0].cells[0]
        shade(c, NAVY)
        write(c, text, size=10, bold=True, color="FFFFFF")
        keep = c.paragraphs[0]
        keep.paragraph_format.keep_with_next = True
        return t

    def spacer(pts=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        for r in p.runs:
            r.font.size = Pt(1)
        p.paragraph_format.line_spacing = Pt(pts)

    def field(par, code):
        run = par.add_run()
        for typ, txt in (("begin", None), (None, code), ("end", None)):
            if typ:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), typ)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
            run._r.append(e)
        run.font.size = Pt(8)

    # en-tête / pied de page
    hp = sec.header.paragraphs[0]
    hp.style = doc.styles["Normal"]
    hp.text = f"{P['projet']} — {P['element']} ({P['niveau']})\tNote de calcul — Voile RPA 2024"
    hp.paragraph_format.tab_stops.add_tab_stop(Mm(180), alignment=2)
    for r in hp.runs:
        r.font.size = Pt(8)
    fp = sec.footer.paragraphs[0]
    fp.style = doc.styles["Normal"]
    fp.paragraph_format.tab_stops.add_tab_stop(Mm(180), alignment=2)
    r = fp.add_run(f"Établie le {P['date']}" + (f" par {P['auteur']}" if P["auteur"] else "") + "\tPage ")
    r.font.size = Pt(8)
    field(fp, "PAGE")
    r2 = fp.add_run(" / ")
    r2.font.size = Pt(8)
    field(fp, "NUMPAGES")

    for b in blocks:
        t = b["t"]
        if t == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            r = p.add_run(b["text"])
            r.bold = True
            r.font.size = Pt(20)
            r.font.color.rgb = RGBColor.from_string(NAVY)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b["sub"].replace("<br/>", "\n"))
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            spacer(8)
        elif t == "info":
            tb = doc.add_table(rows=len(b["rows"]), cols=4)
            borders(tb)
            for i, row in enumerate(b["rows"]):
                for j, x in enumerate(row):
                    write(tb.rows[i].cells[j], x, size=8.5, bold=(j % 2 == 0))
                    if j % 2 == 0:
                        shade(tb.rows[i].cells[j], "E8EEF5")
            set_widths(tb, [24, 66, 24, 66])
            spacer(8)
        elif t == "banner":
            tb = doc.add_table(rows=1, cols=1)
            set_widths(tb, [180])
            c = tb.rows[0].cells[0]
            shade(c, "1B7A3D" if b["ok"] else "B3261E")
            write(c, b["text"], size=9.5, bold=True, color="FFFFFF")
            spacer(6)
        elif t == "list":
            for k in b["items"]:
                p = doc.add_paragraph()
                r = p.add_run("✗ ")
                r.font.color.rgb = RGBColor(0xB3, 0x26, 0x1E)
                p.add_run(k).font.size = Pt(8.5)
            spacer(6)
        elif t == "bar":
            bar(b["text"])
            spacer(6)
        elif t == "para":
            p = doc.add_paragraph()
            r = p.add_run(b["text"].replace("<br/>", "\n"))
            r.italic = True
            r.font.size = Pt(8)
            spacer(6)
        elif t == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "image":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(io.BytesIO(b["png"]), width=Mm(b["w"]))
        elif t == "table":
            if b.get("title"):
                bar(b["title"])
            if b.get("intro"):
                p = doc.add_paragraph()
                r = p.add_run(b["intro"])
                r.italic = True
                r.font.size = Pt(8)
            fs = b.get("fs", 7.6) + 0.4
            al = b["align"]
            ncol = len(b["head"])
            tb = doc.add_table(rows=1 + len(b["rows"]), cols=ncol)
            borders(tb)
            for j, h in enumerate(b["head"]):
                c = tb.rows[0].cells[j]
                write(c, h, size=fs, bold=True, align="c" if al[j] != "l" else "l")
                shade(c, "E8EEF5")
            repeat_header(tb.rows[0])
            for i, row in enumerate(b["rows"], start=1):
                no_split(tb.rows[i])
                for j, c in enumerate(row):
                    txt, ok = _cell_split(c)
                    cell = tb.rows[i].cells[j]
                    if b.get("italic_col") == j:
                        write(cell, txt, size=fs - 0.6, italic=True, color="444444")
                    elif ok is None:
                        write(cell, txt, size=fs, align=al[j])
                    else:
                        write(cell, txt, size=fs, bold=True, color="1B7A3D" if ok else "B3261E", align=al[j])
                        shade(cell, "E6F4EA" if ok else "FDE7E9")
                if b.get("span_eq") and len(row) == 3 and row[2] == "=":
                    a = tb.rows[i].cells[1].merge(tb.rows[i].cells[2])
                    for extra in a.paragraphs[1:]:
                        extra._element.getparent().remove(extra._element)
            set_widths(tb, b["widths"])
            spacer(8)
    doc.core_properties.title = f"Note de calcul — {P['element']}"
    doc.core_properties.author = P["auteur"] or ""
    doc.save(path)
    return path


def build_pdf(P, res, path, multi=None):
    return render_pdf(report_blocks(P, res, multi), P, path)


def build_docx(P, res, path, multi=None):
    return render_docx(report_blocks(P, res, multi), P, path)


# ---------------------------------------------------------------------------
#  6. INTERFACE TKINTER
# ---------------------------------------------------------------------------
GENERAL_FIELDS = [
    ("he", "he — Hauteur d'étage (libre)", "m"), ("lw", "lw — Longueur du voile", "m"),
    ("bw", "bw — Épaisseur du voile", "m"), ("d", "d — Enrobage des armatures", "cm"),
    ("fc", "fc28 — Résistance caractéristique du béton", "MPa"), ("fe", "fe — Limite élastique de l'acier", "MPa"),
    ("gb", "γb — Coefficient de sécurité du béton", "-"), ("gs", "γs — Coefficient de sécurité de l'acier", "-"),
    ("hcr", "hcr — Hauteur de la zone critique", "m"), ("R", "R — Coefficient de comportement", "-"),
    ("Qf", "Qf — Facteur de qualité", "-"), ("T0", "T0 — Période fondamentale / caractéristique", "s"),
    ("T2", "T2 — Période caractéristique du site", "s"), ("N", "N — Effort normal sollicitant", "kN"),
    ("MED", "MED — Moment fléchissant sollicitant", "kN.m"), ("MRD", "MRD — Moment résistant", "kN.m"),
    ("V", "V — Effort tranchant sismique (ETABS)", "kN"),
]
# bouton « valeur du 1er niveau » placé à côté de la valeur de calcul qu'il renseigne
FIRST_BTN = {"MED": ("← |M3| du 1er niveau (ETABS)", "a"), "N": ("← |P min| du 1er niveau (ETABS)", "b"),
             "V": ("← |V2| du 1er niveau (ETABS)", "c")}
SIDE_FIELDS = [
    ("phi_l", "Øl — Diamètre des barres verticales de rive", "mm"),
    ("lc", "lc,choisi — Longueur de l'élément de rive", "cm"),
    ("bc", "bc,choisi — Largeur de l'élément de rive", "cm"),
    ("Avcalc", "Avcalc — Section calculée (RDM, flexion composée)", "cm²"),
    ("phi_c", "Øc — Diamètre des cadres de confinement", "mm"),
    ("nbr", "Nbr — Nombre de branches de confinement", "-"),
    ("n_conf", "n — Nombre de barres transversales confinées", "-"),
    ("bi", "bi — Entraxe des barres confinées", "cm"),
    ("L_conf", "L — Longueur cumulée des branches", "cm"),
    ("nbc", "nbc — Barres par petite face (schéma, ≥ 2)", "-"),
    ("nl", "nl — Barres par grande face (0 = auto)", "-"),
]
WEB_FIELDS = [
    ("phi_v", "Øv — Diamètre des barres verticales de l'âme", "mm"),
    ("st_v", "St,v — Espacement des barres verticales", "cm"),
    ("phi_h", "Øh — Diamètre des barres horizontales de l'âme", "mm"),
    ("st_h", "St,h — Espacement des barres horizontales", "cm"),
]
STORY_COLS = [("name", "Niveau", 12), ("h", "h (m)", 7), ("L", "L (m)", 7), ("e", "e (m)", 7),
              ("N", "N [P] (kN)", 10), ("V2", "V2 (kN)", 10), ("M3", "M3 ETABS (kN.m)", 13), ("Pmin", "P min (kN)", 10)]
REF_LABELS = {"base": "base de l'étage (sécuritaire)", "sommet": "sommet de l'étage (comme le classeur)"}
V_LABELS = {"kV": "kV × V2 ETABS (comportement d'origine)", "enveloppe": "V corrigé de l'enveloppe (RPA 2024 Art. 7.7.2)"}
DIAMETRES = ["8", "10", "12", "14", "16", "20", "25", "32", "40"]
MODE_LABELS = {"auto": "Automatique selon l'élancement (Art. 7.7.4)", "enveloppe": "Enveloppe (classeur)", "max": "max(enveloppe ; M3 ETABS)", "etabs": "M3 ETABS brut"}


def stories_to_floats(raw):
    """Convertit la liste d'étages saisie (chaînes) en nombres ; lève ValueError si invalide."""
    out = []
    for i, s in enumerate(raw, start=1):
        d = {"name": str(s["name"]).strip() or f"Niv {i}"}
        for k in ("h", "L", "e", "N", "V2", "M3", "Pmin"):
            v = _num(s.get(k, 0))
            if v is None:
                raise ValueError(f"Étage {d['name']} : valeur invalide pour « {k} »")
            d[k] = v
        d["sec"] = str(s.get("sec", "") or "").strip()          # section adoptée (facultative)
        try:
            parse_section(d["sec"])
        except ValueError as e:
            raise ValueError(f"Niveau {d['name']} : {e}")
        out.append(d)
    return out


def run_gui():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

    class App(tk.Tk):
        def __init__(self):
            super().__init__()
            self.title("Voile RPA 2024 — Enveloppes M / V, contraintes, moment résistant, ferraillage et ductilité locale")
            self.geometry("1280x830")
            self.minsize(1040, 680)
            self.vars = {k: tk.StringVar(value=str(v)) for k, v in DEFAULTS.items()}
            self.srows = []            # lignes de la grille d'étages (dict de StringVar)
            self.etabs = None
            self.res = self.multi = self.mr = None
            self.mr_err = ""
            self.path = None                  # fichier de projet courant (« Enregistrer » y réécrit)
            self._job = None
            self._build()
            self.after(100, self.recalc)

        # ------------------------------------------------------------ construction
        def _build(self):
            style = ttk.Style(self)
            try:
                style.theme_use("clam")
            except tk.TclError:
                pass
            style.configure("Treeview", rowheight=22)
            style.configure("Banner.TLabel", font=("Segoe UI", 10, "bold"), padding=6)
            bar = ttk.Frame(self, padding=(8, 6))
            bar.pack(fill="x")
            for txt, cmd in (("Nouveau", self.new), ("Ouvrir…", self.open), ("Enregistrer  (Ctrl+S)", self.save),
                             ("Enregistrer sous…", self.save_as)):
                ttk.Button(bar, text=txt, command=cmd).pack(side="left", padx=2)
            ttk.Separator(bar, orient="vertical").pack(side="left", fill="y", padx=8)
            ttk.Button(bar, text="↻ Actualiser le calcul  (F5)", command=self.refresh).pack(side="left", padx=2)
            ttk.Separator(bar, orient="vertical").pack(side="left", fill="y", padx=8)
            self.bind("<Control-s>", lambda e: self.save())
            self.bind("<Control-S>", lambda e: self.save_as())
            self.bind("<F5>", lambda e: self.refresh())
            ttk.Button(bar, text="Note de calcul A4 (PDF)…", command=lambda: self.make_report("pdf")).pack(side="left", padx=2)
            ttk.Button(bar, text="Note de calcul Word (.docx)…", command=lambda: self.make_report("docx")).pack(side="left", padx=2)
            ttk.Button(bar, text="Contact", command=self.show_contact).pack(side="right", padx=2)
            ttk.Button(bar, text="Aide", command=self.open_help).pack(side="right", padx=2)
            self.banner = ttk.Label(self, text="", style="Banner.TLabel", anchor="center")
            self.banner.pack(fill="x", padx=8)
            self.nb = ttk.Notebook(self)
            self.nb.pack(fill="both", expand=True, padx=8, pady=6)
            self.gen_ent = {}
            self._tab_seisme()
            self._tab_mat()
            self._tab_general()
            self._tab_stories()
            self._tab_env()
            self._tab_envv()
            self._tab_stress()
            self._tab_mr()
            self._tab_elev()
            self._tab_dims()
            self._tab_rive()
            self._tab_results()
            self._tab_schema()
            self.status = ttk.Label(self, text="Prêt.", anchor="w", padding=(8, 2))
            self.status.pack(fill="x")

        def _entry(self, parent, key, row, col, width=12, left=False):
            e = ttk.Entry(parent, textvariable=self.vars[key], width=width, justify="left" if left else "right")
            e.grid(row=row, column=col, padx=4, pady=2, sticky="e")
            self.vars[key].trace_add("write", lambda *_: self.schedule())
            return e

        def _combo(self, parent, key, values, row, col, width=24):
            c = ttk.Combobox(parent, textvariable=self.vars[key], values=values, width=width, state="readonly")
            c.grid(row=row, column=col, padx=4, pady=2, sticky="w")
            c.bind("<<ComboboxSelected>>", lambda *_: self.schedule())
            return c

        # ---- données sismiques (RPA 2024)
        def _tab_seisme(self):
            f = ttk.Frame(self.nb, padding=10)
            self.nb.add(f, text="  Données sismiques  ")
            self.sz_pq_lbl = []

            def row(box, i, label, key, unit="", combo=None, width=11):
                ttk.Label(box, text=label).grid(row=i, column=0, sticky="w", pady=2)
                if combo:
                    self._combo(box, key, combo, i, 1, width=width)
                else:
                    self._entry(box, key, i, 1, width=width)
                ttk.Label(box, text=unit, foreground="gray").grid(row=i, column=2, sticky="w", padx=4)

            a = ttk.LabelFrame(f, text="A. Zone, importance et site (Tables 3.3 et 3.10)", padding=8)
            a.grid(row=0, column=0, sticky="nsew", padx=5, pady=4)
            row(a, 0, "Zone sismique", "sz_zone", "", list(RPA_ZONES), 30)
            row(a, 1, "Groupe d'importance", "sz_groupe", "", list(RPA_GROUPES), 30)
            row(a, 2, "Classe de site", "sz_site", "", list(RPA_SITES), 30)
            b = ttk.LabelFrame(f, text="B. Système de contreventement (Table 3.18)", padding=8)
            b.grid(row=1, column=0, sticky="nsew", padx=5, pady=4)
            row(b, 0, "Système (n° 1 à 9)", "sz_sys", "", SYS_LABELS, 60)
            c = ttk.LabelFrame(f, text="C. Facteur de qualité QF = 1 + ΣPq (Table 3.19) — critère observé ?", padding=8)
            c.grid(row=0, column=1, rowspan=2, sticky="nsew", padx=5, pady=4)
            for i, (cat, lab, v) in enumerate(RPA_PQ, start=1):
                lb = ttk.Label(c, text=f"({cat}) {lab}   [Pq = {fr(v, 2)}]")
                lb.grid(row=i - 1, column=0, sticky="w", pady=2)
                self.sz_pq_lbl.append(lb)
                self._combo(c, f"sz_pq{i}", ["Oui", "Non"], i - 1, 1, width=6)
            ttk.Label(c, text="Seuls les critères de la catégorie du système retenu s'appliquent.",
                      foreground="gray").grid(row=8, column=0, columnspan=2, sticky="w", pady=(4, 0))
            d = ttk.LabelFrame(f, text="D. Période fondamentale  T = CT · hN^(3/4)", padding=8)
            d.grid(row=2, column=0, sticky="nsew", padx=5, pady=4)
            row(d, 0, "CT", "sz_ct", "-")
            row(d, 1, "hN — hauteur du bâtiment", "sz_hn", "m")
            ttk.Button(d, text="← hw (somme des hauteurs d'étage)", command=self.sz_hn_from_stories).grid(
                row=1, column=3, padx=6)
            e = ttk.LabelFrame(f, text="E. Amortissement critique ξ (Table 3.5)", padding=8)
            e.grid(row=2, column=1, sticky="nsew", padx=5, pady=4)
            row(e, 0, "Type de structure", "sz_struct", "", list(RPA_AMORT), 30)
            row(e, 1, "Remplissage", "sz_remp", "", ["Léger", "Dense"], 12)
            g = ttk.LabelFrame(f, text="Récapitulatif et liaison avec les Données générales", padding=8)
            g.grid(row=3, column=0, columnspan=2, sticky="nsew", padx=5, pady=4)
            ttk.Checkbutton(g, text="Lier à Données générales : R → R ; QF → Qf ; T2 → T2 ; T → T0",
                            variable=self.vars["seis_link"], onvalue="1", offvalue="0",
                            command=self.schedule).pack(anchor="w")
            self.sz_lbl = ttk.Label(g, text="", font=("Consolas", 10), justify="left")
            self.sz_lbl.pack(anchor="w", pady=4)
            ttk.Label(f, text="La régularité en plan s'indique ici par Oui / Non (la méthode géométrique du classeur n'est "
                              "pas reprise). T0 = T = CT·hN^(3/4) ; T2 dépend de la classe de site.",
                      foreground="gray", wraplength=1000).grid(row=4, column=0, columnspan=2, sticky="w", padx=5)

        # ---- matériaux
        def _tab_mat(self):
            f = ttk.Frame(self.nb, padding=10)
            self.nb.add(f, text="  Matériaux  ")

            def row(box, i, label, key, unit="", combo=None, width=11):
                ttk.Label(box, text=label).grid(row=i, column=0, sticky="w", pady=2)
                if combo:
                    self._combo(box, key, combo, i, 1, width=width)
                else:
                    self._entry(box, key, i, 1, width=width)
                ttk.Label(box, text=unit, foreground="gray").grid(row=i, column=2, sticky="w", padx=4)

            a = ttk.LabelFrame(f, text="Béton", padding=8)
            a.grid(row=0, column=0, sticky="nsew", padx=5, pady=4)
            row(a, 0, "fc28 — Fondation", "mt_fc_fond", "MPa")
            row(a, 1, "fc28 — Poteaux, voiles, poutres, dalles", "mt_fc", "MPa")
            row(a, 2, "Densité du béton", "mt_dens", "T/m³")
            b = ttk.LabelFrame(f, text="Aciers", padding=8)
            b.grid(row=0, column=1, sticky="nsew", padx=5, pady=4)
            row(b, 0, "fe — Barres HA (acier de calcul)", "mt_fe_ha", "MPa")
            row(b, 1, "fe — Treillis soudés HA", "mt_fe_tr", "MPa")
            row(b, 2, "fe — Aciers doux", "mt_fe_doux", "MPa")
            c = ttk.LabelFrame(f, text="Coefficients de sécurité (situation de calcul)", padding=8)
            c.grid(row=1, column=0, columnspan=2, sticky="nsew", padx=5, pady=4)
            row(c, 0, "Situation", "mt_sit", "", list(MAT_SITUATIONS), 26)
            row(c, 1, "γb manuel (situation « Manuel »)", "mt_gb_m", "-")
            row(c, 2, "γs manuel (situation « Manuel »)", "mt_gs_m", "-")
            g = ttk.LabelFrame(f, text="Récapitulatif et liaison avec les Données générales", padding=8)
            g.grid(row=2, column=0, columnspan=2, sticky="nsew", padx=5, pady=4)
            ttk.Checkbutton(g, text="Lier à Données générales : fc28 → fc ; fe → fe ; γb, γs  "
                                    "(et fc28, fe de « Moment résistant » et des cas « Durable » / « Accidentelle »)",
                            variable=self.vars["mat_link"], onvalue="1", offvalue="0",
                            command=self.schedule).pack(anchor="w")
            self.mt_lbl = ttk.Label(g, text="", font=("Consolas", 10), justify="left")
            self.mt_lbl.pack(anchor="w", pady=4)
            ttk.Label(f, text="Ev = 3700·fc28^(1/3) ; Ei = 11000·fc28^(1/3). Durable : γb = 1,5 ; γs = 1,15 — "
                              "Accidentelle (séisme) : γb = 1,2 ; γs = 1,0 (valeurs de l'application).",
                      foreground="gray", wraplength=1000).grid(row=3, column=0, columnspan=2, sticky="w", padx=5)

        def sz_hn_from_stories(self):
            hs = [_num(r["h"].get()) for r in self.srows]
            if not hs or None in hs:
                messagebox.showinfo("hN", "Aucun niveau valide dans « Étages / ETABS ».")
                return
            self.vars["sz_hn"].set(str(round(sum(hs), 3)))
            self.status.config(text=f"hN = Σ h = {fr(sum(hs))} m")

        def _apply_links(self):
            """Impose aux Données générales les valeurs des rubriques sismique / matériaux (si liaison active)."""
            try:
                P = self.read()
                lv = link_values(P)
            except ValueError as e:
                self.status.config(text=f"⚠ Hypothèses : {e}")
                return
            for k, val in lv.items():
                cur = _num(self.vars[k].get())
                if cur is None or abs(cur - val) > 1e-9:
                    self.vars[k].set(str(val))
            for keys, flag in ((("R", "Qf", "T2", "T0"), P["seis_link"]), (("fc", "fe", "gb", "gs"), P["mat_link"])):
                for k in keys:
                    if k in self.gen_ent:
                        self.gen_ent[k].configure(state="readonly" if int(flag) else "normal")
            self.fill_hyp(P)

        def fill_hyp(self, P):
            try:
                sp = seismic_params(P)
                self.sz_lbl.config(text=(
                    f"A = {fr(sp['A'], 2)}   I = {fr(sp['I'], 1)}   S = {fr(sp['S'], 2)}   T1 = {fr(sp['T1'], 2)} s   "
                    f"T2 = {fr(sp['T2'], 2)} s   T3 = {fr(sp['T3'], 2)} s\n"
                    f"R = {fr(sp['R'], 2)} (système {sp['n_sys']}, catégorie {sp['cat']})   QF = {fr(sp['QF'], 2)}   "
                    f"ξ = {fr(sp['xi'], 0)} %   η = {fr(sp['eta'], 3)}\n"
                    f"T = {fr(P['sz_ct'], 2)} × {fr(P['sz_hn'])}^(3/4) = {fr(sp['T'], 3)} s   "
                    f"0,2·A·I = {fr(sp['Sa_min'], 3)} g"), foreground="black")
                for lb, q in zip(self.sz_pq_lbl, sp["pq"]):
                    lb.config(foreground="black" if q["actif"] else "gray")
            except ValueError as e:
                self.sz_lbl.config(text=f"⚠ {e}", foreground="#b3261e")
            try:
                mp = material_params(P)
                self.mt_lbl.config(text=(
                    f"fc28 voiles = {fr(mp['fc'], 0)} MPa   Ev = {fr(mp['Ev'], 0)} MPa   Ei = {fr(mp['Ei'], 0)} MPa   "
                    f"(fondation : Ev = {fr(mp['Ev_fond'], 0)} ; Ei = {fr(mp['Ei_fond'], 0)})\n"
                    f"fe = {fr(mp['fe'], 0)} MPa   γb = {fr(mp['gb'], 2)}   γs = {fr(mp['gs'], 2)}   "
                    f"fbu = {fr(mp['fbu'], 2)} MPa   σs = fe/γs = {fr(mp['sigs'], 1)} MPa"), foreground="black")
            except ValueError as e:
                self.mt_lbl.config(text=f"⚠ {e}", foreground="#b3261e")

        # ---- données générales
        def _tab_general(self):
            f = ttk.Frame(self.nb, padding=12)
            self.nb.add(f, text="  Données générales  ")
            box = ttk.LabelFrame(f, text="Projet", padding=8)
            box.grid(row=0, column=0, sticky="nsew", padx=6, pady=4)
            for i, (k, lab) in enumerate((("projet", "Projet"), ("element", "Élément"), ("niveau", "Niveau"),
                                          ("auteur", "Établi par"), ("date", "Date"))):
                ttk.Label(box, text=lab).grid(row=i, column=0, sticky="w", pady=2)
                self._entry(box, k, i, 1, width=34, left=True)
            box2 = ttk.LabelFrame(f, text="Données d'entrée du voile (ductilité locale — ETABS / RDM)", padding=8)
            box2.grid(row=0, column=1, sticky="nsew", padx=6, pady=4)
            for i, (k, lab, u) in enumerate(GENERAL_FIELDS):
                ttk.Label(box2, text=lab).grid(row=i, column=0, sticky="w", pady=2)
                self.gen_ent[k] = self._entry(box2, k, i, 1)
                ttk.Label(box2, text=u, foreground="gray").grid(row=i, column=2, sticky="w")
                if k in FIRST_BTN:
                    ttk.Button(box2, text=FIRST_BTN[k][0], width=34,
                               command=lambda w=FIRST_BTN[k][1]: self.take_first(w)).grid(row=i, column=3, padx=6)
                elif k == "MRD":
                    ttk.Button(box2, text="← Mru (Moment résistant)", width=34,
                               command=self.mr_to_mrd).grid(row=i, column=3, padx=6)
            ttk.Label(f, text="R, Qf, T2, T0 (rubrique « Données sismiques ») et fc28, fe, γb, γs (rubrique « Matériaux ») "
                              "sont reliés : champs grisés = valeurs imposées ; décocher la liaison pour les saisir librement.",
                      foreground="#1f3a5f", wraplength=900).grid(row=2, column=0, columnspan=2, sticky="w", padx=6, pady=4)
            ttk.Label(f, text="Ṽ = 1,4 × V est calculé automatiquement. N, MED, V peuvent être repris d'un étage "
                              "(onglet « Contraintes »).", foreground="gray").grid(row=1, column=0, columnspan=2,
                                                                                    sticky="w", padx=6)

        # ---- étages / ETABS
        def _tab_stories(self):
            f = ttk.Frame(self.nb, padding=8)
            self.nb.add(f, text="  Étages / ETABS  ")
            top = ttk.Frame(f)
            top.pack(fill="x")
            ttk.Button(top, text="Importer un export ETABS (xlsx / xlsm / csv)…", command=self.import_etabs).pack(side="left")
            ttk.Label(top, text="  Pier :").pack(side="left")
            self.pier_cb = ttk.Combobox(top, textvariable=self.vars["pier"], values=[], width=10, state="readonly")
            self.pier_cb.pack(side="left", padx=4)
            ttk.Button(top, text="Charger les efforts du pier", command=self.load_pier).pack(side="left", padx=4)
            self.etabs_lbl = ttk.Label(top, text="Aucun fichier importé — saisie manuelle possible.", foreground="gray")
            self.etabs_lbl.pack(side="left", padx=10)
            top2 = ttk.Frame(f)
            top2.pack(fill="x", pady=(6, 2))
            ttk.Button(top2, text="Ajouter un niveau (en haut)", command=self.add_story).pack(side="left", padx=2)
            ttk.Button(top2, text="Supprimer le dernier", command=self.del_story).pack(side="left", padx=2)
            ttk.Button(top2, text="Appliquer lw et bw (Données générales) à tous les niveaux",
                       command=self.apply_geom).pack(side="left", padx=2)
            box1 = ttk.LabelFrame(f, text="Valeurs du 1er niveau (1re ligne du tableau)", padding=6)
            box1.pack(fill="x", pady=(2, 4))
            ttk.Button(box1, text="Tout appliquer : MED, N, V (Données générales), Nu (Moment résistant), bi (Rives et âme)",
                       command=lambda: self.take_first("all")).pack(side="left", padx=3)
            ttk.Label(box1, text="Chaque valeur a aussi son bouton « ← 1er niveau » à côté du champ concerné.",
                      foreground="gray").pack(side="left", padx=10)
            ttk.Label(f, text="Efforts de la section du bas de chaque niveau — N < 0 = compression. N [P] = P "
                              "correspondant au M3 maximal. Niveaux listés du bas (haut du tableau) vers le haut. "
                              "Tout est modifiable.", foreground="gray", wraplength=1150).pack(anchor="w")
            cont = ttk.Frame(f)
            cont.pack(fill="both", expand=True, pady=4)
            cv = tk.Canvas(cont, highlightthickness=0)
            sb = ttk.Scrollbar(cont, orient="vertical", command=cv.yview)
            self.grid_frame = ttk.Frame(cv)
            self.grid_frame.bind("<Configure>", lambda e: cv.configure(scrollregion=cv.bbox("all")))
            cv.create_window((0, 0), window=self.grid_frame, anchor="nw")
            cv.configure(yscrollcommand=sb.set)
            cv.pack(side="left", fill="both", expand=True)
            sb.pack(side="right", fill="y")
            self._grid_header()

        def _grid_header(self):
            for j, (_, lab, w) in enumerate(STORY_COLS):
                ttk.Label(self.grid_frame, text=lab, font=("Segoe UI", 9, "bold")).grid(row=0, column=j, padx=3, pady=2)

        def set_stories(self, stories):
            for w in self.grid_frame.winfo_children():
                w.destroy()
            self._grid_header()
            self.srows = []
            for i, s in enumerate(stories, start=1):
                row = {}
                for j, (k, _, w) in enumerate(STORY_COLS):
                    var = tk.StringVar(value=str(s.get(k, 0 if k != "name" else f"Niv {i}")))
                    var.trace_add("write", lambda *_: self.schedule())
                    ttk.Entry(self.grid_frame, textvariable=var, width=w,
                              justify="left" if k == "name" else "right").grid(row=i, column=j, padx=3, pady=1)
                    row[k] = var
                sv = tk.StringVar(value=str(s.get("sec", "") or ""))     # section adoptée (saisie dans « Contraintes »)
                sv.trace_add("write", lambda *_: self.schedule())
                row["sec"] = sv
                self.srows.append(row)
            self.schedule()

        def raw_stories(self):
            return [{k: v.get() for k, v in r.items()} for r in self.srows]

        def add_story(self):
            raw = self.raw_stories()
            last = raw[-1] if raw else {"name": "Niv 0", "h": self.vars["he"].get(), "L": self.vars["lw"].get(),
                                       "e": self.vars["bw"].get(), "N": 0, "V2": 0, "M3": 0, "Pmin": 0}
            new = dict(last)
            new["name"] = f"Niv {len(raw) + 1}"
            new["sec"] = ""
            if len(raw) >= 40:
                return
            self.set_stories(raw + [new])

        def del_story(self):
            raw = self.raw_stories()
            if raw:
                self.set_stories(raw[:-1])

        def apply_geom(self):
            raw = self.raw_stories()
            for r in raw:
                r["L"], r["e"] = self.vars["lw"].get(), self.vars["bw"].get()
            self.set_stories(raw)

        def import_etabs(self):
            path = filedialog.askopenfilename(
                filetypes=[("Export ETABS", "*.xlsx *.xlsm *.csv *.txt"), ("Tous", "*.*")])
            if not path:
                return
            try:
                self.etabs = load_etabs(path)
            except Exception as e:  # noqa
                messagebox.showerror("Import ETABS", str(e))
                return
            self.vars["etabs_path"].set(path)
            self.pier_cb.configure(values=self.etabs["piers"])
            if self.vars["pier"].get() not in self.etabs["piers"]:
                self.vars["pier"].set(self.etabs["piers"][0])
            self.etabs_lbl.config(text=f"{os.path.basename(path)} — feuille « {self.etabs['sheet']} » — "
                                       f"{len(self.etabs['piers'])} piers, {len(self.etabs['stories'])} niveaux "
                                       f"({'brut' if self.etabs['kind'] == 'raw' else 'synthèse'})")
            self.load_pier()

        def load_pier(self):
            if not self.etabs:
                messagebox.showinfo("ETABS", "Importer d'abord un fichier ETABS.")
                return
            try:
                P = self.read()
                self.set_stories(stories_from_etabs(self.etabs, self.vars["pier"].get(), P))
            except Exception as e:  # noqa
                messagebox.showerror("ETABS", str(e))
                return
            self.status.config(text=f"Efforts du pier {self.vars['pier'].get()} chargés "
                                    f"({len(self.srows)} niveaux). Ajuster L, e et h si nécessaire.")

        # ---- enveloppe
        def _tab_env(self):
            f = ttk.Frame(self.nb, padding=10)
            self.nb.add(f, text="  Enveloppe M  ")
            box = ttk.LabelFrame(f, text="Paramètres de l'enveloppe (RPA 2024 — éq. 7.22 et 7.23)", padding=8)
            box.pack(fill="x")
            auto = ttk.Checkbutton(box, text="M base et M sommet issus des efforts ETABS (niveaux bas / haut)",
                                   variable=self.vars["env_auto"], onvalue="1", offvalue="0",
                                   command=self.schedule)
            auto.grid(row=0, column=0, columnspan=3, sticky="w")
            ttk.Label(box, text="M base (valeur absolue)").grid(row=1, column=0, sticky="w")
            self._entry(box, "env_mbase", 1, 1)
            ttk.Label(box, text="kN.m — utilisé si la case est décochée", foreground="gray").grid(row=1, column=2, sticky="w")
            ttk.Label(box, text="M sommet (valeur absolue)").grid(row=2, column=0, sticky="w")
            self._entry(box, "env_msom", 2, 1)
            ttk.Label(box, text="Hauteur de référence du niveau").grid(row=3, column=0, sticky="w", pady=(6, 0))
            self.ref_cb = ttk.Combobox(box, values=list(REF_LABELS.values()), width=36, state="readonly")
            self.ref_cb.grid(row=3, column=1, columnspan=2, sticky="w", pady=(6, 0))
            self.ref_cb.bind("<<ComboboxSelected>>", lambda *_: self._set_ref())
            ttk.Label(box, text="Moment de calcul retenu").grid(row=4, column=0, sticky="w")
            self.mode_cb = ttk.Combobox(box, values=list(MODE_LABELS.values()), width=36, state="readonly")
            self.mode_cb.grid(row=4, column=1, columnspan=2, sticky="w")
            self.mode_cb.bind("<<ComboboxSelected>>", lambda *_: self._set_mode())
            ttk.Label(box, text="« Base » applique l'enveloppe à la section du bas de chaque étage (cohérent avec les "
                                "efforts ETABS « Bottom »). « Sommet » reproduit à l'identique la colonne « M corrigé » "
                                "du classeur d'origine.", foreground="gray", wraplength=1100).grid(
                row=5, column=0, columnspan=3, sticky="w", pady=(4, 0))
            self.env_lbl = ttk.Label(f, text="", font=("Consolas", 10), justify="left")
            self.env_lbl.pack(anchor="w", pady=6)
            ttk.Button(f, text="Envoyer hcr → Données générales (zone critique du voile)", command=self.send_hcr).pack(anchor="w")
            cols = ("h", "H", "ml", "mc", "z")
            self.env_tree = ttk.Treeview(f, columns=cols, show="tree headings", height=14)
            for c, t, w in (("h", "h (m)", 70), ("H", "H cumulée (m)", 110), ("ml", "M linéaire (kN.m)", 130),
                            ("mc", "M corrigé (kN.m)", 130), ("z", "Zone", 220)):
                self.env_tree.heading(c, text=t)
                self.env_tree.column(c, width=w, anchor="center")
            self.env_tree.heading("#0", text="Niveau")
            self.env_tree.column("#0", width=160)
            self.env_tree.pack(fill="both", expand=True, pady=6)
            self._sync_combos()

        def _sync_combos(self):
            self.ref_cb.set(REF_LABELS.get(self.vars["env_ref"].get(), REF_LABELS["base"]))
            self.mode_cb.set(MODE_LABELS.get(self.vars["m_mode"].get(), MODE_LABELS["auto"]))
            if hasattr(self, "v_cb"):
                self.v_cb.set(V_LABELS.get(self.vars["v_mode"].get(), V_LABELS["kV"]))

        def _set_ref(self):
            for k, v in REF_LABELS.items():
                if v == self.ref_cb.get():
                    self.vars["env_ref"].set(k)
            self.schedule()

        def _set_mode(self):
            for k, v in MODE_LABELS.items():
                if v == self.mode_cb.get():
                    self.vars["m_mode"].set(k)
            self.schedule()

        # ---- enveloppe de l'effort tranchant
        def _tab_envv(self):
            f = ttk.Frame(self.nb, padding=10)
            self.nb.add(f, text="  Enveloppe V  ")
            box = ttk.LabelFrame(f, text="Enveloppe de l'effort tranchant (RPA 2024 — Art. 7.7.2)", padding=8)
            box.pack(fill="x")
            ttk.Label(box, text="V de calcul pour τb (méthode des contraintes)").grid(row=0, column=0, sticky="w")
            self.v_cb = ttk.Combobox(box, values=list(V_LABELS.values()), width=46, state="readonly")
            self.v_cb.grid(row=0, column=1, sticky="w", padx=6)
            self.v_cb.bind("<<ComboboxSelected>>", lambda *_: self._set_v())
            ttk.Label(box, text="Zone critique : V × 1,4 (h ≤ hcr) ; transition linéaire de (hcr ; 1,4·V(hcr)) à "
                                "(hw/3 ; V(hw/3)) ; au-delà de hw/3 : V de l'analyse. « kV × V2 » conserve le calcul "
                                "d'origine (kV dans l'onglet Contraintes) ; « V corrigé » utilise cette enveloppe "
                                "(kV = 1, le 1,4 étant déjà inclus en zone critique).", foreground="gray",
                      wraplength=1100).grid(row=1, column=0, columnspan=2, sticky="w", pady=(4, 0))
            self.v_lbl = ttk.Label(f, text="", font=("Consolas", 10), justify="left")
            self.v_lbl.pack(anchor="w", pady=6)
            cols = ("H", "Ti", "T14", "Vc", "z")
            self.v_tree = ttk.Treeview(f, columns=cols, show="tree headings", height=14)
            for c, t, w in (("H", "H cumulée (m)", 110), ("Ti", "V analyse (kN)", 120), ("T14", "1,4 × V (kN)", 120),
                            ("Vc", "V corrigé (kN)", 120), ("z", "Zone", 220)):
                self.v_tree.heading(c, text=t)
                self.v_tree.column(c, width=w, anchor="center")
            self.v_tree.heading("#0", text="Niveau")
            self.v_tree.column("#0", width=160)
            self.v_tree.pack(fill="both", expand=True, pady=6)
            self._sync_combos()

        def _set_v(self):
            for k, v in V_LABELS.items():
                if v == self.v_cb.get():
                    self.vars["v_mode"].set(k)
            self.schedule()

        # ---- contraintes
        def _tab_stress(self):
            f = ttk.Frame(self.nb, padding=8)
            self.nb.add(f, text="  Contraintes (tous étages)  ")
            box = ttk.LabelFrame(f, text="Cas de calcul et paramètres de la méthode", padding=6)
            box.pack(fill="x")
            for j, h in enumerate(("Cas", "fc28", "fe", "γb", "γs", "θ")):
                ttk.Label(box, text=h, font=("Segoe UI", 9, "bold")).grid(row=0, column=j, padx=3)
            for i in (1, 2, 3):
                self._entry(box, f"c{i}_name", i, 0, width=22, left=True)
                for j, k in enumerate(("fc", "fe", "gb", "gs", "th"), start=1):
                    self._entry(box, f"c{i}_{k}", i, j, width=8)
            ttk.Label(box, text="Cas actif ►").grid(row=1, column=7, padx=(20, 2))
            self.cas_cb = ttk.Combobox(box, textvariable=self.vars["cas"], width=22, state="readonly")
            self.cas_cb.grid(row=1, column=8, padx=2)
            self.cas_cb.bind("<<ComboboxSelected>>", lambda *_: self.schedule())
            self.cas_cb.configure(values=[DEFAULTS[f"c{i}_name"] for i in (1, 2, 3)])
            for k in (1, 2, 3):
                self.vars[f"c{k}_name"].trace_add("write", lambda *_: self._refresh_cas())
            prm = ttk.Frame(box)
            prm.grid(row=4, column=0, columnspan=9, sticky="w", pady=(6, 0))
            col = 0
            for k, lab in (("kV", "kV"), ("kd", "kd"), ("ktau", "kτ"), ("rho_t", "ρmin tendue"), ("rho_g", "ρmin S.E.C")):
                ttk.Label(prm, text=lab).grid(row=0, column=col, padx=(8, 2))
                self._entry(prm, k, 0, col + 1, width=7)
                col += 2
            ttk.Label(prm, text="d = kd ×").grid(row=0, column=col, padx=(12, 2))
            self._combo(prm, "dref", ["h étage", "L voile"], 0, col + 1, width=9)
            act = ttk.Frame(f)
            act.pack(fill="x", pady=4)
            ttk.Label(act, text="Envoyer à « Données générales » (ductilité locale) le niveau :").pack(side="left")
            self.send_cb = ttk.Combobox(act, values=[], width=14, state="readonly")
            self.send_cb.pack(side="left", padx=4)
            ttk.Button(act, text="N = |P min|, MED = |M3|, V = |V2|  →", command=self.send_efforts).pack(side="left", padx=4)
            self.sum_lbl = ttk.Label(act, text="", foreground="#333333")
            self.sum_lbl.pack(side="left", padx=14)
            act2 = ttk.Frame(f)
            act2.pack(fill="x", pady=(0, 4))
            ttk.Label(act2, text="Section adoptée (facultatif) : double-clic dans la colonne « Section adoptée » puis "
                                 "saisir par ex. 12HA20 ou 6HA20+4HA16 — la section choisie est calculée.",
                      foreground="#1f4e79").pack(side="left")
            ttk.Label(act2, text="   Proposer avec Ø").pack(side="left")
            self.phi_cb = ttk.Combobox(act2, values=DIAMETRES, width=4, state="readonly")
            self.phi_cb.set("16")
            self.phi_cb.pack(side="left", padx=3)
            ttk.Button(act2, text="Proposer (cases vides)", command=self.propose_sections).pack(side="left", padx=3)
            ttk.Button(act2, text="Effacer les choix", command=self.clear_sections).pack(side="left", padx=3)
            cols = ("N", "V2", "M", "Le", "s1", "s2", "typ", "Lt", "Asc", "Asm", "As", "sec", "Asch", "vas", "sc", "tau", "st")
            tf = ttk.Frame(f)
            tf.pack(fill="both", expand=True)
            self.stress_tree = ttk.Treeview(tf, columns=cols, show="tree headings", height=11)
            heads = (("N", "N (kN)", 60), ("V2", "V2 (kN)", 60), ("M", "M calc. (kN.m)", 90), ("Le", "L×e (m)", 80),
                     ("s1", "σ1 (MPa)", 62), ("s2", "σ2 (MPa)", 62), ("typ", "Type", 50), ("Lt", "Lt (m)", 48),
                     ("Asc", "As cal. (cm²)", 78), ("Asm", "As min (cm²)", 78), ("As", "As requise (cm²)", 92),
                     ("sec", "Section adoptée ✎", 115), ("Asch", "As choisie (cm²)", 92), ("vas", "As ch. ≥ req.", 75),
                     ("sc", "σc/fbu", 56), ("tau", "τb/τlim", 56), ("st", "Statut", 48))
            for c, t, w in heads:
                self.stress_tree.heading(c, text=t)
                self.stress_tree.column(c, width=w, anchor="center", stretch=(c != "sec"))
            self.stress_tree.heading("#0", text="Niveau")
            self.stress_tree.column("#0", width=85)
            self.stress_tree.tag_configure("ok", background="#e6f4ea")
            self.stress_tree.tag_configure("ko", background="#fde7e9")
            self.stress_tree.bind("<Double-1>", self._stress_edit)
            xs = ttk.Scrollbar(tf, orient="horizontal", command=self.stress_tree.xview)
            self.stress_tree.configure(xscrollcommand=xs.set)
            self.stress_tree.pack(side="top", fill="both", expand=True)
            xs.pack(side="bottom", fill="x")
            self.comp_tree = ttk.Treeview(f, columns=("g", "s", "fbu", "sa", "As", "ng", "sc", "tau"), show="tree headings",
                                          height=3)
            for c, t, w in (("g", "γb", 50), ("s", "γs", 50), ("fbu", "fbu (MPa)", 80), ("sa", "σs (MPa)", 80),
                            ("As", "ΣAs (cm²)", 90), ("ng", "Niveaux NG", 90), ("sc", "σc/fbu max", 90),
                            ("tau", "τ/τlim max", 90)):
                self.comp_tree.heading(c, text=t)
                self.comp_tree.column(c, width=w, anchor="center")
            self.comp_tree.heading("#0", text="Comparaison des cas")
            self.comp_tree.column("#0", width=180)
            self.comp_tree.pack(fill="x", pady=(6, 0))

        def _refresh_cas(self):
            names = [self.vars[f"c{i}_name"].get() for i in (1, 2, 3)]
            self.cas_cb.configure(values=names)

        # ---- moment résistant
        def _tab_mr(self):
            f = ttk.Frame(self.nb, padding=8)
            self.nb.add(f, text="  Moment résistant  ")
            top = ttk.Frame(f)
            top.pack(fill="x")

            def field(box, i, label, key, unit="", combo=None, width=11, btn=None):
                ttk.Label(box, text=label).grid(row=i, column=0, sticky="w", pady=1)
                if combo:
                    self._combo(box, key, combo, i, 1, width=width + 10)
                else:
                    self._entry(box, key, i, 1, width=width)
                ttk.Label(box, text=unit, foreground="gray").grid(row=i, column=2, sticky="w")
                if btn:
                    ttk.Button(box, text=btn[0], command=btn[1]).grid(row=i, column=3, padx=4)

            b1 = ttk.LabelFrame(top, text="Matériaux et situation (BAEL 91 / RPA 2024)", padding=6)
            b1.grid(row=0, column=0, sticky="nsew", padx=4)
            field(b1, 0, "Situation de calcul", "mr_sit", "", MR_SITUATIONS, width=14)
            field(b1, 1, "θ — durée d'application", "mr_theta", "-")
            field(b1, 2, "γb manuel (saisie manuelle)", "mr_gb_m", "-")
            field(b1, 3, "γs manuel (saisie manuelle)", "mr_gs_m", "-")
            field(b1, 4, "fc28", "mr_fc", "MPa")
            field(b1, 5, "fe", "mr_fe", "MPa")
            field(b1, 6, "Es", "mr_Es", "MPa")
            field(b1, 7, "n — coefficient d'équivalence (ELS)", "mr_n", "-")
            b2 = ttk.LabelFrame(top, text="Section rectangulaire (fibre sup. = comprimée)", padding=6)
            b2.grid(row=0, column=1, sticky="nsew", padx=4)
            field(b2, 0, "b — Largeur", "mr_b", "m")
            field(b2, 1, "h — Hauteur", "mr_h", "m")
            field(b2, 2, "d' — cdg A' depuis fibre sup.", "mr_dp", "m")
            field(b2, 3, "c — cdg A depuis fibre inf.", "mr_c", "m")
            field(b2, 4, "A' — Armatures supérieures", "mr_Ap", "cm²",
                  btn=("← Avcalc rive gauche", lambda: self.mr_from_avcalc("Ap")))
            field(b2, 5, "A — Armatures inférieures", "mr_A", "cm²",
                  btn=("← Avcalc rive droite", lambda: self.mr_from_avcalc("A")))
            ttk.Button(b2, text="A' et A ← Avcalc (rive G / rive D)",
                       command=lambda: self.mr_from_avcalc("both")).grid(row=6, column=0, columnspan=4, sticky="w",
                                                                         pady=(4, 0))
            ttk.Label(b2, text="Si A ≠ A', permuter A/A' (et d'/c)\npour le moment de sens opposé.",
                      foreground="gray").grid(row=7, column=0, columnspan=3, sticky="w", pady=(4, 0))
            b3 = ttk.LabelFrame(top, text="Sollicitations et vérification", padding=6)
            b3.grid(row=0, column=2, sticky="nsew", padx=4)
            field(b3, 0, "Calcul à l'ELU ?", "mr_elu", "", ["Oui", "Non"], width=4)
            field(b3, 1, "Nu (compression > 0)", "mr_Nu", "kN",
                  btn=("← −N [P] du 1er niveau", lambda: self.take_first("d")))
            field(b3, 2, "Calcul à l'ELS ?", "mr_els", "", ["Oui", "Non"], width=4)
            field(b3, 3, "Ns (compression > 0)", "mr_Ns", "kN")
            field(b3, 4, "Fissuration", "mr_fiss", "", MR_FISSURATION, width=12)
            field(b3, 5, "η — coefficient de fissuration", "mr_eta", "-")
            field(b3, 6, "Mu (facultatif, vide = non utilisé)", "mr_Mu", "kN.m")
            field(b3, 7, "Ms (facultatif, vide = non utilisé)", "mr_Ms", "kN.m")
            field(b3, 8, "νlim — effort normal réduit limite", "mr_nulim", "-")
            bar = ttk.Frame(f)
            bar.pack(fill="x", pady=(6, 2))
            ttk.Button(bar, text="Reprendre lw, bw, fc28, fe, N, MED des Données générales",
                       command=self.mr_from_general).pack(side="left", padx=2)
            ttk.Button(bar, text="Prendre Mru calculé comme MRD (Données générales)",
                       command=self.mr_to_mrd).pack(side="left", padx=2)
            ttk.Checkbutton(bar, text="Inclure dans la note de calcul", variable=self.vars["mr_report"],
                            onvalue="1", offvalue="0", command=self.schedule).pack(side="left", padx=14)
            self.mr_lbl = ttk.Label(f, text="", foreground="#b3261e")
            self.mr_lbl.pack(anchor="w")
            bot = ttk.Frame(f)
            bot.pack(fill="both", expand=True)
            tf = ttk.Frame(bot)
            tf.pack(side="left", fill="both", expand=True)
            self.mr_tree = ttk.Treeview(tf, columns=("expr", "val", "u"), show="tree headings", selectmode="browse")
            self.mr_tree.heading("#0", text="Résultat")
            self.mr_tree.column("#0", width=290)
            for c, t, w, a in (("expr", "Expression", 200, "w"), ("val", "Valeur", 150, "center"), ("u", "Unité", 50, "center")):
                self.mr_tree.heading(c, text=t)
                self.mr_tree.column(c, width=w, anchor=a)
            sb = ttk.Scrollbar(tf, orient="vertical", command=self.mr_tree.yview)
            self.mr_tree.configure(yscrollcommand=sb.set)
            self.mr_tree.pack(side="left", fill="both", expand=True)
            sb.pack(side="right", fill="y")
            self.mr_tree.tag_configure("sec", background="#1f3a5f", foreground="white", font=("Segoe UI", 9, "bold"))
            self.mr_tree.tag_configure("ok", background="#e6f4ea")
            self.mr_tree.tag_configure("ko", background="#fde7e9")
            self.fig_mr = Figure(figsize=(5.6, 3.6), dpi=100)
            self.canvas_mr = FigureCanvasTkAgg(self.fig_mr, master=bot)
            self.canvas_mr.get_tk_widget().pack(side="left", fill="both", expand=True)

        # ---- élévation
        def _tab_elev(self):
            f = ttk.Frame(self.nb)
            self.nb.add(f, text="  Élévation  ")
            self.el_lbl = ttk.Label(f, text="", font=("Segoe UI", 10, "bold"), anchor="w", padding=(8, 4),
                                    justify="left", wraplength=1100)
            self.el_lbl.pack(fill="x")
            self.fig_el = Figure(figsize=(8, 8.5), dpi=100)
            self.canvas_el = FigureCanvasTkAgg(self.fig_el, master=f)
            self.canvas_el.get_tk_widget().pack(fill="both", expand=True)

        # ---- rives et âme
        def _tab_rive(self):
            f = ttk.Frame(self.nb, padding=12)
            self.nb.add(f, text="  Rives et âme  ")
            box = ttk.LabelFrame(f, text="Éléments de rive", padding=8)
            box.pack(fill="x", padx=4, pady=4)
            ttk.Label(box, text="Désignation", font=("Segoe UI", 9, "bold")).grid(row=0, column=0, sticky="w")
            ttk.Label(box, text="Rive gauche", font=("Segoe UI", 9, "bold")).grid(row=0, column=1)
            ttk.Label(box, text="Rive droite", font=("Segoe UI", 9, "bold")).grid(row=0, column=2)
            for i, (k, lab, u) in enumerate(SIDE_FIELDS, start=1):
                ttk.Label(box, text=lab).grid(row=i, column=0, sticky="w", pady=2)
                self._entry(box, k + "g", i, 1)
                self._entry(box, k + "d", i, 2)
                ttk.Label(box, text=u, foreground="gray").grid(row=i, column=3, sticky="w")
                if k == "bi":
                    ttk.Button(box, text="← entraxe des barres (Schéma, 1er niveau)",
                               command=lambda: self.take_first("e")).grid(row=i, column=4, padx=6)
            ttk.Button(box, text="Copier gauche → droite", command=self.copy_gd).grid(
                row=len(SIDE_FIELDS) + 1, column=1, columnspan=2, pady=6)
            box2 = ttk.LabelFrame(f, text="Âme (partie courante)", padding=8)
            box2.pack(fill="x", padx=4, pady=4)
            for i, (k, lab, u) in enumerate(WEB_FIELDS):
                ttk.Label(box2, text=lab).grid(row=i, column=0, sticky="w", pady=2)
                self._entry(box2, k, i, 1)
                ttk.Label(box2, text=u, foreground="gray").grid(row=i, column=2, sticky="w")

        def _tab_results(self):
            f = ttk.Frame(self.nb)
            self.nb.add(f, text="  Résultats (ductilité)  ")
            cols = ("expr", "g", "d", "u")
            self.tree = ttk.Treeview(f, columns=cols, show="tree headings", selectmode="browse")
            self.tree.heading("#0", text="Désignation")
            for c, t, w in (("expr", "Expression", 270), ("g", "Gauche / Vert.", 130), ("d", "Droite / Hor.", 130),
                            ("u", "Unité", 60)):
                self.tree.heading(c, text=t)
                self.tree.column(c, width=w, anchor="center" if c != "expr" else "w")
            self.tree.column("#0", width=380)
            sb = ttk.Scrollbar(f, orient="vertical", command=self.tree.yview)
            self.tree.configure(yscrollcommand=sb.set)
            self.tree.pack(side="left", fill="both", expand=True)
            sb.pack(side="right", fill="y")
            self.tree.tag_configure("sec", background="#1f3a5f", foreground="white", font=("Segoe UI", 9, "bold"))
            self.tree.tag_configure("ok", background="#e6f4ea")
            self.tree.tag_configure("ko", background="#fde7e9")

        def _tab_schema(self):
            f = ttk.Frame(self.nb)
            self.nb.add(f, text="  Schéma (plan)  ")
            box = ttk.LabelFrame(f, text="Cadres de confinement — types, nombre et longueur (choix de l'ingénieur) "
                                         "→ longueur cumulée L", padding=6)
            box.pack(fill="x", padx=6, pady=4)
            self.cad_lbl, self.cad_tot, self.cad_warn = {}, {}, {}
            for c, (sd, nm) in enumerate((("g", "Rive gauche"), ("d", "Rive droite"))):
                sub = ttk.Frame(box)
                sub.grid(row=0, column=c, padx=10, sticky="n")
                ttk.Label(sub, text=nm, font=("Segoe UI", 9, "bold")).grid(row=0, column=0, columnspan=4)
                for j, h in enumerate(("Type de cadre", "Nombre", "Longueur (cm)", "Σ (cm)")):
                    ttk.Label(sub, text=h, foreground="gray").grid(row=1, column=j)
                for i in range(1, CAD_ROWS + 1):
                    self._combo(sub, f"cad_t{i}{sd}", [""] + CADRE_TYPES, i + 1, 0, width=21)
                    self._entry(sub, f"cad_n{i}{sd}", i + 1, 1, width=7)
                    self._entry(sub, f"cad_l{i}{sd}", i + 1, 2, width=10)
                    lb = ttk.Label(sub, text="", width=9, anchor="e")
                    lb.grid(row=i + 1, column=3)
                    self.cad_lbl[(sd, i)] = lb
                self.cad_tot[sd] = ttk.Label(sub, text="", font=("Segoe UI", 9, "bold"), foreground="#1f4e79")
                self.cad_tot[sd].grid(row=CAD_ROWS + 2, column=0, columnspan=4, sticky="e", pady=(3, 0))
                self.cad_warn[sd] = ttk.Label(sub, text="", foreground="#b3261e")
                self.cad_warn[sd].grid(row=CAD_ROWS + 3, column=0, columnspan=4, sticky="w")
            bt = ttk.Frame(box)
            bt.grid(row=1, column=0, columnspan=2, sticky="w", pady=(6, 0))
            ttk.Button(bt, text="Longueurs automatiques (cases vides)", command=self.cad_auto_len).pack(side="left", padx=3)
            ttk.Button(bt, text="Envoyer L → Rives et âme", command=self.cad_send).pack(side="left", padx=3)
            ttk.Checkbutton(bt, text="Envoi automatique", variable=self.vars["cad_auto"], onvalue="1", offvalue="0",
                            command=self.schedule).pack(side="left", padx=8)
            ttk.Button(bt, text="Effacer", command=self.cad_clear).pack(side="left", padx=3)
            ttk.Label(bt, text="Longueur auto = périmètre (ou portée) à l'axe des cadres + 2 crochets de 10·Øc ; "
                               "modifiable.", foreground="gray").pack(side="left", padx=8)
            self.fig = Figure(figsize=(8, 7), dpi=100)
            self.canvas = FigureCanvasTkAgg(self.fig, master=f)
            self.canvas.get_tk_widget().pack(fill="both", expand=True)

        # ------------------------------------------------------------ données
        def read(self):
            P = {}
            for k in DEFAULTS:
                v = self.vars[k].get().strip()
                if k in TEXT_KEYS:
                    P[k] = v
                    continue
                num = _num(v)
                if num is None:
                    raise ValueError(f"Valeur numérique invalide pour « {k} » : '{v}'")
                P[k] = num
            P["env_auto"] = int(P["env_auto"])
            P["mr_report"] = int(P["mr_report"])
            P["cad_auto"] = int(P["cad_auto"])
            return P

        def read_all(self):
            P = self.read()
            P["stories"] = stories_to_floats(self.raw_stories()) if self.srows else []
            return P

        def schedule(self):
            if self._job:
                self.after_cancel(self._job)
            self._job = self.after(350, self.recalc)

        def recalc(self):
            self._job = None
            self._apply_links()
            try:
                P = self.read_all()
            except (ValueError, ZeroDivisionError) as e:
                self.status.config(text=f"⚠ {e}")
                return
            if P["cad_auto"]:                              # L cumulée des cadres → Rives et âme (si renseignée)
                for sd in ("g", "d"):
                    tot = round(cadres_cumul(P, sd)["total"], 1)
                    if tot > 0 and abs(tot - P["L_conf" + sd]) > 0.05:
                        P["L_conf" + sd] = tot
                        self.vars["L_conf" + sd].set(str(tot))
            try:                                           # moment résistant : indépendant du reste
                self.mr, self.mr_err = compute_mr(P), ""
            except (ValueError, ZeroDivisionError) as e:
                self.mr, self.mr_err = None, str(e)
            try:
                self.res = compute(P)
                self.multi = compute_multi(P) if P["stories"] else None
            except (ValueError, ZeroDivisionError) as e:
                self.status.config(text=f"⚠ {e}")
                self.fill_mr(P)
                return
            self.P = P
            self.status.config(text="Calcul à jour.")
            bad = [n for n, ok in self.res["checks"] if not ok]
            if self.mr and P["mr_report"]:
                bad = [n for n, ok in self.mr["checks"] if not ok] + bad
            if self.multi:
                bad = ([f"section adoptée {r['name']}" for r in self.multi["rows"] if r["ok_as"] is False] + bad)
                bad = [f"contraintes {r['name']}" for r in self.multi["rows"] if not r["ok"]] + bad
            if not bad:
                self.banner.config(text="✔ Toutes les vérifications sont satisfaites", background="#1b7a3d",
                                   foreground="white")
            else:
                self.banner.config(text=f"✖ {len(bad)} vérification(s) non satisfaite(s) : " + " ; ".join(bad[:3]) +
                                   (" …" if len(bad) > 3 else ""), background="#b3261e", foreground="white")
            self.fill_tree()
            draw_schema(self.fig, P, self.res, layout="landscape")
            self.canvas.draw_idle()
            self.fill_multi()
            self.fill_mr(P)
            self.fill_cad(P)

        def fill_mr(self, P):
            self.mr_tree.delete(*self.mr_tree.get_children())
            if not self.mr:
                self.mr_lbl.config(text=f"⚠ {self.mr_err}" if self.mr_err else "")
                self.fig_mr.clear()
                self.canvas_mr.draw_idle()
                return
            self.mr_lbl.config(text="")
            for title, rows in self.mr["sections"][1:]:
                self.mr_tree.insert("", "end", text=title, values=("", "", ""), tags=("sec",))
                for lab, ex, val, un, ok in rows:
                    self.mr_tree.insert("", "end", text="   " + lab, values=(ex, val, un),
                                        tags=() if ok is None else (("ok",) if ok else ("ko",)))
            try:
                draw_mr(self.fig_mr, P, self.mr)
                self.canvas_mr.draw_idle()
            except Exception as e:  # noqa
                self.status.config(text=f"⚠ dessin du moment résistant : {e}")

        def fill_tree(self):
            self.tree.delete(*self.tree.get_children())
            for title, heads, rows in self.res["sections"]:
                self.tree.insert("", "end", text=title, values=("", heads[0], heads[1], ""), tags=("sec",))
                for r in rows:
                    t1, t2 = r.texts()
                    oks = [o for o in (r.ok1, r.ok2) if o is not None]
                    tag = () if not oks else (("ok",) if all(oks) else ("ko",))
                    self.tree.insert("", "end", text="   " + r.label, values=(r.expr, t1, t2, r.unit), tags=tag)

        def fill_multi(self):
            self.env_tree.delete(*self.env_tree.get_children())
            self.v_tree.delete(*self.v_tree.get_children())
            self.stress_tree.delete(*self.stress_tree.get_children())
            self.comp_tree.delete(*self.comp_tree.get_children())
            if not self.multi:
                self.env_lbl.config(text="Aucun étage : importer un fichier ETABS (onglet « Étages / ETABS »).")
                self.v_lbl.config(text="Aucun étage : importer un fichier ETABS (onglet « Étages / ETABS »).")
                self.sum_lbl.config(text="")
                self.send_cb.configure(values=[])
                self.fig_el.clear()
                self.canvas_el.draw_idle()
                self.el_lbl.config(text="Aucun étage : importer un fichier ETABS ou saisir les niveaux.", foreground="gray")
                return
            m, env, S = self.multi, self.multi["env"], self.P["stories"]
            self.env_lbl.config(text=(
                f"n = {env['n']}   hw = {fr(env['hw'])} m   hcr1 = max(lw ; hw/6) = {fr(env['hcr1'])} m   "
                f"limite = {fr(env['limite'])} m\n"
                f"hcr retenue = {fr(env['hcr'])} m   M base = {fr(env['mbase'], 1)}   M sommet = {fr(env['msom'], 1)} kN.m   "
                f"pente = {fr(env['pente'], 1)} kN.m/m   M tête décalée = {fr(env['m_top'], 1)} kN.m"))
            for t in env["table"]:
                lab = "Base du voile (H = 0)" if t["k"] == 0 else S[t["k"] - 1]["name"]
                hh = "" if t["k"] == 0 else fr(S[t["k"] - 1]["h"])
                self.env_tree.insert("", "end", text=lab, values=(hh, fr(t["H"]), fr(t["M_lin"], 1),
                                                                    fr(t["M_corr"], 1), t["zone"]))
            ev = m["envv"]
            self.v_lbl.config(text=(
                f"hcr = {fr(ev['hcr'])} m   hw/3 = {fr(ev['h3'])} m   V(hcr) = {fr(ev['v_hcr'], 1)} kN   "
                f"V(hw/3) = {fr(ev['v_h3'], 1)} kN   facteur = {fr(ev['k'])}\n"
                f"V de calcul pour τb : {V_LABELS.get(m['v_mode'], m['v_mode'])}"))
            for t in ev["table"]:
                self.v_tree.insert("", "end", text=S[t["k"]]["name"], values=(
                    fr(t["H"]), fr(t["Ti"], 1), fr(t["T14"], 1), fr(t["Vc"], 1), t["zone"]))
            for k, r in enumerate(m["rows"]):
                if r["ok_as"] is None:
                    sec_v = (r["sec"], "—", "—")
                else:
                    sec_v = (r["sec"], fr(r["As_ch"], 1), "OK" if r["ok_as"] else "NG")
                tag = "ok" if (r["ok"] and r["ok_as"] is not False) else "ko"
                self.stress_tree.insert("", "end", iid=str(k), text=r["name"], tags=(tag,), values=(
                    fr(r["N"], 0), fr(r["V2"], 0), fr(r["M"], 0), f"{fr(r['L'])}×{fr(r['e'])}", fr(r["s1"]), fr(r["s2"]),
                    r["type"], fr(r["Lt"]), fr(r["As_calc"], 1), fr(r["As_min"], 1), fr(r["As"], 1), *sec_v,
                    fr(r["sc_fbu"], 3), fr(r["tau_lim"], 3), "OK" if r["ok"] else "NG"))
            for c in m["comp"]:
                self.comp_tree.insert("", "end", text=c["cas"]["name"], values=(
                    fr(c["cas"]["gb"]), fr(c["cas"]["gs"]), fr(c["cas"]["fbu"]), fr(c["cas"]["sigs"], 1),
                    fr(c["sumAs"], 1), c["n_ng"], fr(c["max_sc"], 3), fr(c["max_tau"], 3)))
            self.sum_lbl.config(text=f"Cas actif : {m['cas']['name']} — As requise max = {fr(m['As_max'], 1)} cm² — "
                                     f"{m['n_ng']} niveau(x) NG"
                                     + (f" — sections adoptées : {m['n_sec']} ({m['n_as_ko']} insuffisante(s))"
                                        if m["has_sec"] else ""))
            names = [s["name"] for s in S]
            self.send_cb.configure(values=names)
            if self.send_cb.get() not in names:
                self.send_cb.set(names[0])
            el = m["elanc"]
            self.el_lbl.config(
                text=(f"Art. 7.7.4 — hw/lw = {fr(env['hw'])} / {fr(self.P['lw'])} = {fr(el['ratio'])}  "
                      f"{'>' if el['elance'] else '≤'} {fr(el['lim'], 1)}  →  {el['cls'].upper()} : {el['regle']}.\n"
                      f"Moment de calcul retenu : {MODE_LABELS.get(m['mode'], m['mode'])}"
                      + (f"  →  « {MODE_LABELS.get(m['mode_eff'])} »" if m["mode"] == "auto" else "")),
                foreground="#b3261e" if el["elance"] else "#1e6b34")
            try:
                draw_elevation(self.fig_el, self.P, self.res, m)
                self.canvas_el.draw_idle()
            except Exception as e:  # noqa
                self.status.config(text=f"⚠ élévation : {e}")

        # ------------------------------------------------------------ actions
        def send_efforts(self):
            if not self.multi:
                return
            names = [s["name"] for s in self.P["stories"]]
            if self.send_cb.get() not in names:
                return
            s = self.P["stories"][names.index(self.send_cb.get())]
            self.vars["N"].set(str(round(abs(s["Pmin"]), 2)))
            self.vars["MED"].set(str(round(abs(s["M3"]), 2)))
            self.vars["V"].set(str(round(abs(s["V2"]), 2)))
            self.vars["lw"].set(str(s["L"]))
            self.vars["bw"].set(str(s["e"]))
            self.vars["niveau"].set(s["name"])
            self.status.config(text=f"Niveau {s['name']} envoyé : N, MED, V, lw, bw mis à jour "
                                    "(MRD, Avcalc… restent à renseigner depuis le calcul RDM).")

        def _stress_edit(self, event):
            """Double-clic sur la colonne « Section adoptée » : saisie directe (ex. 12HA20, 6HA20+4HA16)."""
            tree = self.stress_tree
            if tree.identify_region(event.x, event.y) != "cell":
                return
            ci = int(tree.identify_column(event.x)[1:]) - 1
            cols = list(tree["columns"])
            iid = tree.identify_row(event.y)
            if ci < 0 or cols[ci] != "sec" or not iid or int(iid) >= len(self.srows):
                return
            var = self.srows[int(iid)]["sec"]
            x, y, w, h = tree.bbox(iid, f"#{ci + 1}")
            ent = ttk.Entry(tree, justify="center")
            ent.insert(0, var.get())
            ent.place(x=x, y=y, width=max(w, 110), height=h)
            ent.focus_set()
            ent.select_range(0, "end")
            done = {"v": False}

            def commit(_e=None, keep=True):
                if done["v"]:
                    return
                done["v"] = True
                txt = ent.get().strip()
                ent.destroy()
                if keep and txt != var.get().strip():
                    var.set(txt)

            ent.bind("<Return>", commit)
            ent.bind("<FocusOut>", commit)
            ent.bind("<Escape>", lambda e: commit(keep=False))

        def propose_sections(self):
            if not self.multi:
                return
            phi = _num(self.phi_cb.get())
            n = 0
            for k, r in enumerate(self.multi["rows"]):
                if k < len(self.srows) and not self.srows[k]["sec"].get().strip():
                    self.srows[k]["sec"].set(propose_section(r["As"], phi))
                    n += 1
            self.status.config(text=f"{n} section(s) proposée(s) avec Ø{phi:g} (cases vides uniquement) — modifiables.")

        def clear_sections(self):
            for r in self.srows:
                r["sec"].set("")
            self.status.config(text="Sections adoptées effacées.")

        def mr_from_general(self):
            for dst, src in (("mr_b", "bw"), ("mr_h", "lw"), ("mr_fc", "fc"), ("mr_fe", "fe"), ("mr_Nu", "N"),
                             ("mr_Mu", "MED")):
                self.vars[dst].set(self.vars[src].get())
            self.status.config(text="Moment résistant : b ← bw, h ← lw, fc28, fe, Nu ← N, Mu ← MED (Données générales).")

        def refresh(self):
            """Actualise immédiatement tous les calculs, tableaux et schémas."""
            if self._job:
                self.after_cancel(self._job)
            self.recalc()

        def fill_cad(self, P):
            for sd in ("g", "d"):
                c = cadres_cumul(P, sd)
                subs = {r["i"]: r["sub"] for r in c["rows"]}
                for i in range(1, CAD_ROWS + 1):
                    self.cad_lbl[(sd, i)].config(text=fr(subs[i], 1) if i in subs else "")
                self.cad_tot[sd].config(text=f"L cumulée = {fr(c['total'], 1)} cm" if c["rows"] else "")
                self.cad_warn[sd].config(text=("⚠ " + " ; ".join(c["warn"])) if c["warn"] else "")

        def cad_auto_len(self):
            if not self.res:
                return
            n = 0
            for sd, r in (("g", self.res["G"]), ("d", self.res["D"])):
                for i in range(1, CAD_ROWS + 1):
                    if self.vars[f"cad_l{i}{sd}"].get().strip():
                        continue
                    ln = cadre_longueur_auto(self.vars[f"cad_t{i}{sd}"].get(), r)
                    if ln is not None:
                        self.vars[f"cad_l{i}{sd}"].set(str(round(ln, 1)))
                        n += 1
            self.status.config(text=f"{n} longueur(s) de cadre proposée(s) (cases vides uniquement).")
            self.refresh()

        def cad_send(self):
            try:
                P = self.read()
            except ValueError as e:
                self.status.config(text=f"⚠ {e}")
                return
            done = []
            for sd, nm in (("g", "gauche"), ("d", "droite")):
                tot = cadres_cumul(P, sd)["total"]
                if tot > 0:
                    self.vars["L_conf" + sd].set(str(round(tot, 1)))
                    done.append(f"rive {nm} : L = {fr(tot, 1)} cm")
            self.status.config(text=("Envoyé vers Rives et âme — " + " ; ".join(done)) if done
                               else "Aucune longueur de cadre renseignée : rien à envoyer.")
            self.refresh()

        def cad_clear(self):
            for sd in ("g", "d"):
                for i in range(1, CAD_ROWS + 1):
                    for k in "tnl":
                        self.vars[f"cad_{k}{i}{sd}"].set("")
            self.refresh()

        def take_first(self, which):
            """Reprend des valeurs de la 1re ligne du tableau Étages / ETABS vers les données de calcul."""
            if not self.srows:
                messagebox.showinfo("Étages / ETABS", "Aucun niveau dans le tableau.")
                return
            first = self.srows[0]
            val = {k: _num(first[k].get()) for k in ("M3", "Pmin", "V2", "N")}
            bad = [k for k, v in val.items() if v is None]
            if bad:
                messagebox.showerror("Étages / ETABS", f"Valeur invalide au 1er niveau : {', '.join(bad)}")
                return
            msg = []
            if which in ("a", "all"):
                self.vars["MED"].set(str(round(abs(val["M3"]), 2)))
                msg.append(f"MED = {fr(abs(val['M3']), 1)} kN.m")
            if which in ("b", "all"):
                self.vars["N"].set(str(round(abs(val["Pmin"]), 2)))
                msg.append(f"N = {fr(abs(val['Pmin']), 1)} kN")
            if which in ("c", "all"):
                self.vars["V"].set(str(round(abs(val["V2"]), 2)))
                msg.append(f"V = {fr(abs(val['V2']), 1)} kN")
            if which in ("d", "all"):
                self.vars["mr_Nu"].set(str(round(-val["N"], 2)))
                msg.append(f"Nu = {fr(-val['N'], 1)} kN")
            if which in ("e", "all"):
                if not self.res:
                    messagebox.showinfo("bi", "Calcul indisponible : corriger d'abord les données.")
                else:
                    for sd, r in (("g", self.res["G"]), ("d", self.res["D"])):
                        bi = max(r["lay"]["el"], r["lay"]["eb"])
                        self.vars["bi" + sd].set(str(round(bi, 2)))
                    msg.append(f"bi = {fr(max(self.res['G']['lay']['el'], self.res['G']['lay']['eb']))} / "
                               f"{fr(max(self.res['D']['lay']['el'], self.res['D']['lay']['eb']))} cm")
            self.status.config(text=f"1er niveau ({first['name'].get()}) → " + " ; ".join(msg))
            self.refresh()

        def mr_from_avcalc(self, which):
            """A' (armatures supérieures) ← Avcalc de la rive gauche ; A (armatures inférieures) ← Avcalc de la rive droite."""
            g, d = _num(self.vars["Avcalcg"].get()), _num(self.vars["Avcalcd"].get())
            msg = []
            if which in ("Ap", "both"):
                if g is None:
                    messagebox.showerror("Avcalc", "Avcalc (rive gauche) invalide ou vide — rubrique « Rives et âme ».")
                    return
                self.vars["mr_Ap"].set(str(round(g, 2)))
                msg.append(f"A' = {fr(g, 2)} cm² (rive gauche)")
            if which in ("A", "both"):
                if d is None:
                    messagebox.showerror("Avcalc", "Avcalc (rive droite) invalide ou vide — rubrique « Rives et âme ».")
                    return
                self.vars["mr_A"].set(str(round(d, 2)))
                msg.append(f"A = {fr(d, 2)} cm² (rive droite)")
            self.status.config(text="Moment résistant ← Avcalc : " + " ; ".join(msg))
            self.refresh()

        def open_help(self):
            if not HELP_URL:
                messagebox.showinfo("Aide", "Le site d'aide n'est pas encore en ligne.\n\nL'adresse sera renseignée "
                                            "dans la constante HELP_URL du programme (ou la variable d'environnement "
                                            "VOILE_HELP_URL).")
                return
            webbrowser.open(HELP_URL)

        def show_contact(self):
            c = CONTACT
            win = tk.Toplevel(self)
            win.title("Contact — " + c["bureau"])
            win.resizable(False, False)
            fr_ = ttk.Frame(win, padding=16)
            fr_.pack()
            ttk.Label(fr_, text=c["bureau"], font=("Segoe UI", 14, "bold")).grid(row=0, column=0, columnspan=2, sticky="w")
            ttk.Label(fr_, text="Bureau d'études — structures", foreground="gray").grid(row=1, column=0, columnspan=2,
                                                                                         sticky="w", pady=(0, 8))
            rows = [("Ingénieur", c["ingenieur"]), ("E-mail", c["mail"])] + \
                   [("Téléphone" if i == 0 else "", t) for i, t in enumerate(c["tels"])]
            for i, (a, b) in enumerate(rows, start=2):
                ttk.Label(fr_, text=a, foreground="gray").grid(row=i, column=0, sticky="w", padx=(0, 12), pady=1)
                ttk.Label(fr_, text=b).grid(row=i, column=1, sticky="w", pady=1)
            bb = ttk.Frame(fr_)
            bb.grid(row=len(rows) + 2, column=0, columnspan=2, pady=(12, 0), sticky="w")
            ttk.Button(bb, text="Écrire un e-mail", command=lambda: webbrowser.open("mailto:" + c["mail"])).pack(side="left")

            def copy():
                self.clipboard_clear()
                self.clipboard_append(c["mail"])
                self.status.config(text="Adresse e-mail copiée.")
            ttk.Button(bb, text="Copier l'e-mail", command=copy).pack(side="left", padx=6)
            ttk.Button(bb, text="Fermer", command=win.destroy).pack(side="left")

        # ---- dimensions des voiles (ETABS « Pier Section Properties »)
        def _tab_dims(self):
            f = ttk.Frame(self.nb, padding=8)
            self.nb.add(f, text="  Dimensions (ETABS)  ")
            self.dims = None
            top = ttk.Frame(f)
            top.pack(fill="x")
            ttk.Button(top, text="Importer « Pier Section Properties » (xlsx / csv)…", command=self.dim_import).pack(side="left")
            ttk.Label(top, text="  Pier :").pack(side="left")
            self.dim_pier = tk.StringVar()
            self.dim_cb = ttk.Combobox(top, textvariable=self.dim_pier, values=[], width=10, state="readonly")
            self.dim_cb.pack(side="left", padx=4)
            self.dim_cb.bind("<<ComboboxSelected>>", lambda *_: self.dim_fill_pier())
            self.dim_lbl = ttk.Label(top, text="Aucun fichier importé.", foreground="gray")
            self.dim_lbl.pack(side="left", padx=10)
            bar = ttk.Frame(f)
            bar.pack(fill="x", pady=4)
            ttk.Button(bar, text="Appliquer h, L, e du pier aux niveaux (Étages / ETABS)",
                       command=self.dim_apply).pack(side="left", padx=2)
            ttk.Button(bar, text="Créer les niveaux depuis ce pier", command=self.dim_create).pack(side="left", padx=2)
            ttk.Button(bar, text="lw, bw, he (Données générales) ← niveau le plus bas du pier",
                       command=self.dim_to_general).pack(side="left", padx=2)
            ttk.Button(bar, text="Exporter les tableaux (Excel)…", command=self.dim_export).pack(side="left", padx=2)
            self.dim_info = ttk.Label(f, text="", font=("Segoe UI", 10, "bold"), anchor="w", justify="left")
            self.dim_info.pack(fill="x", pady=2)
            nbk = ttk.Notebook(f)
            nbk.pack(fill="both", expand=True)

            def tree(parent, title, cols):
                fr2 = ttk.Frame(parent)
                nbk.add(fr2, text=title)
                tv = ttk.Treeview(fr2, columns=[c[0] for c in cols], show="headings")
                for k, t, w in cols:
                    tv.heading(k, text=t)
                    tv.column(k, width=w, anchor="center")
                sb = ttk.Scrollbar(fr2, orient="vertical", command=tv.yview)
                tv.configure(yscrollcommand=sb.set)
                tv.pack(side="left", fill="both", expand=True)
                sb.pack(side="right", fill="y")
                tv.tag_configure("el", background="#fde7e9")
                tv.tag_configure("tot", background="#e8eef7", font=("Segoe UI", 9, "bold"))
                return tv

            self.dim_t_pier = tree(nbk, "Pier sélectionné", [("n", "Niveau", 110), ("zb", "Cote base (m)", 100),
                                   ("zt", "Cote sommet (m)", 110), ("h", "h (m)", 80), ("L", "L (m)", 80),
                                   ("e", "e (m)", 80), ("S", "L × e (m²)", 100)])
            self.dim_t_lev = tree(nbk, "Synthèse par niveau", [("n", "Niveau", 110), ("zb", "Cote base", 90),
                                  ("zt", "Cote sommet", 90), ("h", "h (m)", 70), ("np", "Nb piers", 70),
                                  ("sL", "Σ L (m)", 90), ("e1", "e min", 70), ("e2", "e max", 70),
                                  ("sS", "Σ L×e (m²)", 100)])
            self.dim_t_rec = tree(nbk, "Récapitulatif par pier", [("p", "Pier", 70), ("n", "Nb niv.", 70),
                                  ("L1", "L min", 70), ("L2", "L max", 70), ("e1", "e min", 70), ("e2", "e max", 70),
                                  ("hw", "hw = Σh (m)", 90), ("lw", "lw base (m)", 90), ("r", "hw/lw", 70),
                                  ("c", "Art. 7.7.4", 110), ("S", "Σ L×e (m²)", 100)])

        def dim_import(self):
            path = filedialog.askopenfilename(
                filetypes=[("Pier Section Properties", "*.xlsx *.xlsm *.csv *.txt"), ("Tous", "*.*")])
            if not path:
                return
            try:
                self.dims = load_pier_sections(path)
                T = dimension_tables(self.dims)
            except Exception as e:  # noqa
                messagebox.showerror("Dimensions", str(e))
                return
            self.dim_cb.configure(values=self.dims["piers"])
            if self.dim_pier.get() not in self.dims["piers"]:
                self.dim_pier.set(self.vars["pier"].get() if self.vars["pier"].get() in self.dims["piers"]
                                  else self.dims["piers"][0])
            self.dim_lbl.config(text=f"{os.path.basename(path)} — {len(self.dims['piers'])} piers, "
                                     f"{len(self.dims['stories'])} niveaux")
            self.dim_info.config(text=f"Hauteur totale : {fr(T['hw_total'])} m   —   Σ surface L×e : "
                                      f"{fr(T['S_total'], 3)} m²")
            for tv in (self.dim_t_lev, self.dim_t_rec):
                tv.delete(*tv.get_children())
            for l in T["levels"]:
                self.dim_t_lev.insert("", "end", values=(l["name"], fr(l["zb"]), fr(l["zt"]), fr(l["h"]), l["n"],
                                                         fr(l["sumL"]), fr(l["emin"]), fr(l["emax"]), fr(l["sumS"], 3)))
            for q in T["recap"]:
                self.dim_t_rec.insert("", "end", tags=("el",) if q["elance"] else (),
                                      values=(q["pier"], q["n"], fr(q["Lmin"]), fr(q["Lmax"]), fr(q["emin"]),
                                              fr(q["emax"]), fr(q["hw"]), fr(q["lw0"]), fr(q["ratio"]),
                                              "élancé" if q["elance"] else "court", fr(q["S"], 3)))
            self.dim_fill_pier()
            self.status.config(text="Dimensions importées : L = max(largeur bas, haut), e = max(épaisseur bas, haut), "
                                    "h = cote sommet − cote base.")

        def dim_fill_pier(self):
            self.dim_t_pier.delete(*self.dim_t_pier.get_children())
            if not self.dims or not self.dim_pier.get():
                return
            for r in pier_levels(self.dims, self.dim_pier.get()):
                self.dim_t_pier.insert("", "end", values=(r["story"], fr(r["zb"]), fr(r["zt"]), fr(r["h"]), fr(r["L"]),
                                                          fr(r["e"]), fr(r["S"], 3)))

        def _dim_ok(self):
            if not self.dims or not self.dim_pier.get():
                messagebox.showinfo("Dimensions", "Importer d'abord un fichier « Pier Section Properties » "
                                                  "et choisir un pier.")
                return False
            return True

        def dim_apply(self):
            if not self._dim_ok():
                return
            raw = self.raw_stories()
            if not raw:
                messagebox.showinfo("Dimensions", "Aucun niveau : utiliser « Créer les niveaux depuis ce pier ».")
                return
            n = apply_dims_to_stories(raw, self.dims, self.dim_pier.get())
            self.set_stories(raw)
            self.status.config(text=f"h, L, e du pier {self.dim_pier.get()} appliqués à {n} niveau(x) sur {len(raw)} "
                                    f"(correspondance par nom de niveau).")

        def dim_create(self):
            if not self._dim_ok():
                return
            self.set_stories(stories_from_dims(self.dims, self.dim_pier.get()))
            self.status.config(text=f"Niveaux créés depuis le pier {self.dim_pier.get()} (efforts à 0 : importer "
                                    f"l'export d'efforts ETABS puis « Charger les efforts du pier »).")

        def dim_to_general(self):
            if not self._dim_ok():
                return
            r = pier_levels(self.dims, self.dim_pier.get())[0]
            self.vars["lw"].set(str(r["L"]))
            self.vars["bw"].set(str(r["e"]))
            self.vars["he"].set(str(r["h"]))
            self.status.config(text=f"lw = {fr(r['L'])} m, bw = {fr(r['e'])} m, he = {fr(r['h'])} m "
                                    f"(niveau {r['story']}, pier {self.dim_pier.get()}).")
            self.refresh()

        def dim_export(self):
            if not self.dims:
                messagebox.showinfo("Dimensions", "Importer d'abord un fichier « Pier Section Properties ».")
                return
            path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="Dimensions_voiles_ETABS.xlsx",
                                                filetypes=[("Excel", "*.xlsx")])
            if path:
                try:
                    export_dimensions_xlsx(self.dims, path)
                    self.status.config(text=f"Tableaux exportés : {path}")
                except Exception as e:  # noqa
                    messagebox.showerror("Export", str(e))

        def mr_to_mrd(self):
            if self.mr and self.mr["elu"]:
                self.vars["MRD"].set(str(round(self.mr["elu"]["Mru"], 1)))
                self.status.config(text=f"MRD ← Mru = {fr(self.mr['elu']['Mru'], 1)} kN.m (Données générales).")
                self.refresh()
            else:
                messagebox.showinfo("Moment résistant", "Aucun Mru disponible (calcul ELU désactivé ou invalide).")

        def send_hcr(self):
            if self.multi:
                self.vars["hcr"].set(str(round(self.multi["env"]["hcr"], 3)))
                self.status.config(text="hcr de l'enveloppe envoyé aux Données générales.")

        def copy_gd(self):
            for k, *_ in SIDE_FIELDS:
                self.vars[k + "d"].set(self.vars[k + "g"].get())

        def _update_title(self):
            base = "Voile RPA 2024 — Enveloppes M / V, contraintes, moment résistant, ferraillage et ductilité locale"
            self.title(base + (f"  —  {os.path.basename(self.path)}" if self.path else ""))

        def new(self):
            for k, v in DEFAULTS.items():
                self.vars[k].set(str(v))
            self.set_stories([])
            self._sync_combos()
            self.path = None
            self._update_title()

        def _write(self, path):
            data = {k: self.vars[k].get() for k in DEFAULTS}
            data["stories"] = self.raw_stories()
            with open(path, "w", encoding="utf-8") as fh:
                json.dump(data, fh, indent=2, ensure_ascii=False)
            self.path = path
            self._update_title()
            self.status.config(text=f"Projet enregistré : {path}")

        def save(self):
            """Enregistre dans le fichier courant ; à défaut, demande un nom (comme « Enregistrer sous »)."""
            if self.path:
                try:
                    self._write(self.path)
                except OSError as e:
                    messagebox.showerror("Enregistrer", str(e))
            else:
                self.save_as()

        def save_as(self):
            """Enregistre sous un nouveau nom : le fichier d'origine (modèle) reste intact."""
            init = os.path.basename(self.path) if self.path else "projet_voile.json"
            path = filedialog.asksaveasfilename(defaultextension=".json", initialfile=init,
                                                filetypes=[("Projet voile", "*.json")])
            if path:
                try:
                    self._write(path)
                except OSError as e:
                    messagebox.showerror("Enregistrer sous", str(e))

        def open(self):
            path = filedialog.askopenfilename(filetypes=[("Projet voile", "*.json")])
            if path:
                with open(path, encoding="utf-8") as fh:
                    data = json.load(fh)
                for k, v in data.items():
                    if k in self.vars:
                        self.vars[k].set(str(v))
                if "seis_link" not in data:            # projet antérieur : ses valeurs R, Qf, T0, fc… sont conservées
                    self.vars["seis_link"].set("0")
                if "mat_link" not in data:
                    self.vars["mat_link"].set("0")
                self.set_stories(data.get("stories", []))
                self._sync_combos()
                self._refresh_cas()
                self.path = path
                self._update_title()
                self.status.config(text=f"Projet ouvert : {path}")

        def make_report(self, fmt):
            try:
                P = self.read_all()
                res = compute(P)
                multi = compute_multi(P) if P["stories"] else None
            except (ValueError, ZeroDivisionError) as e:
                messagebox.showerror("Données invalides", str(e))
                return
            if not (res["all_ok"] and (multi is None or multi["all_ok"])) and not messagebox.askyesno(
                    "Vérifications non satisfaites",
                    "Certaines vérifications ne sont pas satisfaites.\nGénérer quand même la note de calcul ?"):
                return
            ext = ".pdf" if fmt == "pdf" else ".docx"
            init = f"Note_calcul_{P['element'].replace(' ', '_')}{ext}"
            path = filedialog.asksaveasfilename(
                defaultextension=ext, initialfile=init,
                filetypes=[("PDF", "*.pdf")] if fmt == "pdf" else [("Document Word", "*.docx")])
            if not path:
                return
            try:
                (build_pdf if fmt == "pdf" else build_docx)(P, res, path, multi)
            except Exception as e:  # noqa
                messagebox.showerror("Erreur lors de la génération", str(e))
                return
            self.status.config(text=f"Note de calcul générée : {path}")
            if messagebox.askyesno("Note de calcul", "Note de calcul générée.\nL'ouvrir maintenant ?"):
                try:
                    if sys.platform.startswith("win"):
                        os.startfile(path)  # type: ignore
                    elif sys.platform == "darwin":
                        os.system(f'open "{path}"')
                    else:
                        os.system(f'xdg-open "{path}"')
                except Exception:
                    pass

    return App


# ---------------------------------------------------------------------------
#  7. POINT D'ENTRÉE
# ---------------------------------------------------------------------------
def _arg(argv, name):
    return argv[argv.index(name) + 1] if name in argv and argv.index(name) + 1 < len(argv) else None


def main(argv):
    if "--pdf" in argv or "--docx" in argv:                 # mode console (sans interface)
        P = dict(DEFAULTS)
        pj = _arg(argv, "--projet")
        if pj:
            with open(pj, encoding="utf-8") as fh:
                data = json.load(fh)
            for k, v in data.items():
                if k == "stories":
                    P["stories"] = stories_to_floats(v)
                elif k in DEFAULTS:
                    P[k] = v if k in TEXT_KEYS else _num(v)
            for k in ("seis_link", "mat_link"):          # projet antérieur sans hypothèses : on garde ses valeurs
                if k not in data:
                    P[k] = 0
        et = _arg(argv, "--etabs")
        if et:
            data = load_etabs(et)
            pier = _arg(argv, "--pier") or data["piers"][0]
            P["etabs_path"], P["pier"] = et, pier
            P["stories"] = stories_from_etabs(data, pier, P)
        dm = _arg(argv, "--dims")                           # dimensions L, e, h depuis « Pier Section Properties »
        if dm and P.get("stories"):
            dd = load_pier_sections(dm)
            pr = _arg(argv, "--pier") or P.get("pier") or dd["piers"][0]
            apply_dims_to_stories(P["stories"], dd, pr)
        P["env_auto"] = int(P["env_auto"])
        P["mr_report"] = int(P["mr_report"])
        P["cad_auto"] = int(P["cad_auto"])
        P["seis_link"], P["mat_link"] = int(P["seis_link"]), int(P["mat_link"])
        apply_links(P)
        for sd in ("g", "d"):                               # L cumulée des cadres, si renseignée
            tot = cadres_cumul(P, sd)["total"]
            if P["cad_auto"] and tot > 0:
                P["L_conf" + sd] = round(tot, 1)
        res = compute(P)
        multi = compute_multi(P) if P.get("stories") else None
        ok = res["all_ok"] and (multi is None or multi["all_ok"])
        for opt, fn in (("--pdf", build_pdf), ("--docx", build_docx)):
            out = _arg(argv, opt)
            if out:
                fn(P, res, out, multi)
                print(f"Note de calcul : {out}")
        print("Toutes les vérifications sont satisfaites." if ok else "Des vérifications ne sont pas satisfaites.")
        return
    run_gui()().mainloop()


if __name__ == "__main__":
    main(sys.argv[1:])
